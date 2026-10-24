"""
============================================================
MRS Speech Analytics Professional
Vista: Dashboard
============================================================
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional
from collections import Counter
import json
import shutil
import re
from datetime import datetime

import customtkinter as ctk
from tkinter import messagebox

from app.core import project_manager
from app.services.dashboard_service import DashboardService
from app.ui.themes import Theme
from app.ui.widgets.chart_card import ChartCard


class DashboardView(ctk.CTkFrame):
    """
    Dashboard general del proyecto activo.

    Los datos provienen de los análisis previamente guardados
    en la carpeta 04_Reportes/Analytics.
    """

    def __init__(self, parent):
        super().__init__(
            parent,
            fg_color=Theme.colors.BACKGROUND,
            corner_radius=0,
        )

        self.dashboard_data = {}
        self.filter_gestor = "TODOS"
        self.filter_supervisor = "TODOS"
        self.filter_resultado = "TODOS"
        self.filter_auditoria = "TODOS"
        self.filter_score = "TODOS"
        self.detail_selection = None

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(2, weight=1)

        self._build_header()
        self._build_tabs()
        self._build_content()
        self._build_detail_content()

        self.load_dashboard()

    # =========================================================
    # HEADER
    # =========================================================

    def _build_header(self):
        self.header = ctk.CTkFrame(
            self,
            fg_color=Theme.colors.SURFACE,
            corner_radius=15,
            border_width=1,
            border_color=Theme.colors.BORDER_LIGHT,
        )

        self.header.grid(
            row=0,
            column=0,
            padx=20,
            pady=(16, 8),
            sticky="ew",
        )

        self.header.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            self.header,
            text="DASHBOARD  |  VISIÓN GENERAL",
            text_color=Theme.colors.PRIMARY,
            font=ctk.CTkFont(
                family="Segoe UI",
                size=24,
                weight="bold",
            ),
            anchor="w",
        ).grid(
            row=0,
            column=0,
            padx=24,
            pady=(15, 3),
            sticky="w",
        )

        self.project_label = ctk.CTkLabel(
            self.header,
            text="Proyecto: Sin proyecto abierto",
            text_color=Theme.colors.TEXT_SECONDARY,
            font=ctk.CTkFont(
                family="Segoe UI",
                size=12,
            ),
            anchor="w",
        )

        self.project_label.grid(
            row=1,
            column=0,
            padx=24,
            pady=(0, 15),
            sticky="w",
        )

        self.refresh_button = ctk.CTkButton(
            self.header,
            text="Actualizar datos",
            width=126,
            height=36,
            fg_color=Theme.colors.PRIMARY,
            hover_color=Theme.colors.PRIMARY_HOVER,
            corner_radius=9,
            command=self.load_dashboard,
        )

        self.refresh_button.grid(
            row=0,
            column=1,
            rowspan=2,
            padx=20,
            pady=15,
            sticky="e",
        )

    # =========================================================
    # NAVEGACIÓN DEL DASHBOARD
    # =========================================================

    def _build_tabs(self):
        self.tab_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.tab_frame.grid(row=1, column=0, padx=24, pady=(0, 7), sticky="ew")

        self.tab_selector = ctk.CTkSegmentedButton(
            self.tab_frame,
            values=["VISIÓN GENERAL", "GESTORES", "SUPERVISORES", "CALIDAD", "PALABRAS", "ESTATUS", "SALIDAS"],
            command=self._change_tab,
            fg_color=Theme.colors.SURFACE,
            selected_color=Theme.colors.PRIMARY,
            selected_hover_color=Theme.colors.PRIMARY_HOVER,
            unselected_color=Theme.colors.SURFACE,
            unselected_hover_color=Theme.colors.CARD_PURPLE,
            text_color=Theme.colors.TEXT_PRIMARY,
            corner_radius=9,
            height=34,
            font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
        )
        self.tab_selector.pack(side="left")
        self.tab_selector.set("VISIÓN GENERAL")

    def _change_tab(self, value):
        if value == "VISIÓN GENERAL":
            self.detail_content.grid_remove()
            self.content.grid()
            return

        self.content.grid_remove()
        self.detail_content.grid()
        self._render_detail_tab(value)

    def _build_detail_content(self):
        self.detail_content = ctk.CTkScrollableFrame(self, fg_color="transparent")
        self.detail_content.grid(row=2, column=0, padx=20, pady=(0, 12), sticky="nsew")
        self.detail_content.grid_columnconfigure(0, weight=1)
        self.detail_content.grid_remove()

    def _clear_detail_content(self):
        for widget in self.detail_content.winfo_children():
            widget.destroy()

    def _section_title(self, parent, title, subtitle=""):
        card = ctk.CTkFrame(
            parent, fg_color=Theme.colors.SURFACE, corner_radius=14,
            border_width=1, border_color=Theme.colors.BORDER_LIGHT,
        )
        card.pack(fill="x", padx=5, pady=(0, 10))
        ctk.CTkLabel(
            card, text=title, text_color=Theme.colors.TEXT_PRIMARY,
            font=ctk.CTkFont(family="Segoe UI", size=18, weight="bold"), anchor="w",
        ).pack(fill="x", padx=18, pady=(13, 2))
        if subtitle:
            ctk.CTkLabel(
                card, text=subtitle, text_color=Theme.colors.TEXT_SECONDARY,
                font=ctk.CTkFont(family="Segoe UI", size=11), anchor="w",
            ).pack(fill="x", padx=18, pady=(0, 12))
        return card

    def _metric_strip(self, parent, values):
        frame = ctk.CTkFrame(parent, fg_color="transparent")
        frame.pack(fill="x", padx=0, pady=(0, 10))
        for i, (title, value, bg) in enumerate(values):
            frame.grid_columnconfigure(i, weight=1, uniform="detailkpi")
            card = ctk.CTkFrame(frame, fg_color=bg, corner_radius=13,
                                border_width=1, border_color=Theme.colors.BORDER_LIGHT)
            card.grid(row=0, column=i, padx=5, sticky="nsew")
            ctk.CTkLabel(card, text=title, text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold"),
                         anchor="w").pack(fill="x", padx=15, pady=(12, 2))
            ctk.CTkLabel(card, text=str(value), text_color=Theme.colors.TEXT_PRIMARY,
                         font=ctk.CTkFont(family="Segoe UI", size=20, weight="bold"),
                         anchor="w").pack(fill="x", padx=15, pady=(0, 12))

    def _ranking_table(self, parent, title, rows, name_header):
        card = ctk.CTkFrame(parent, fg_color=Theme.colors.SURFACE, corner_radius=14,
                            border_width=1, border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x", padx=5, pady=(0, 12))
        ctk.CTkLabel(card, text=title, text_color=Theme.colors.TEXT_PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=15, weight="bold"),
                     anchor="w").pack(fill="x", padx=16, pady=(13, 8))
        header = ctk.CTkFrame(card, fg_color=Theme.colors.CARD_PURPLE, corner_radius=8)
        header.pack(fill="x", padx=12, pady=(0, 4))
        for col, text, width in [(0, name_header, 280), (1, "LLAMADAS", 90),
                                 (2, "PROMEDIO", 90), (3, "CUMPLE", 90),
                                 (4, "PARCIAL", 90), (5, "NO CUMPLE", 90)]:
            header.grid_columnconfigure(col, weight=1 if col == 0 else 0, minsize=width)
            ctk.CTkLabel(header, text=text, text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold"),
                         anchor="w" if col == 0 else "center").grid(
                             row=0, column=col, padx=9, pady=7, sticky="ew")
        if not rows:
            ctk.CTkLabel(card, text="Sin datos disponibles", text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=11)).pack(pady=28)
        for idx, row in enumerate(rows[:20]):
            line = ctk.CTkFrame(card, fg_color="transparent", corner_radius=0)
            line.pack(fill="x", padx=12)
            vals = [row["name"], row["calls"], f'{row["avg"]:.0f}%',
                    row["cumple"], row["parcial"], row["no_cumple"]]
            for col, val in enumerate(vals):
                line.grid_columnconfigure(col, weight=1 if col == 0 else 0,
                                           minsize=280 if col == 0 else 90)
                ctk.CTkLabel(line, text=str(val), text_color=Theme.colors.TEXT_PRIMARY,
                             font=ctk.CTkFont(family="Segoe UI", size=10,
                                              weight="bold" if col == 0 else "normal"),
                             anchor="w" if col == 0 else "center").grid(
                                 row=0, column=col, padx=9, pady=7, sticky="ew")
            if idx < min(len(rows), 20) - 1:
                ctk.CTkFrame(card, height=1, fg_color=Theme.colors.BORDER_LIGHT).pack(fill="x", padx=16)
        ctk.CTkLabel(card, text="Ordenado por número de llamadas analizadas",
                     text_color=Theme.colors.TEXT_SECONDARY,
                     font=ctk.CTkFont(family="Segoe UI", size=9), anchor="w").pack(
                         fill="x", padx=16, pady=(7, 11))

    def _render_detail_tab(self, tab):
        self._clear_detail_content()
        data = getattr(self, "operational_data", {}) or {}
        if not data:
            self._section_title(self.detail_content, tab, "Abra un proyecto con Analytics para consultar esta sección.")
            return

        self._build_filter_bar(tab)
        filtered = self._filtered_operational_data()

        if tab == "GESTORES":
            rows = self._aggregate_records(filtered.get("records", []), "gestor")
            self._section_title(self.detail_content, "GESTORES | DESEMPEÑO",
                                "Seleccione un gestor para consultar su calidad, distribución y penalizaciones.")
            self._metric_strip(self.detail_content, [
                ("GESTORES EVALUADOS", len(rows), Theme.colors.CARD_BLUE),
                ("LLAMADAS FILTRADAS", filtered.get("total", 0), Theme.colors.CARD_PURPLE),
                ("PROMEDIO", f'{filtered.get("average", 0):.0f}%', Theme.colors.CARD_GREEN),
                ("CON PENALIZACIÓN", filtered.get("penalized_calls", 0), Theme.colors.CARD_ORANGE),
            ])
            self._build_entity_selector("GESTOR", rows, filtered, "gestor")
            self._ranking_table(self.detail_content, "Detalle por gestor", rows, "GESTOR")

        elif tab == "SUPERVISORES":
            rows = self._aggregate_records(filtered.get("records", []), "supervisor")
            self._section_title(self.detail_content, "SUPERVISORES | VISIÓN DE EQUIPO",
                                "Seleccione un supervisor para consultar sus gestores y resultados de calidad.")
            self._metric_strip(self.detail_content, [
                ("SUPERVISORES", len(rows), Theme.colors.CARD_BLUE),
                ("LLAMADAS FILTRADAS", filtered.get("total", 0), Theme.colors.CARD_PURPLE),
                ("PROMEDIO", f'{filtered.get("average", 0):.0f}%', Theme.colors.CARD_GREEN),
                ("NO CUMPLE", filtered.get("no_cumple", 0), Theme.colors.CARD_ORANGE),
            ])
            self._build_entity_selector("SUPERVISOR", rows, filtered, "supervisor")
            self._ranking_table(self.detail_content, "Detalle por supervisor", rows, "SUPERVISOR")

        elif tab == "PALABRAS":
            self._render_word_audit()

        elif tab == "ESTATUS":
            self._render_excel_status_audit()

        elif tab == "SALIDAS":
            self._render_final_outputs()

        elif tab == "CALIDAD":
            self._section_title(self.detail_content, "CALIDAD | PUNTOS DE CONTROL",
                                "Los filtros superiores se aplican a bloques, penalizaciones e indicadores.")
            self._metric_strip(self.detail_content, [
                ("LLAMADAS EVALUADAS", filtered.get("total", 0), Theme.colors.CARD_BLUE),
                ("PROMEDIO", f'{filtered.get("average", 0):.0f}%', Theme.colors.CARD_GREEN),
                ("CON PENALIZACIÓN", filtered.get("penalized_calls", 0), Theme.colors.CARD_ORANGE),
                ("PENALIZACIONES", filtered.get("penalty_count", 0), Theme.colors.CARD_PURPLE),
            ])
            self._render_quality_cards(filtered)

    def _build_filter_bar(self, tab):
        # Tarjeta alta y estable: evita que las etiquetas queden ocultas bajo las pestañas.
        card = ctk.CTkFrame(self.detail_content, fg_color=Theme.colors.SURFACE, corner_radius=14,
                            border_width=1, border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x", padx=5, pady=(0, 12))
        card.grid_columnconfigure((0, 1, 2, 3, 4), weight=1, uniform="filters")

        ctk.CTkLabel(card, text="FILTROS DE CONSULTA", text_color=Theme.colors.PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=11, weight="bold"),
                     anchor="w").grid(row=0, column=0, columnspan=3, padx=16, pady=(12, 5), sticky="w")
        ctk.CTkButton(card, text="Limpiar filtros", width=105, height=27, corner_radius=8,
                      fg_color=Theme.colors.CARD_PURPLE, text_color=Theme.colors.PRIMARY,
                      hover_color=Theme.colors.BORDER_LIGHT,
                      command=lambda t=tab: self._clear_filters(t)).grid(
                          row=0, column=4, padx=16, pady=(9, 4), sticky="e")

        records = (getattr(self, "operational_data", {}) or {}).get("records", [])
        gestores = ["TODOS"] + sorted({r.get("gestor", "SIN IDENTIFICAR") for r in records})
        supervisores = ["TODOS"] + sorted({r.get("supervisor", "SIN IDENTIFICAR") for r in records})
        specs = [
            ("Gestor", gestores, self.filter_gestor, "filter_gestor"),
            ("Supervisor", supervisores, self.filter_supervisor, "filter_supervisor"),
            ("Resultado", ["TODOS", "APROBADO", "APROBADO SIN ETIQUETAR", "MENOR A 80%", "MAL DICTAMINADO", "NEGOCIACIÓN INVÁLIDA", "FALLA GRAVE", "REVISAR"], self.filter_resultado, "filter_resultado"),
            ("Auditoría", ["TODOS", "NEGOCIACIÓN INVÁLIDA", "CONVENIO INVÁLIDO", "MAL DICTAMINADO", "CONFRONTACIÓN", "FALLA GRAVE", "REFERENCIA DISTINTA"], self.filter_auditoria, "filter_auditoria"),
            ("Cumplimiento", ["TODOS", "80-100%", "60-79%", "0-59%"], self.filter_score, "filter_score"),
        ]
        for col, (label, values, current, attr) in enumerate(specs):
            box = ctk.CTkFrame(card, fg_color="transparent")
            box.grid(row=1, column=col, padx=(16 if col == 0 else 6, 16 if col == 4 else 6),
                     pady=(2, 13), sticky="ew")
            ctk.CTkLabel(box, text=label.upper(), text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold"),
                         anchor="w").pack(fill="x", pady=(0, 4))
            combo = ctk.CTkComboBox(box, values=values, height=30, state="readonly",
                                    command=lambda value, a=attr, t=tab: self._on_filter_change(a, value, t),
                                    font=ctk.CTkFont(family="Segoe UI", size=9))
            combo.set(current if current in values else "TODOS")
            combo.pack(fill="x")

    def _on_filter_change(self, attr, value, tab):
        setattr(self, attr, value)
        self.detail_selection = None
        self._render_detail_tab(tab)

    def _clear_filters(self, tab):
        self.filter_gestor = "TODOS"
        self.filter_supervisor = "TODOS"
        self.filter_resultado = "TODOS"
        self.filter_auditoria = "TODOS"
        self.filter_score = "TODOS"
        self.detail_selection = None
        self._render_detail_tab(tab)

    def _filtered_operational_data(self):
        records = (getattr(self, "operational_data", {}) or {}).get("records", [])
        out = []
        for r in records:
            if self.filter_gestor != "TODOS" and r.get("gestor") != self.filter_gestor:
                continue
            if self.filter_supervisor != "TODOS" and r.get("supervisor") != self.filter_supervisor:
                continue
            # RESULTADO operativo v3.6.1: categorías separadas de auditoría.
            result_label = str(r.get("operational_class") or "REVISAR")
            if self.filter_resultado != "TODOS" and result_label != self.filter_resultado:
                continue
            if self.filter_auditoria != "TODOS":
                audit_flags = set(r.get("audit_flags") or [])
                if self.filter_auditoria not in audit_flags:
                    continue
            score = r.get("score", 0)
            if self.filter_score == "80-100%" and score < 80: continue
            if self.filter_score == "60-79%" and not (60 <= score < 80): continue
            if self.filter_score == "0-59%" and score >= 60: continue
            out.append(r)
        return self._summarize_records(out)

    @staticmethod
    def _aggregate_records(records, key):
        groups = {}
        for r in records:
            name = r.get(key) or "SIN IDENTIFICAR"
            row = groups.setdefault(name, {"name": name, "calls": 0, "score_sum": 0.0,
                                           "cumple": 0, "parcial": 0, "no_cumple": 0})
            row["calls"] += 1
            row["score_sum"] += float(r.get("score", 0) or 0)
            row[r.get("status", "no_cumple")] += 1
        rows = []
        for row in groups.values():
            row["avg"] = row["score_sum"] / row["calls"] if row["calls"] else 0.0
            rows.append(row)
        rows.sort(key=lambda r: (-r["calls"], r["name"]))
        return rows

    @staticmethod
    def _summarize_records(records):
        block_points = Counter(); block_max = Counter(); penalty_reasons = Counter()
        score_sum = 0.0; penalized_calls = 0; penalty_count = 0; statuses = Counter()
        for r in records:
            score_sum += float(r.get("score", 0) or 0); statuses[r.get("status", "no_cumple")] += 1
            penalties = r.get("penalties", []) or []
            if penalties: penalized_calls += 1
            for pen in penalties:
                reason = str(pen.get("reason") or pen.get("category") or "Penalización").strip()
                penalty_reasons[reason] += 1; penalty_count += 1
            for name, values in (r.get("blocks", {}) or {}).items():
                if not isinstance(values, dict): continue
                try:
                    block_points[str(name)] += float(values.get("points", 0) or 0)
                    block_max[str(name)] += float(values.get("max_points", 0) or 0)
                except (TypeError, ValueError): pass
        labels = {"identificacion": "Identificación", "proteccion_datos": "Protección de datos",
                  "negociacion": "Negociación", "cierre": "Cierre"}
        blocks=[]
        for name, maximum in block_max.items():
            pct=(block_points[name]/maximum*100.0) if maximum else 0.0
            blocks.append({"name": labels.get(name, str(name).replace("_", " ").title()), "key": name, "pct": pct})
        blocks.sort(key=lambda b:b["name"])
        total=len(records)
        return {"records": records, "total": total, "average": score_sum/total if total else 0.0,
                "quality_blocks": blocks, "penalties": [{"reason":r,"count":c} for r,c in penalty_reasons.most_common(15)],
                "penalized_calls": penalized_calls, "penalty_count": penalty_count,
                "no_cumple": statuses.get("no_cumple",0)}

    def _build_entity_selector(self, label, rows, filtered, key):
        card = ctk.CTkFrame(self.detail_content, fg_color=Theme.colors.SURFACE, corner_radius=14,
                            border_width=1, border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x", padx=5, pady=(0, 10))
        top = ctk.CTkFrame(card, fg_color="transparent"); top.pack(fill="x", padx=15, pady=(11, 7))
        ctk.CTkLabel(top, text=f"DETALLE DE {label}", text_color=Theme.colors.TEXT_PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold")).pack(side="left")
        values=["Seleccione..."]+[r["name"] for r in rows]
        combo=ctk.CTkComboBox(top, values=values, width=260, height=29, state="readonly",
                              command=lambda value, f=filtered, k=key: self._show_entity_detail(value, f, k),
                              font=ctk.CTkFont(family="Segoe UI", size=9))
        combo.set(self.detail_selection if self.detail_selection in values else "Seleccione...")
        combo.pack(side="right")
        if self.detail_selection and self.detail_selection in values:
            self._entity_detail_body(card, self.detail_selection, filtered, key)
        else:
            ctk.CTkLabel(card, text="Seleccione un elemento para ver promedio, distribución, bloques y penalizaciones.",
                         text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9)).pack(fill="x", padx=15, pady=(0,11))

    def _show_entity_detail(self, value, filtered, key):
        if value == "Seleccione...": return
        self.detail_selection=value
        self._render_detail_tab(self.tab_selector.get())

    def _entity_detail_body(self, card, name, filtered, key):
        records = [r for r in filtered.get("records", []) if r.get(key) == name]
        summary = self._summarize_records(records)
        dist = Counter(r.get("status") for r in records)

        line = ctk.CTkFrame(card, fg_color=Theme.colors.CARD_PURPLE, corner_radius=10)
        line.pack(fill="x", padx=15, pady=(0, 10))
        txt = (f"{len(records)} llamadas   |   Promedio {summary.get('average', 0):.0f}%   |   "
               f"Cumple {dist.get('cumple', 0)}   |   Parcial {dist.get('parcial', 0)}   |   "
               f"No cumple {dist.get('no_cumple', 0)}")
        ctk.CTkLabel(line, text=txt, text_color=Theme.colors.TEXT_PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold")).pack(padx=12, pady=9)

        body = ctk.CTkFrame(card, fg_color="transparent")
        body.pack(fill="x", padx=15, pady=(0, 12))
        body.grid_columnconfigure(0, weight=3)
        body.grid_columnconfigure(1, weight=2)

        block_box = ctk.CTkFrame(body, fg_color=Theme.colors.BACKGROUND, corner_radius=10)
        block_box.grid(row=0, column=0, padx=(0, 6), sticky="nsew")
        ctk.CTkLabel(block_box, text="BLOQUES DE CALIDAD", text_color=Theme.colors.PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"), anchor="w").pack(
                         fill="x", padx=12, pady=(10, 5))
        blocks = summary.get("quality_blocks", [])
        if blocks:
            for block in blocks:
                row = ctk.CTkFrame(block_box, fg_color="transparent")
                row.pack(fill="x", padx=12, pady=4)
                ctk.CTkLabel(row, text=block["name"], width=125, anchor="w",
                             text_color=Theme.colors.TEXT_SECONDARY,
                             font=ctk.CTkFont(family="Segoe UI", size=9)).pack(side="left")
                bar = ctk.CTkProgressBar(row, height=9, corner_radius=5,
                                         progress_color=Theme.colors.PRIMARY)
                bar.pack(side="left", fill="x", expand=True, padx=8)
                bar.set(max(0.0, min(1.0, block["pct"] / 100.0)))
                ctk.CTkLabel(row, text=f'{block["pct"]:.0f}%', width=38, anchor="e",
                             text_color=Theme.colors.TEXT_PRIMARY,
                             font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold")).pack(side="right")
        else:
            ctk.CTkLabel(block_box, text="Sin información por bloque.",
                         text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9)).pack(padx=12, pady=(4, 12), anchor="w")

        penalty_box = ctk.CTkFrame(body, fg_color=Theme.colors.BACKGROUND, corner_radius=10)
        penalty_box.grid(row=0, column=1, padx=(6, 0), sticky="nsew")
        ctk.CTkLabel(penalty_box, text="TOP PENALIZACIONES", text_color=Theme.colors.PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI", size=10, weight="bold"), anchor="w").pack(
                         fill="x", padx=12, pady=(10, 5))
        penalties = summary.get("penalties", [])[:4]
        if penalties:
            for i, pen in enumerate(penalties, 1):
                row = ctk.CTkFrame(penalty_box, fg_color="transparent")
                row.pack(fill="x", padx=12, pady=3)
                ctk.CTkLabel(row, text=str(i), width=20, height=20, corner_radius=10,
                             fg_color=Theme.colors.CARD_PURPLE, text_color=Theme.colors.PRIMARY,
                             font=ctk.CTkFont(family="Segoe UI", size=8, weight="bold")).pack(side="left", padx=(0, 6))
                ctk.CTkLabel(row, text=pen["reason"], anchor="w", justify="left", wraplength=320,
                             text_color=Theme.colors.TEXT_SECONDARY,
                             font=ctk.CTkFont(family="Segoe UI", size=8)).pack(side="left", fill="x", expand=True)
                ctk.CTkLabel(row, text=str(pen["count"]), width=28, anchor="e",
                             text_color=Theme.colors.TEXT_PRIMARY,
                             font=ctk.CTkFont(family="Segoe UI", size=9, weight="bold")).pack(side="right")
        else:
            ctk.CTkLabel(penalty_box, text="Sin penalizaciones detectadas.",
                         text_color=Theme.colors.TEXT_SECONDARY,
                         font=ctk.CTkFont(family="Segoe UI", size=9)).pack(padx=12, pady=(4, 12), anchor="w")

    def _render_quality_cards(self, data):
        blocks=data.get("quality_blocks",[]); penalties=data.get("penalties",[])
        block_card=ctk.CTkFrame(self.detail_content,fg_color=Theme.colors.SURFACE,corner_radius=14,
                                border_width=1,border_color=Theme.colors.BORDER_LIGHT)
        block_card.pack(fill="x",padx=5,pady=(0,12))
        ctk.CTkLabel(block_card,text="Cumplimiento por bloque · seleccione para ver causas",text_color=Theme.colors.TEXT_PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI",size=15,weight="bold"),anchor="w").pack(fill="x",padx=16,pady=(13,7))
        for block in blocks:
            row=ctk.CTkFrame(block_card,fg_color="transparent"); row.pack(fill="x",padx=16,pady=4)
            ctk.CTkButton(row,text=block["name"],width=190,height=27,anchor="w",fg_color="transparent",
                          hover_color=Theme.colors.CARD_PURPLE,text_color=Theme.colors.TEXT_PRIMARY,
                          font=ctk.CTkFont(family="Segoe UI",size=10,weight="bold"),
                          command=lambda b=block: self._show_block_causes(b)).pack(side="left")
            bar=ctk.CTkProgressBar(row,height=10,corner_radius=5,progress_color=Theme.colors.PRIMARY); bar.pack(side="left",fill="x",expand=True,padx=10)
            bar.set(max(0.0,min(1.0,block["pct"]/100.0)))
            ctk.CTkLabel(row,text=f'{block["pct"]:.0f}%',width=50,text_color=Theme.colors.TEXT_PRIMARY,
                         font=ctk.CTkFont(family="Segoe UI",size=10,weight="bold")).pack(side="right")
        self.block_detail=ctk.CTkLabel(block_card,text="Haga clic en un bloque para ver las penalizaciones relacionadas.",
                                       text_color=Theme.colors.TEXT_SECONDARY,
                                       font=ctk.CTkFont(family="Segoe UI",size=9),anchor="w",wraplength=900)
        self.block_detail.pack(fill="x",padx=16,pady=(7,12))
        pen_card=ctk.CTkFrame(self.detail_content,fg_color=Theme.colors.SURFACE,corner_radius=14,
                              border_width=1,border_color=Theme.colors.BORDER_LIGHT); pen_card.pack(fill="x",padx=5,pady=(0,12))
        ctk.CTkLabel(pen_card,text="Penalizaciones más frecuentes",text_color=Theme.colors.TEXT_PRIMARY,
                     font=ctk.CTkFont(family="Segoe UI",size=15,weight="bold"),anchor="w").pack(fill="x",padx=16,pady=(13,7))
        if not penalties:
            ctk.CTkLabel(pen_card,text="Sin penalizaciones registradas",text_color=Theme.colors.TEXT_SECONDARY).pack(pady=24)
        for item in penalties[:10]:
            row=ctk.CTkFrame(pen_card,fg_color="transparent"); row.pack(fill="x",padx=16,pady=4)
            ctk.CTkLabel(row,text=item["reason"],anchor="w",text_color=Theme.colors.TEXT_PRIMARY,
                         font=ctk.CTkFont(family="Segoe UI",size=10)).pack(side="left",fill="x",expand=True)
            ctk.CTkLabel(row,text=str(item["count"]),width=55,text_color=Theme.colors.PRIMARY,
                         font=ctk.CTkFont(family="Segoe UI",size=10,weight="bold")).pack(side="right")

    def _show_block_causes(self, block):
        data=self._filtered_operational_data(); key=block.get("key",""); reasons=Counter()
        for r in data.get("records",[]):
            for pen in r.get("penalties",[]) or []:
                cat=str(pen.get("category") or "").strip().lower()
                normalized=cat.replace("ó","o").replace("í","i").replace(" ","_")
                if key and (key in normalized or normalized in key):
                    reasons[str(pen.get("reason") or pen.get("category") or "Penalización")]+=1
        text=" · ".join(f"{reason} ({count})" for reason,count in reasons.most_common(4))
        self.block_detail.configure(text=(f"{block['name']}: " + (text if text else "No hay penalizaciones categorizadas para este bloque.")))

    # =========================================================
    # CONTENIDO
    # =========================================================

    def _build_content(self):
        self.content = ctk.CTkScrollableFrame(
            self,
            fg_color="transparent",
        )

        self.content.grid(
            row=2,
            column=0,
            padx=20,
            pady=(0, 12),
            sticky="nsew",
        )

        self.content.grid_columnconfigure(0, weight=1)

        self._build_kpi_section()
        self._build_primary_charts()
        self._build_secondary_charts()

        self.status_label = ctk.CTkLabel(
            self.content,
            text="",
            text_color=Theme.colors.TEXT_SECONDARY,
            font=ctk.CTkFont(
                family="Segoe UI",
                size=11,
            ),
            anchor="w",
        )

        self.status_label.grid(
            row=3,
            column=0,
            padx=4,
            pady=(12, 4),
            sticky="ew",
        )

    # =========================================================
    # KPIs
    # =========================================================

    def _build_kpi_section(self):
        self.kpi_frame = ctk.CTkFrame(
            self.content,
            fg_color="transparent",
        )

        self.kpi_frame.grid(
            row=0,
            column=0,
            sticky="ew",
        )

        for column in range(4):
            self.kpi_frame.grid_columnconfigure(
                column,
                weight=1,
                uniform="kpi",
            )

        self.total_calls_value = self._create_kpi_card(
            0,
            "LLAMADAS ANALIZADAS",
            Theme.colors.CARD_BLUE,
            "0",
        )

        self.compliance_value = self._create_kpi_card(
            1,
            "CUMPLIMIENTO PROMEDIO",
            Theme.colors.CARD_GREEN,
            "0%",
        )

        self.main_reason_value = self._create_kpi_card(
            2,
            "MOTIVO PRINCIPAL",
            Theme.colors.CARD_PURPLE,
            "Sin datos",
        )

        self.main_sentiment_value = self._create_kpi_card(
            3,
            "SENTIMIENTO PRINCIPAL",
            Theme.colors.CARD_ORANGE,
            "Sin datos",
        )

    def _create_kpi_card(
        self,
        column,
        title,
        background,
        value,
    ):
        card = ctk.CTkFrame(
            self.kpi_frame,
            fg_color=background,
            corner_radius=15,
            border_width=1,
            border_color=Theme.colors.BORDER_LIGHT,
        )

        card.grid(
            row=0,
            column=column,
            padx=5,
            sticky="nsew",
        )

        ctk.CTkLabel(
            card,
            text=title,
            text_color=Theme.colors.TEXT_SECONDARY,
            font=ctk.CTkFont(
                family="Segoe UI",
                size=10,
                weight="bold",
            ),
            anchor="w",
        ).pack(
            anchor="w",
            padx=18,
            pady=(16, 5),
        )

        value_label = ctk.CTkLabel(
            card,
            text=value,
            text_color=Theme.colors.TEXT_PRIMARY,
            font=ctk.CTkFont(
                family="Segoe UI",
                size=23,
                weight="bold",
            ),
            anchor="w",
        )

        value_label.pack(
            anchor="w",
            padx=18,
            pady=(0, 17),
        )

        return value_label

    # =========================================================
    # GRÁFICAS PRINCIPALES
    # =========================================================

    def _build_primary_charts(self):
        self.primary_charts = ctk.CTkFrame(
            self.content,
            fg_color="transparent",
        )

        self.primary_charts.grid(
            row=1,
            column=0,
            pady=(12, 0),
            sticky="ew",
        )

        self.primary_charts.grid_columnconfigure(
            0,
            weight=1,
            uniform="primary",
        )

        self.primary_charts.grid_columnconfigure(
            1,
            weight=1,
            uniform="primary",
        )

        self.sentiment_chart = ChartCard(
            self.primary_charts,
            title="Sentimiento de las llamadas",
            data={
                "positivo": 0,
                "neutral": 0,
                "negativo": 0,
            },
            height=220,
            bar_color=Theme.colors.PRIMARY,
        )

        self.sentiment_chart.grid(
            row=0,
            column=0,
            padx=(5, 6),
            sticky="nsew",
        )

        self.compliance_chart = ChartCard(
            self.primary_charts,
            title="Estado de cumplimiento",
            data={
                "cumple": 0,
                "parcial": 0,
                "no_cumple": 0,
            },
            height=220,
            bar_color=Theme.colors.SUCCESS,
        )

        self.compliance_chart.grid(
            row=0,
            column=1,
            padx=(6, 5),
            sticky="nsew",
        )

    # =========================================================
    # GRÁFICAS SECUNDARIAS
    # =========================================================

    def _build_secondary_charts(self):
        self.secondary_charts = ctk.CTkFrame(
            self.content,
            fg_color="transparent",
        )

        self.secondary_charts.grid(
            row=2,
            column=0,
            pady=(12, 0),
            sticky="ew",
        )

        self.secondary_charts.grid_columnconfigure(
            0,
            weight=1,
            uniform="secondary",
        )

        self.secondary_charts.grid_columnconfigure(
            1,
            weight=1,
            uniform="secondary",
        )

        # -----------------------------------------------------
        # MOTIVOS
        # -----------------------------------------------------

        self.reasons_chart = ChartCard(
            self.secondary_charts,
            title="Motivos de llamada",
            data={},
            height=245,
            bar_color=Theme.colors.BLUE,
        )

        self.reasons_chart.grid(
            row=0,
            column=0,
            padx=(5, 6),
            sticky="nsew",
        )

        # -----------------------------------------------------
        # PALABRAS CLAVE
        # -----------------------------------------------------

        self.keywords_chart = ChartCard(
            self.secondary_charts,
            title="Palabras clave principales",
            data={},
            height=245,
            bar_color=Theme.colors.WARNING,
        )

        self.keywords_chart.grid(
            row=0,
            column=1,
            padx=(6, 5),
            sticky="nsew",
        )

    # =========================================================
    # CARGAR DASHBOARD
    # =========================================================

    def load_dashboard(self):
        project_path = self._get_active_project_path()

        if project_path is None:
            self.project_label.configure(
                text="Proyecto: Sin proyecto abierto"
            )

            self._clear_dashboard()

            self.status_label.configure(
                text="Abra un proyecto para consultar el Dashboard."
            )

            return

        self.project_label.configure(
            text=f"Proyecto: {self._get_active_project_name()}"
        )

        self.operational_data = self._load_operational_dashboard(project_path)
        if not self.operational_data.get("records"):
            self.operational_data = self._load_standalone_analytics(project_path)

        try:
            self.refresh_button.configure(
                state="disabled",
                text="Actualizando...",
            )

            self.update_idletasks()

            data = DashboardService.get_dashboard_data(
                project_path
            )

            # El procesamiento por lotes conserva Analytics en el checkpoint.
            # Si el servicio tradicional no encuentra reportes guardados,
            # reutilizamos esos resultados SIN reprocesar los audios.
            if int(data.get("total_calls", 0) or 0) <= 0:
                checkpoint_data = self._load_batch_checkpoint_dashboard(project_path)
                if int(checkpoint_data.get("total_calls", 0) or 0) > 0:
                    data = checkpoint_data

            self.dashboard_data = data

            self._display_dashboard(
                data
            )

            current_tab = self.tab_selector.get()
            if current_tab != "VISIÓN GENERAL":
                self._render_detail_tab(current_tab)

            total = int(
                data.get(
                    "total_calls",
                    0,
                )
                or 0
            )

            if total <= 0:
                self.status_label.configure(
                    text=(
                        "No existen análisis guardados. "
                        "Analice llamadas desde Analytics de Llamadas."
                    )
                )

            else:
                suffix = "" if total == 1 else "s"

                self.status_label.configure(
                    text=(
                        f"Dashboard actualizado · "
                        f"{total} llamada{suffix} analizada{suffix}."
                    )
                )

        except Exception as error:
            self._clear_dashboard()

            self.status_label.configure(
                text="No fue posible cargar los indicadores."
            )

            messagebox.showerror(
                "Dashboard",
                "No fue posible cargar el Dashboard.\n\n"
                f"Detalle: {error}",
            )

        finally:
            self.refresh_button.configure(
                state="normal",
                text="Actualizar datos",
            )

    # =========================================================
    # MOSTRAR DATOS
    # =========================================================

    def _display_dashboard(
        self,
        data,
    ):
        analyses = data.get(
            "analyses",
            [],
        ) or []

        # =====================================================
        # KPIs
        # =====================================================

        total_calls = int(
            data.get(
                "total_calls",
                0,
            )
            or 0
        )

        average_compliance = float(
            data.get(
                "average_compliance",
                0.0,
            )
            or 0.0
        )

        self.total_calls_value.configure(
            text=str(total_calls)
        )

        self.compliance_value.configure(
            text=f"{average_compliance:.0f}%"
        )

        main_reason = data.get("main_reason")
        if not main_reason:
            main_reason = (
                DashboardService.get_main_reason(analyses)
                if analyses else "Sin datos"
            )

        main_sentiment = data.get("main_sentiment")
        if not main_sentiment:
            main_sentiment = (
                DashboardService.get_main_sentiment(analyses)
                if analyses else "Sin datos"
            )

        self.main_reason_value.configure(
            text=str(main_reason)
        )

        self.main_sentiment_value.configure(
            text=str(main_sentiment).capitalize()
        )

        # =====================================================
        # SENTIMIENTOS
        # =====================================================

        sentiments = data.get(
            "sentiments",
            {},
        ) or {}

        self.sentiment_chart.set_data(
            {
                "positivo": self._safe_int(
                    sentiments.get("positivo", 0)
                ),
                "neutral": self._safe_int(
                    sentiments.get("neutral", 0)
                ),
                "negativo": self._safe_int(
                    sentiments.get("negativo", 0)
                ),
            }
        )

        # =====================================================
        # CUMPLIMIENTO
        # =====================================================

        compliance = data.get(
            "compliance",
            {},
        ) or {}

        self.compliance_chart.set_data(
            {
                "cumple": self._safe_int(
                    compliance.get("cumple", 0)
                ),
                "parcial": self._safe_int(
                    compliance.get("parcial", 0)
                ),
                "no_cumple": self._safe_int(
                    compliance.get("no_cumple", 0)
                ),
            }
        )

        # =====================================================
        # MOTIVOS
        # =====================================================

        reasons = data.get(
            "call_reasons",
            {},
        ) or {}

        reason_data = {}

        if isinstance(
            reasons,
            dict,
        ):
            for reason, count in reasons.items():
                reason_data[
                    str(reason)
                ] = self._safe_int(
                    count
                )

        self.reasons_chart.set_data(
            reason_data
        )

        # =====================================================
        # PALABRAS CLAVE
        # =====================================================

        keywords = data.get(
            "top_keywords",
            [],
        ) or []

        keyword_data = {}

        if isinstance(
            keywords,
            list,
        ):
            for item in keywords:

                if not isinstance(
                    item,
                    dict,
                ):
                    continue

                keyword = str(
                    item.get(
                        "keyword",
                        "",
                    )
                    or ""
                ).strip()

                if not keyword:
                    continue

                keyword_data[
                    keyword
                ] = self._safe_int(
                    item.get(
                        "count",
                        0,
                    )
                )

        self.keywords_chart.set_data(
            keyword_data
        )
        try:
            mentions = sum(keyword_data.values())
            self.keywords_chart.total_label.configure(text=f"{mentions:,} menciones")
        except Exception:
            pass

    # =========================================================
    # FALLBACK: ANALYTICS DEL CHECKPOINT DE LOTE
    # =========================================================

    @staticmethod
    def _load_batch_checkpoint_dashboard(project_path: Path) -> dict:
        """Construye indicadores usando resultados ya guardados por BatchService."""
        checkpoint = (
            Path(project_path)
            / "05_Datos"
            / "batch_state"
            / "batch_checkpoint.json"
        )

        if not checkpoint.exists():
            return {}

        try:
            with checkpoint.open("r", encoding="utf-8") as handle:
                state = json.load(handle)
        except (OSError, ValueError, TypeError):
            return {}

        items = state.get("items", []) if isinstance(state, dict) else []
        if not isinstance(items, list):
            return {}

        analyses = []
        sentiments = Counter()
        compliance = Counter()
        reasons = Counter()
        keywords = Counter()
        scores = []

        for item in items:
            if not isinstance(item, dict):
                continue

            analytics = item.get("analytics")
            if not isinstance(analytics, dict):
                continue

            # Solo cuenta registros que realmente tienen un Analytics terminado.
            score = analytics.get("compliance_score")
            if score is None:
                score = item.get("compliance_score")

            try:
                score_value = float(score)
            except (TypeError, ValueError):
                score_value = None

            if score_value is None:
                continue

            analyses.append(analytics)
            scores.append(score_value)

            sentiment = str(analytics.get("sentiment", "") or "").strip().lower()
            if sentiment in {"positivo", "positiva", "positive"}:
                sentiments["positivo"] += 1
            elif sentiment in {"negativo", "negativa", "negative"}:
                sentiments["negativo"] += 1
            else:
                sentiments["neutral"] += 1

            status = str(analytics.get("compliance_status", "") or "").strip().lower()
            status = status.replace(" ", "_").replace("-", "_")
            if status in {"cumple", "completo", "aprobado", "aprobada"}:
                compliance["cumple"] += 1
            elif status in {"parcial", "cumple_parcial", "parcialmente"}:
                compliance["parcial"] += 1
            else:
                # Si el Analytics no trae un estado reconocido, usamos el score
                # únicamente para representarlo en las tres bandas del dashboard.
                if score_value >= 80:
                    compliance["cumple"] += 1
                elif score_value >= 60:
                    compliance["parcial"] += 1
                else:
                    compliance["no_cumple"] += 1

            reason = str(analytics.get("call_reason", "") or "").strip()
            if reason:
                reasons[reason] += 1

            raw_keywords = analytics.get("keywords", []) or []
            if isinstance(raw_keywords, (list, tuple)):
                for keyword in raw_keywords:
                    word = str(keyword or "").strip().lower()
                    if word:
                        keywords[word] += 1

        total = len(analyses)
        if total <= 0:
            return {}

        average = sum(scores) / len(scores) if scores else 0.0
        main_reason = reasons.most_common(1)[0][0] if reasons else "Sin datos"
        main_sentiment = sentiments.most_common(1)[0][0] if sentiments else "Sin datos"

        return {
            "source": "batch_checkpoint",
            "analyses": analyses,
            "total_calls": total,
            "average_compliance": average,
            "main_reason": main_reason,
            "main_sentiment": main_sentiment,
            "sentiments": {
                "positivo": sentiments.get("positivo", 0),
                "neutral": sentiments.get("neutral", 0),
                "negativo": sentiments.get("negativo", 0),
            },
            "compliance": {
                "cumple": compliance.get("cumple", 0),
                "parcial": compliance.get("parcial", 0),
                "no_cumple": compliance.get("no_cumple", 0),
            },
            "call_reasons": dict(reasons.most_common(8)),
            "top_keywords": [
                {"keyword": word, "count": count}
                for word, count in keywords.most_common(10)
            ],
        }

    # =========================================================
    # DATOS OPERATIVOS: GESTORES / SUPERVISORES / CALIDAD
    # =========================================================

    @staticmethod
    def _load_operational_dashboard(project_path: Path) -> dict:
        checkpoint = Path(project_path) / "05_Datos" / "batch_state" / "batch_checkpoint.json"
        if not checkpoint.exists():
            return {}
        try:
            with checkpoint.open("r", encoding="utf-8") as handle:
                state = json.load(handle)
        except (OSError, ValueError, TypeError):
            return {}
        items = state.get("items", []) if isinstance(state, dict) else []
        if not isinstance(items, list):
            return {}

        groups_g = {}
        groups_s = {}
        block_points = Counter()
        block_max = Counter()
        penalty_reasons = Counter()
        total = 0
        score_sum = 0.0
        penalized_calls = 0
        penalty_count = 0
        overall_status = Counter()
        records = []

        def add_group(container, name, score, status):
            key = name or "SIN IDENTIFICAR"
            row = container.setdefault(key, {"name": key, "calls": 0, "score_sum": 0.0,
                                             "cumple": 0, "parcial": 0, "no_cumple": 0})
            row["calls"] += 1
            row["score_sum"] += score
            row[status] += 1

        for item in items:
            if not isinstance(item, dict):
                continue
            analytics = item.get("analytics")
            if not isinstance(analytics, dict):
                continue
            score = analytics.get("compliance_score", item.get("compliance_score"))
            try:
                score = float(score)
            except (TypeError, ValueError):
                continue

            status_raw = str(analytics.get("compliance_status", "") or "").strip().lower().replace(" ", "_").replace("-", "_")
            if status_raw in {"cumple", "completo", "aprobado", "aprobada"}:
                status = "cumple"
            elif status_raw in {"parcial", "cumple_parcial", "parcialmente"}:
                status = "parcial"
            elif status_raw in {"no_cumple", "incumple", "rechazado", "rechazada"}:
                status = "no_cumple"
            else:
                status = "cumple" if score >= 80 else ("parcial" if score >= 60 else "no_cumple")

            relation = item.get("relation") if isinstance(item.get("relation"), dict) else {}
            metadata = analytics.get("metadata") if isinstance(analytics.get("metadata"), dict) else {}
            context = metadata.get("operational_context") if isinstance(metadata.get("operational_context"), dict) else {}
            gestor = str(relation.get("gestor") or context.get("gestor") or "SIN IDENTIFICAR").strip().upper()
            supervisor = str(relation.get("supervisor") or context.get("supervisor") or "SIN IDENTIFICAR").strip().upper()

            add_group(groups_g, gestor, score, status)
            add_group(groups_s, supervisor, score, status)
            overall_status[status] += 1
            total += 1
            score_sum += score

            blocks = metadata.get("quality_blocks") if isinstance(metadata.get("quality_blocks"), dict) else {}
            for name, values in blocks.items():
                if not isinstance(values, dict):
                    continue
                try:
                    pts = float(values.get("points", 0) or 0)
                    mx = float(values.get("max_points", 0) or 0)
                except (TypeError, ValueError):
                    continue
                block_points[str(name)] += pts
                block_max[str(name)] += mx

            penalties = metadata.get("penalties") if isinstance(metadata.get("penalties"), list) else []
            relation_status = str(relation.get("estatus_grabacion") or "SIN ESTATUS").strip().upper()
            status_validation = metadata.get("status_validation") if isinstance(metadata.get("status_validation"), dict) else {}
            bad_dictation = status_validation.get("valid") is False
            bad_reasons = status_validation.get("reasons") if isinstance(status_validation.get("reasons"), list) else []
            detected_status = str(status_validation.get("detected_status") or "").strip().upper()
            result_data = item.get("result") if isinstance(item.get("result"), dict) else {}
            transcription_text = str(result_data.get("texto") or analytics.get("transcription") or "")
            audio_path = str(item.get("path") or result_data.get("ruta_audio") or "")
            negotiation = metadata.get("negotiation_validation") if isinstance(metadata.get("negotiation_validation"), dict) else {}
            treatment = metadata.get("treatment_validation") if isinstance(metadata.get("treatment_validation"), dict) else {}
            failures = metadata.get("critical_failures") if isinstance(metadata.get("critical_failures"), list) else []
            negotiation_invalid = negotiation.get("valid") is False and negotiation.get("is_negotiation") is True
            convenio_invalid = relation_status == "CONVENIO" and negotiation_invalid
            confrontation = "CONFRONTACION DEL GESTOR" in str(treatment.get("confrontation") or "").upper()
            grave = confrontation or any(isinstance(f, dict) and f.get("invalidates") is True for f in failures)
            audit_flags = []
            if negotiation_invalid: audit_flags.append("NEGOCIACIÓN INVÁLIDA")
            if convenio_invalid: audit_flags.append("CONVENIO INVÁLIDO")
            if bad_dictation: audit_flags.append("MAL DICTAMINADO")
            if confrontation: audit_flags.append("CONFRONTACIÓN")
            if grave: audit_flags.append("FALLA GRAVE")
            # Auditoría transversal: no depende de clasificación operativa.
            card_info = metadata.get("card_validation") if isinstance(metadata.get("card_validation"), dict) else {}
            from app.services.payment_reference_validation import find_reference
            expected_card = str(relation.get("numero_de_tarjeta") or "")
            ref_check = find_reference(transcription_text, expected_card)
            if ref_check["incorrect"] or "REFERENCIA DISTINTA" in str(card_info.get("result") or ""):
                audit_flags.append("REFERENCIA DISTINTA")
            checkpoint_classification = item.get("classification") if isinstance(item.get("classification"), dict) else {}
            physical_class = str(checkpoint_classification.get("clasificacion") or "").strip().upper()
            operational_label_map = {
                "APROBADO": "APROBADO",
                "APROBADOS": "APROBADO",
                "APROBADO_SIN_ETIQUETAR": "APROBADO SIN ETIQUETAR",
                "APROBADO_CON_EXCEPCION": "APROBADO CON EXCEPCIÓN",
                "NO_APROBADO_MENOR_80": "MENOR A 80%",
                "MAL_DICTAMINADO": "MAL DICTAMINADO",
                "NEGOCIACION_INVALIDA": "NEGOCIACIÓN INVÁLIDA",
                "FALLA_GRAVE": "FALLA GRAVE",
                "REVISAR": "REVISAR",
            }
            if physical_class:
                operational_class = operational_label_map.get(physical_class, "REVISAR")
            elif grave:
                operational_class = "FALLA GRAVE"
            elif bad_dictation:
                operational_class = "MAL DICTAMINADO"
            elif negotiation_invalid:
                operational_class = "NEGOCIACIÓN INVÁLIDA"
            elif score < 80:
                operational_class = "MENOR A 80%"
            else:
                operational_class = "REVISAR"
            final_status = "mal_dictaminado" if bad_dictation else status
            records.append({"gestor": gestor, "supervisor": supervisor, "score": score,
                            "status": status, "final_status": final_status, "operational_class": operational_class, "blocks": blocks, "penalties": penalties,
                            "audit_flags": audit_flags, "negotiation_invalid": negotiation_invalid,
                            "convenio_invalid": convenio_invalid, "confrontation_flag": confrontation, "grave_flag": grave,
                            "file_name": str(item.get("file_name") or relation.get("archivo_original") or ""),
                            "audio_path": audio_path, "transcription_text": transcription_text,
                            "excel_status": relation_status, "detected_status": detected_status,
                            "bad_dictation": bad_dictation, "bad_reasons": bad_reasons})
            if penalties:
                penalized_calls += 1
            for pen in penalties:
                if not isinstance(pen, dict):
                    continue
                reason = str(pen.get("reason") or pen.get("category") or "Penalización").strip()
                penalty_reasons[reason] += 1
                penalty_count += 1

        def finalize(container):
            rows = []
            for row in container.values():
                calls = row["calls"]
                row["avg"] = row["score_sum"] / calls if calls else 0.0
                rows.append(row)
            rows.sort(key=lambda r: (-r["calls"], r["name"]))
            return rows

        blocks = []
        labels = {"identificacion": "Identificación", "proteccion_datos": "Protección de datos",
                  "negociacion": "Negociación", "cierre": "Cierre"}
        for name, maximum in block_max.items():
            pct = (block_points[name] / maximum * 100.0) if maximum else 0.0
            blocks.append({"name": labels.get(name, str(name).replace("_", " ").title()), "pct": pct})
        blocks.sort(key=lambda b: b["name"])

        return {
            "total": total,
            "average": score_sum / total if total else 0.0,
            "gestores": finalize(groups_g),
            "supervisores": finalize(groups_s),
            "quality_blocks": blocks,
            "penalties": [{"reason": r, "count": c} for r, c in penalty_reasons.most_common(10)],
            "penalized_calls": penalized_calls,
            "penalty_count": penalty_count,
            "no_cumple": overall_status.get("no_cumple", 0),
            "records": records,
        }

    @staticmethod
    def _load_standalone_analytics(project_path: Path) -> dict:
        """Carga análisis individuales para reportes, sin checkpoint ni Excel."""
        base = Path(project_path)
        directory = base / "04_Reportes" / "Analytics"
        if not directory.is_dir():
            return {}
        records = []
        for path in directory.glob("*_analytics.json"):
            try:
                with path.open("r", encoding="utf-8") as stream:
                    data = json.load(stream)
            except (OSError, ValueError):
                continue
            if not isinstance(data, dict):
                continue
            metadata = data.get("metadata") or {}
            shared = metadata.get("shared_classification") or {}
            score = float(data.get("compliance_score") or 0)
            classification = shared.get("clasificacion") or ("APROBADO" if score >= 80 else "REVISAR")
            transcription = ""
            transcript_path = data.get("transcription_file")
            if transcript_path:
                try:
                    with Path(transcript_path).open("r", encoding="utf-8") as stream:
                        original = json.load(stream)
                    transcription = str(original.get("texto") or original.get("text") or "")
                except (OSError, ValueError, TypeError):
                    try:
                        transcription = Path(transcript_path).read_text(encoding="utf-8")
                    except (OSError, ValueError, TypeError):
                        pass
            if not transcription:
                transcription = str(data.get("transcription") or "")
            records.append({
                "gestor": str((metadata.get("operational_context") or {}).get("gestor") or "SIN IDENTIFICAR"),
                "supervisor": str((metadata.get("operational_context") or {}).get("supervisor") or "SIN IDENTIFICAR"),
                "score": score, "status": "cumple" if score >= 80 else "parcial" if score >= 60 else "no_cumple",
                "final_status": classification, "operational_class": classification,
                "blocks": metadata.get("quality_blocks") or {},
                "penalties": metadata.get("penalties") or [], "audit_flags": [],
                "file_name": data.get("audio_file") or path.stem,
                "audio_path": "", "transcription_text": transcription,
                "excel_status": "SIN EXCEL", "detected_status": (metadata.get("status_validation") or {}).get("detected_status") or "",
                "bad_dictation": False, "bad_reasons": [],
            })
        if not records:
            return {}
        return DashboardView._summarize_records(records)

    # =========================================================
    # AUDITORÍA: PALABRAS / ESTATUS / SALIDAS FINALES
    # =========================================================

    def _audit_records(self):
        return list((getattr(self, "operational_data", {}) or {}).get("records", []) or [])

    def _render_word_audit(self):
        self._section_title(self.detail_content, "BÚSQUEDA POR PALABRA O FRASE",
                            "Busca dentro de las transcripciones. Este filtro es independiente del estatus del Excel.")
        card = ctk.CTkFrame(self.detail_content, fg_color=Theme.colors.SURFACE, corner_radius=14,
                            border_width=1, border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x", padx=5, pady=(0, 10))
        self.word_query = ctk.StringVar(value=getattr(self, "_last_word_query", ""))
        entry = ctk.CTkEntry(card, textvariable=self.word_query, height=36,
                             placeholder_text="Ej. convenio, pago mínimo, reestructura...")
        entry.pack(side="left", fill="x", expand=True, padx=(14, 8), pady=14)
        ctk.CTkButton(card, text="Buscar", width=90, height=36, command=self._run_word_search).pack(side="left", padx=4, pady=14)
        ctk.CTkButton(card, text="Escuchar audios", width=118, height=36,
                      fg_color=Theme.colors.CARD_BLUE, text_color=Theme.colors.PRIMARY,
                      command=lambda: self._open_audio_player(getattr(self, "_word_results", []), "Palabras: " + str(getattr(self, "_last_word_query", "") or self.word_query.get()))).pack(side="left", padx=4, pady=14)
        ctk.CTkButton(card, text="Copiar audios", width=110, height=36,
                      fg_color=Theme.colors.CARD_PURPLE, text_color=Theme.colors.PRIMARY,
                      command=lambda: self._copy_audit_records(getattr(self, "_word_results", []), "PALABRAS")).pack(side="left", padx=4, pady=14)
        ctk.CTkButton(card, text="Generar reporte Excel", width=145, height=36,
                      fg_color=Theme.colors.PRIMARY,
                      command=lambda: self._export_audit_report(getattr(self, "_word_results", []), "PALABRAS")).pack(side="left", padx=(4,14), pady=14)
        self._word_results = self._word_matches(self.word_query.get())
        self._render_audit_results(self._word_results, "Resultados de búsqueda")

    def _run_word_search(self):
        self._last_word_query = self.word_query.get().strip()
        self._render_detail_tab("PALABRAS")

    def _word_matches(self, query):
        q = str(query or "").strip().casefold()
        if not q:
            return []
        rows=[]
        import unicodedata
        def normalized(value):
            value = unicodedata.normalize("NFKD", str(value or "").casefold())
            return "".join(c for c in value if not unicodedata.combining(c))
        q = normalized(q)
        for r in self._filtered_operational_data().get("records", []):
            text=str(r.get("transcription_text") or "")
            count=normalized(text).count(q)
            if count:
                x=dict(r); x["matches"]=count; rows.append(x)
        rows.sort(key=lambda r:(-r.get("matches",0), r.get("file_name","")))
        return rows

    def _render_excel_status_audit(self):
        self._section_title(self.detail_content, "FILTRO POR ESTATUS DEL EXCEL",
                            "Usa exactamente el estatus operativo de la relación de audios. Es independiente de la búsqueda por palabras.")
        # IMPORTANTE: la vista ESTATUS debe partir del mismo subconjunto que
        # producen los filtros superiores (Gestor, Supervisor, Resultado y
        # Cumplimiento). De esta forma, por ejemplo, NANCI JUAN + CONVENIO
        # devuelve únicamente los convenios de ese gestor y las acciones
        # Escuchar/Copiar/Excel reciben exactamente ese mismo subconjunto.
        records=list(self._filtered_operational_data().get("records", []) or [])
        statuses=["TODOS"]+sorted({str(r.get("excel_status") or "SIN ESTATUS") for r in records})
        selected=getattr(self,"_excel_status_selected","TODOS")
        if selected not in statuses: selected="TODOS"
        card=ctk.CTkFrame(self.detail_content, fg_color=Theme.colors.SURFACE, corner_radius=14,
                          border_width=1, border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x", padx=5, pady=(0,10))
        ctk.CTkLabel(card,text="Estatus:",text_color=Theme.colors.TEXT_PRIMARY,font=ctk.CTkFont(family="Segoe UI",size=11,weight="bold")).pack(side="left",padx=(14,7),pady=14)
        menu=ctk.CTkOptionMenu(card,values=statuses,width=260,height=36,command=self._on_excel_status_change)
        menu.set(selected); menu.pack(side="left",padx=4,pady=14)
        ctk.CTkButton(card,text="Generar reporte Excel",width=145,height=36,
                      fg_color=Theme.colors.PRIMARY,
                      command=lambda:self._export_audit_report(getattr(self,"_status_results",[]),"ESTATUS")).pack(side="right",padx=(4,14),pady=14)
        ctk.CTkButton(card,text="Copiar audios",width=110,height=36,
                      fg_color=Theme.colors.CARD_PURPLE,text_color=Theme.colors.PRIMARY,
                      command=lambda:self._copy_audit_records(getattr(self,"_status_results",[]),"ESTATUS")).pack(side="right",padx=4,pady=14)
        ctk.CTkButton(card,text="Escuchar audios",width=118,height=36,
                      fg_color=Theme.colors.CARD_BLUE,text_color=Theme.colors.PRIMARY,
                      command=lambda:self._open_audio_player(getattr(self,"_status_results",[]),"Estatus: " + str(getattr(self,"_excel_status_selected","TODOS")))).pack(side="right",padx=4,pady=14)
        self._status_results=records if selected=="TODOS" else [r for r in records if r.get("excel_status")==selected]
        self._render_audit_results(self._status_results,"Audios por estatus")

    def _on_excel_status_change(self,value):
        self._excel_status_selected=value
        self._render_detail_tab("ESTATUS")

    def _render_audit_results(self, rows, title):
        self._metric_strip(self.detail_content, [
            ("RESULTADOS",len(rows),Theme.colors.CARD_BLUE),
            ("MAL DICTAMINADOS",sum(1 for r in rows if r.get("bad_dictation")),Theme.colors.CARD_ORANGE),
            ("PROMEDIO",f'{(sum(float(r.get("score",0) or 0) for r in rows)/len(rows) if rows else 0):.0f}%',Theme.colors.CARD_GREEN),
        ])
        card=ctk.CTkFrame(self.detail_content,fg_color=Theme.colors.SURFACE,corner_radius=14,border_width=1,border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x",padx=5,pady=(0,12))
        ctk.CTkLabel(card,text=title,text_color=Theme.colors.TEXT_PRIMARY,font=ctk.CTkFont(family="Segoe UI",size=14,weight="bold"),anchor="w").pack(fill="x",padx=14,pady=(12,6))
        headers=[("AUDIO ORIGINAL",340),("GESTOR",155),("SUPERVISOR",135),("ESTATUS EXCEL",170),("SCORE",60),("FINAL",120)]
        h=ctk.CTkFrame(card,fg_color=Theme.colors.CARD_PURPLE,corner_radius=7); h.pack(fill="x",padx=10)
        for i,(txt,w) in enumerate(headers):
            h.grid_columnconfigure(i,weight=1 if i==0 else 0,minsize=w)
            ctk.CTkLabel(
                h, text=txt, text_color=Theme.colors.PRIMARY_DARK,
                font=ctk.CTkFont(family="Segoe UI",size=9,weight="bold"),
                anchor="w" if i<4 else "center"
            ).grid(row=0,column=i,padx=7,pady=8,sticky="ew")
        if not rows:
            ctk.CTkLabel(card,text="Sin resultados. Ingrese una palabra o seleccione un estatus.",text_color=Theme.colors.TEXT_SECONDARY).pack(pady=25)
        final_colors={
            "cumple": Theme.colors.SUCCESS,
            "parcial": Theme.colors.WARNING_HOVER,
            "no_cumple": Theme.colors.ERROR,
            "mal_dictaminado": Theme.colors.PRIMARY_DARK,
        }
        for row_no,r in enumerate(rows[:100]):
            row_bg=Theme.colors.TABLE_ROW if row_no % 2 == 0 else Theme.colors.TABLE_ROW_ALT
            line=ctk.CTkFrame(card,fg_color=row_bg,corner_radius=4); line.pack(fill="x",padx=10,pady=1)
            final={"cumple":"CUMPLIDO","parcial":"PARCIAL","no_cumple":"NO CUMPLIDO","mal_dictaminado":"MAL DICTAMINADO"}.get(r.get("final_status"),str(r.get("final_status","")))
            vals=[r.get("file_name",""),r.get("gestor",""),r.get("supervisor",""),r.get("excel_status",""),f'{float(r.get("score",0) or 0):.0f}%',final]
            for i,(val,(_,w)) in enumerate(zip(vals,headers)):
                line.grid_columnconfigure(i,weight=1 if i==0 else 0,minsize=w)
                color=final_colors.get(r.get("final_status"),Theme.colors.TEXT_PRIMARY) if i==5 else Theme.colors.TEXT_PRIMARY
                weight="bold" if i in (0,4,5) else "normal"
                ctk.CTkLabel(
                    line,text=str(val),text_color=color,
                    font=ctk.CTkFont(family="Segoe UI",size=9,weight=weight),
                    anchor="w" if i<4 else "center"
                ).grid(row=0,column=i,padx=7,pady=7,sticky="ew")
        if len(rows)>100:
            ctk.CTkLabel(card,text=f"Mostrando 100 de {len(rows)} resultados. La exportación incluye todos.",text_color=Theme.colors.TEXT_SECONDARY,font=ctk.CTkFont(family="Segoe UI",size=9)).pack(pady=(5,10))

    def _render_final_outputs(self):
        records=self._audit_records()
        counts=Counter(r.get("final_status") for r in records)
        self._section_title(self.detail_content,"SALIDAS FINALES DE AUDITORÍA",
                            "MAL DICTAMINADO tiene prioridad sobre el score cuando la validación de estatus detecta una discrepancia.")
        self._metric_strip(self.detail_content,[
            ("CUMPLIDO",counts.get("cumple",0),Theme.colors.CARD_GREEN),
            ("PARCIAL",counts.get("parcial",0),Theme.colors.CARD_PURPLE),
            ("NO CUMPLIDO",counts.get("no_cumple",0),Theme.colors.CARD_ORANGE),
            ("MAL DICTAMINADO",counts.get("mal_dictaminado",0),Theme.colors.CARD_BLUE),
        ])
        card=ctk.CTkFrame(self.detail_content,fg_color=Theme.colors.SURFACE,corner_radius=14,border_width=1,border_color=Theme.colors.BORDER_LIGHT)
        card.pack(fill="x",padx=5,pady=(0,12))
        ctk.CTkLabel(card,text="Generar carpetas de auditoría",font=ctk.CTkFont(family="Segoe UI",size=15,weight="bold"),anchor="w").pack(fill="x",padx=16,pady=(15,3))
        ctk.CTkLabel(card,text="Copia los audios; nunca mueve ni elimina los originales. También crea un Excel índice.",text_color=Theme.colors.TEXT_SECONDARY,font=ctk.CTkFont(family="Segoe UI",size=10),anchor="w").pack(fill="x",padx=16,pady=(0,10))
        actions=ctk.CTkFrame(card,fg_color="transparent"); actions.pack(fill="x",padx=16,pady=(0,16))
        ctk.CTkButton(actions,text="Escuchar audios",height=38,width=130,fg_color=Theme.colors.CARD_BLUE,text_color=Theme.colors.PRIMARY,
                      command=lambda:self._open_audio_player(records,"Salidas finales")).pack(side="left",padx=(0,8))
        ctk.CTkButton(actions,text="Generar CUMPLIDO / PARCIAL / NO CUMPLIDO / MAL DICTAMINADO",height=38,command=self._generate_final_folders).pack(side="left")

    def _open_audio_player(self, rows, title="Audios filtrados"):
        rows = list(rows or [])
        if not rows:
            messagebox.showinfo("MRS Speech Analytics", "No hay audios en el filtro actual para reproducir.")
            return
        host = self.master
        if hasattr(host, "show_audio_player"):
            host.show_audio_player(rows, title=title)
        else:
            messagebox.showerror("MRS Speech Analytics", "No se encontró la vista del reproductor de audios.")

    def _project_path(self):
        """Devuelve la misma ruta de proyecto que ya usa load_dashboard().

        No usa get_active_project(), porque esa API no existe en todas las
        versiones de ProjectManager y hacía que los botones de auditoría
        fallaran dentro del callback de Tkinter.
        """
        root = self._get_active_project_path()
        if root is None:
            return None
        try:
            root = Path(root).expanduser()
            if root.is_file():
                root = root.parent
            # Si get_project_path devuelve la carpeta contenedora, intenta
            # reconocer el proyecto por su estructura real.
            if not (root / "01_Audios").exists():
                project_label = str(getattr(self, "project_name_label", "") or "")
                # No dependemos de la etiqueta; normalmente root ya es correcto.
            return root
        except Exception:
            return None

    def _resolve_audio_source(self, record):
        """Localiza el audio original sin depender de una sola ruta histórica."""
        file_name = str(record.get("file_name") or "").strip()
        raw_path = str(record.get("audio_path") or "").strip()
        candidates = []
        if raw_path:
            candidates.append(Path(raw_path))
        root = self._project_path()
        if root and file_name:
            candidates.extend([
                root / "01_Audios" / file_name,
                root / file_name,
            ])
        for candidate in candidates:
            try:
                if candidate.is_file():
                    return candidate
            except OSError:
                pass
        # Último recurso: búsqueda por nombre dentro de 01_Audios y, si hace falta,
        # en la carpeta del proyecto. Evita fallar si cambió la ruta del proyecto.
        if root and file_name:
            for base in (root / "01_Audios", root):
                try:
                    if base.exists():
                        found = next((x for x in base.rglob(file_name) if x.is_file()), None)
                        if found:
                            return found
                except (OSError, StopIteration):
                    pass
        return None

    def _audit_output_root(self):
        root = self._project_path()
        if not root:
            return None
        try:
            out = Path(root) / "04_Reportes"
            out.mkdir(parents=True, exist_ok=True)
            return out
        except Exception as exc:
            messagebox.showerror("MRS Speech Analytics", f"No se pudo crear la carpeta de reportes.\n\nProyecto: {root}\nError: {exc}")
            return None

    def _copy_audit_records(self, rows, prefix):
        try:
            if not rows:
                messagebox.showinfo("MRS Speech Analytics", "No hay audios en el resultado actual.")
                return
            reports = self._audit_output_root()
            if not reports:
                messagebox.showwarning("MRS Speech Analytics", "No se pudo determinar la carpeta del proyecto activo.")
                return
            label = getattr(self, "_last_word_query", "") if prefix == "PALABRAS" else getattr(self, "_excel_status_selected", "TODOS")
            safe = re.sub(r'[^A-Za-z0-9_-]+', '_', str(label or "resultado")).strip('_')[:50] or "resultado"
            out = reports / "Auditoria_Filtros" / f"{prefix}_{safe}_{datetime.now():%Y%m%d_%H%M%S}"
            self._copy_rows_and_index(rows, out, quiet=False, prefix=prefix)
        except Exception as exc:
            messagebox.showerror("MRS Speech Analytics", f"Error al copiar audios.\n\n{type(exc).__name__}: {exc}")

    def _export_audit_report(self, rows, prefix):
        if not rows:
            messagebox.showinfo("MRS Speech Analytics", "No hay resultados para exportar.")
            return
        reports = self._audit_output_root()
        if not reports:
            messagebox.showwarning("MRS Speech Analytics", "No se pudo determinar la carpeta del proyecto activo.")
            return
        label = getattr(self, "_last_word_query", "") if prefix == "PALABRAS" else getattr(self, "_excel_status_selected", "TODOS")
        safe = re.sub(r'[^A-Za-z0-9_-]+', '_', str(label or "resultado")).strip('_')[:50] or "resultado"
        out = reports / "Reportes_Auditoria"
        out.mkdir(parents=True, exist_ok=True)
        filename = out / f"Reporte_{prefix}_{safe}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        try:
            self._write_audit_excel(rows, filename, prefix=prefix)
            messagebox.showinfo("MRS Speech Analytics", f"Reporte Excel generado correctamente.\n\nRegistros: {len(rows)}\nArchivo:\n{filename}")
        except Exception as exc:
            messagebox.showerror("MRS Speech Analytics", f"No se pudo generar el reporte Excel.\n\n{exc}")

    def _generate_final_folders(self):
        try:
            rows = self._audit_records()
            reports = self._audit_output_root()
            if not reports or not rows:
                messagebox.showwarning("MRS Speech Analytics", "No hay datos para generar las salidas o no se pudo localizar el proyecto activo.")
                return
            base = reports / "Auditoria_Final"
            names = {"cumple": "CUMPLIDO", "parcial": "PARCIAL", "no_cumple": "NO_CUMPLIDO", "mal_dictaminado": "MAL_DICTAMINADO"}
            total_copied = total_missing = 0
            errors = []
            for key, name in names.items():
                result = self._copy_rows_and_index([r for r in rows if r.get("final_status") == key], base / name, quiet=True, prefix="SALIDAS")
                total_copied += result["copied"]
                total_missing += result["missing"]
                errors.extend(result["errors"])
            msg = f"Auditoría generada.\n\nAudios copiados: {total_copied}\nNo encontrados: {total_missing}\nCarpeta:\n{base}"
            if errors:
                msg += "\n\nSe creó un registro de incidencias en las carpetas correspondientes."
            messagebox.showinfo("MRS Speech Analytics", msg)
        except Exception as exc:
            messagebox.showerror("MRS Speech Analytics", f"Error al generar las carpetas de auditoría.\n\n{type(exc).__name__}: {exc}")

    def _write_audit_excel(self, rows, filename, prefix=""):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "Auditoria"
        headers = ["audio_original", "gestor", "supervisor", "estatus_excel", "estatus_detectado", "score", "resultado_calidad", "resultado_final", "mal_dictaminado", "motivo_mala_dictaminacion"]
        if prefix == "PALABRAS":
            headers += ["palabra_frase", "coincidencias", "fragmento_evidencia"]
        ws.append(headers)
        query = str(getattr(self, "_last_word_query", "") or "").strip()
        for r in rows:
            row = [
                r.get("file_name", ""), r.get("gestor", ""), r.get("supervisor", ""),
                r.get("excel_status", ""), r.get("detected_status", ""), float(r.get("score", 0) or 0),
                r.get("status", ""), r.get("final_status", ""), "SI" if r.get("bad_dictation") else "NO",
                " | ".join(map(str, r.get("bad_reasons", []) or [])),
            ]
            if prefix == "PALABRAS":
                text = str(r.get("transcription_text") or "")
                fragment = ""
                if query:
                    pos = text.casefold().find(query.casefold())
                    if pos >= 0:
                        fragment = text[max(0, pos-100):min(len(text), pos+len(query)+140)].replace("\n", " ")
                row += [query, int(r.get("matches", 0) or 0), fragment]
            ws.append(row)
        fill = PatternFill("solid", fgColor="EDE5FA")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="3B1B5A")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {1: 48, 2: 22, 3: 22, 4: 24, 5: 24, 6: 12, 7: 20, 8: 22, 9: 18, 10: 65, 11: 25, 12: 15, 13: 80}
        for i, width in widths.items():
            if i <= ws.max_column:
                ws.column_dimensions[get_column_letter(i)].width = width
        filename = Path(filename)
        filename.parent.mkdir(parents=True, exist_ok=True)
        wb.save(filename)

    def _copy_rows_and_index(self, rows, out, quiet=False, prefix=""):
        out = Path(out)
        out.mkdir(parents=True, exist_ok=True)
        copied = 0
        missing = 0
        errors = []
        for r in rows:
            src = self._resolve_audio_source(r)
            if src is None:
                missing += 1
                errors.append(f"NO ENCONTRADO: {r.get('file_name','')}")
                continue
            try:
                shutil.copy2(str(src), str(out / src.name))
                copied += 1
            except OSError as exc:
                missing += 1
                errors.append(f"ERROR COPIA: {r.get('file_name','')} | {exc}")
        try:
            self._write_audit_excel(rows, out / "Indice_Auditoria.xlsx", prefix=prefix)
        except Exception as exc:
            errors.append(f"ERROR EXCEL: {exc}")
        if errors:
            try:
                (out / "Audios_No_Encontrados.txt").write_text("\n".join(errors), encoding="utf-8")
            except OSError:
                pass
        if not quiet:
            messagebox.showinfo(
                "MRS Speech Analytics",
                f"Proceso terminado.\n\nAudios copiados: {copied}\nNo encontrados: {missing}\nResultados: {len(rows)}\nCarpeta:\n{out}"
            )
        return {"copied": copied, "missing": missing, "errors": errors, "out": out}

    # =========================================================
    # LIMPIAR
    # =========================================================

    def _clear_dashboard(self):
        self.dashboard_data = {}

        self.total_calls_value.configure(
            text="0"
        )

        self.compliance_value.configure(
            text="0%"
        )

        self.main_reason_value.configure(
            text="Sin datos"
        )

        self.main_sentiment_value.configure(
            text="Sin datos"
        )

        self.sentiment_chart.set_data(
            {
                "positivo": 0,
                "neutral": 0,
                "negativo": 0,
            }
        )

        self.compliance_chart.set_data(
            {
                "cumple": 0,
                "parcial": 0,
                "no_cumple": 0,
            }
        )

        self.reasons_chart.set_data(
            {}
        )

        self.keywords_chart.set_data(
            {}
        )

    # =========================================================
    # UTILIDADES
    # =========================================================

    @staticmethod
    def _safe_int(
        value,
    ) -> int:
        try:
            return max(
                int(
                    value
                    or 0
                ),
                0,
            )

        except (
            TypeError,
            ValueError,
        ):
            return 0

    # =========================================================
    # PROYECTO ACTIVO
    # =========================================================

    def _get_active_project_path(
        self,
    ) -> Optional[Path]:
        try:
            get_path = getattr(
                project_manager,
                "get_project_path",
                None,
            )

            if callable(get_path):
                path = get_path()

                if path:
                    return (
                        Path(path)
                        .expanduser()
                        .resolve()
                    )

            get_current = getattr(
                project_manager,
                "get_current_project",
                None,
            )

            if callable(get_current):
                project = get_current()

                if isinstance(
                    project,
                    dict,
                ):
                    raw_path = project.get(
                        "ruta"
                    )

                    if raw_path:
                        return (
                            Path(raw_path)
                            .expanduser()
                            .resolve()
                        )

        except Exception:
            return None

        return None

    def _get_active_project_name(
        self,
    ) -> str:
        try:
            get_name = getattr(
                project_manager,
                "get_project_name",
                None,
            )

            if callable(get_name):
                name = get_name()

                if name:
                    return str(name)

        except Exception:
            pass

        return "Proyecto activo"