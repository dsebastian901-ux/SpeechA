"""
MRS Speech Analytics Professional
Servicio de relación Audios + Excel.

Responsabilidades:
- Leer la plantilla de relación de audios.
- Validar columnas obligatorias y registros.
- Relacionar archivos de audio con gestor, estatus y nombre final.
- Consultar el puntaje mínimo configurable desde SettingsService.
- Clasificar resultados como APROBADO / REVISAR.
- Evitar sobrescribir audios originales al preparar nombres de salida.
"""

from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
import re
import unicodedata

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "relation_service.py requiere openpyxl para leer archivos Excel (.xlsx)."
    ) from exc

try:
    from app.services.settings_service import SettingsService
except Exception:
    SettingsService = None  # Permite pruebas aisladas del módulo.


@dataclass
class RelationRecord:
    archivo_original: str
    gestor: str = ""
    estatus_grabacion: str = ""
    nombre_final: str = ""
    id_llamada: str = ""
    campana: str = ""
    fecha: str = ""
    supervisor: str = ""
    numero_de_tarjeta: str = ""
    fecha_asignacion: str = ""
    corte: str = ""
    vencimiento: str = ""
    asignacion: str = ""
    pago_minimo: str = ""
    saldo_total: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


class RelationService:
    """Gestiona el modo de procesamiento Audios + Excel."""

    SHEET_RELATION = "Relacion_Audios"
    SHEET_CATALOGS = "Catalogos"

    REQUIRED_COLUMNS = (
        "archivo_original",
        "gestor",
        "estatus_grabacion",
        "nombre_final",
    )

    # La columna nombre_final debe existir en la plantilla, pero su valor
    # solo es obligatorio cuando una llamada vaya a aprobarse/renombrarse.
    # Estatus no aprobables (BUZON, TITULAR CUELGA, REMARCAR, etc.) pueden
    # conservar la relación aunque esta celda esté vacía.
    REQUIRED_ROW_FIELDS = (
        "archivo_original",
        "gestor",
        "estatus_grabacion",
    )

    OPTIONAL_COLUMNS = (
        "id_llamada",
        "campana",
        "fecha",
        "supervisor",
        "numero_de_tarjeta",
        "fecha_asignacion",
        "corte",
        "vencimiento",
        "asignacion",
        "estatus_autorizado",
        "validacion_negociacion",
        "validacion_fila",
        "pago_minimo",
        "saldo_total",
    )

    AUDIO_EXTENSIONS = {
        ".wav",
        ".mp3",
        ".m4a",
        ".flac",
        ".ogg",
        ".aac",
        ".wma",
        ".mp4",
        ".mpeg",
        ".mpg",
        ".webm",
    }

    DEFAULT_MINIMUM_SCORE = 80

    def __init__(self) -> None:
        self.records: List[RelationRecord] = []
        self.authorized_statuses: List[str] = []
        self.source_excel: Optional[Path] = None

    # ------------------------------------------------------------------
    # Configuración
    # ------------------------------------------------------------------
    def get_minimum_score(self) -> int:
        """Obtiene el puntaje mínimo desde SettingsService.

        Si la configuración no está disponible, utiliza 80 como valor seguro.
        El valor siempre queda limitado entre 0 y 100.
        """
        value: Any = self.DEFAULT_MINIMUM_SCORE

        if SettingsService is not None:
            try:
                service = SettingsService()

                if hasattr(service, "get"):
                    value = service.get("minimum_score", self.DEFAULT_MINIMUM_SCORE)
                elif hasattr(service, "get_setting"):
                    value = service.get_setting(
                        "minimum_score", self.DEFAULT_MINIMUM_SCORE
                    )
                elif hasattr(service, "settings"):
                    value = getattr(service, "settings", {}).get(
                        "minimum_score", self.DEFAULT_MINIMUM_SCORE
                    )
            except Exception:
                value = self.DEFAULT_MINIMUM_SCORE

        try:
            value = int(float(value))
        except (TypeError, ValueError):
            value = self.DEFAULT_MINIMUM_SCORE

        return max(0, min(100, value))

    # ------------------------------------------------------------------
    # Lectura de Excel
    # ------------------------------------------------------------------
    def load_excel(self, excel_path: str | Path) -> Dict[str, Any]:
        """Carga y valida el archivo Excel de relación.

        Devuelve un resumen con registros válidos, errores y estatus autorizados.
        """
        path = Path(excel_path)
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {path}")
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("El archivo de relación debe ser .xlsx o .xlsm")

        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)

        if self.SHEET_RELATION not in workbook.sheetnames:
            raise ValueError(
                f"El Excel debe contener la hoja '{self.SHEET_RELATION}'."
            )

        self.authorized_statuses = self._read_authorized_statuses(workbook)
        records, errors = self._read_relation_records(workbook[self.SHEET_RELATION])

        # v3.6.16: si el archivo no incluye hoja Catalogos, deriva los estatus
        # autorizados de las filas que ya traen nombre_final. Esto evita que un
        # Excel operativo válido convierta todos los >=80 en APROBADO_SIN_ETIQUETAR.
        if not self.authorized_statuses:
            derived = []
            seen = set()
            for record in records:
                status = str(record.estatus_grabacion or '').strip()
                if not status or not str(record.nombre_final or '').strip():
                    continue
                key = self._normalize_text(status)
                if key and key not in seen:
                    seen.add(key)
                    derived.append(status)
            self.authorized_statuses = derived

        self.records = records
        self.source_excel = path
        workbook.close()

        return {
            "archivo": str(path),
            "total_registros": len(records),
            "errores": errors,
            "total_errores": len(errors),
            "estatus_autorizados": list(self.authorized_statuses),
            "puntaje_minimo": self.get_minimum_score(),
        }

    def _read_relation_records(
        self, worksheet: Any
    ) -> Tuple[List[RelationRecord], List[Dict[str, Any]]]:
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            raise ValueError("La hoja Relacion_Audios está vacía.")

        headers = [self._normalize_header(value) for value in header_row]
        for name in ('asignacion','mora'):
            if headers.count(name)>1:
                raise ValueError(f'Columna de mora duplicada: {name}.')
        header_map = {name: index for index, name in enumerate(headers) if name}

        missing = [name for name in self.REQUIRED_COLUMNS if name not in header_map]
        if missing:
            raise ValueError(
                "Faltan columnas obligatorias en Relacion_Audios: "
                + ", ".join(missing)
            )

        records: List[RelationRecord] = []
        errors: List[Dict[str, Any]] = []
        seen_files: set[str] = set()

        for excel_row_number, row in enumerate(rows, start=2):
            if self._row_is_empty(row):
                continue

            values: Dict[str, str] = {}
            record_columns = (
                *self.REQUIRED_COLUMNS,
                "id_llamada",
                "campana",
                "fecha",
                "supervisor",
                "numero_de_tarjeta",
                "fecha_asignacion",
                "corte",
                "vencimiento",
                "asignacion",
                "pago_minimo",
                "saldo_total",
            )
            for column in record_columns:
                index = header_map.get(column)
                values[column] = self._clean_cell(
                    row[index] if index is not None and index < len(row) else ""
                )

            row_errors: List[str] = []
            # El Excel operativo usa «mora»; «asignacion» es el nombre interno.
            # Si ambas existen, no escoger silenciosamente valores distintos.
            mora_index=header_map.get('mora')
            mora=self._clean_cell(row[mora_index] if mora_index is not None and mora_index<len(row) else '')
            if values['asignacion'] and mora and self._mora_key(values['asignacion'])!=self._mora_key(mora):
                row_errors.append('Las columnas mora y asignacion no coinciden.')
            elif not values['asignacion']:
                values['asignacion']=mora

            for column in self.REQUIRED_ROW_FIELDS:
                if not values[column]:
                    row_errors.append(f"Campo obligatorio vacío: {column}")

            original_key = values["archivo_original"].casefold()
            if original_key and original_key in seen_files:
                row_errors.append("archivo_original duplicado")
            elif original_key:
                seen_files.add(original_key)

            if values["archivo_original"]:
                suffix = Path(values["archivo_original"]).suffix.lower()
                if suffix and suffix not in self.AUDIO_EXTENSIONS:
                    row_errors.append(
                        f"Extensión de audio no reconocida: {suffix}"
                    )

            if values["nombre_final"]:
                invalid_name_reason = self._validate_output_name(values["nombre_final"])
                if invalid_name_reason:
                    row_errors.append(invalid_name_reason)

            if row_errors:
                errors.append(
                    {
                        "fila": excel_row_number,
                        "archivo_original": values["archivo_original"],
                        "errores": row_errors,
                    }
                )
                continue

            records.append(RelationRecord(**values))

        return records, errors

    def _read_authorized_statuses(self, workbook: Any) -> List[str]:
        """Lee únicamente los estatus activos de la hoja Catalogos.

        Formato esperado:
        - columna: estatus_autorizado / estatus_autorizados / estatus
        - columna opcional: activo

        Si existe la columna ``activo``, solo se consideran autorizados
        los registros marcados como SI. No se leen notas ni otras columnas.
        """
        if self.SHEET_CATALOGS not in workbook.sheetnames:
            return []

        worksheet = workbook[self.SHEET_CATALOGS]
        rows = worksheet.iter_rows(values_only=True)

        try:
            header_row = next(rows)
        except StopIteration:
            return []

        headers = [
            self._normalize_header(value)
            for value in header_row
        ]

        status_aliases = {
            "estatus_autorizado",
            "estatus_autorizados",
            "estatus",
        }

        status_index = next(
            (
                index
                for index, header in enumerate(headers)
                if header in status_aliases
            ),
            None,
        )

        if status_index is None:
            return []

        active_index = next(
            (
                index
                for index, header in enumerate(headers)
                if header in {"activo", "autorizado", "habilitado"}
            ),
            None,
        )

        values: List[str] = []

        for row in rows:
            if status_index >= len(row):
                continue

            status = self._clean_cell(
                row[status_index]
            )

            if not status:
                continue

            if active_index is not None:
                active_value = self._clean_cell(
                    row[active_index]
                    if active_index < len(row)
                    else ""
                )

                normalized_active = self._normalize_text(
                    active_value
                )

                if normalized_active not in {
                    "si",
                    "sí",
                    "s",
                    "1",
                    "true",
                    "verdadero",
                    "activo",
                }:
                    continue

            values.append(status)

        unique: List[str] = []
        seen: set[str] = set()

        for value in values:
            key = self._normalize_text(value)

            if key in seen:
                continue

            seen.add(key)
            unique.append(value)

        return unique

    # ------------------------------------------------------------------
    # Relación Audio -> Excel
    # ------------------------------------------------------------------
    def find_record(self, audio_file: str | Path) -> Optional[RelationRecord]:
        """Busca una relación por nombre de archivo, sin importar mayúsculas."""
        target_name = Path(audio_file).name.casefold()

        for record in self.records:
            if Path(record.archivo_original).name.casefold() == target_name:
                return record

        return None

    def validate_audio_files(
        self, audio_files: Iterable[str | Path]
    ) -> Dict[str, Any]:
        """Compara los audios seleccionados contra el Excel cargado."""
        audio_paths = [Path(item) for item in audio_files]
        existing_names = {path.name.casefold(): path for path in audio_paths}

        matched: List[Dict[str, Any]] = []
        missing_in_excel: List[str] = []
        missing_audio: List[str] = []

        relation_names = {
            Path(record.archivo_original).name.casefold(): record
            for record in self.records
        }

        for path in audio_paths:
            record = relation_names.get(path.name.casefold())
            if record:
                matched.append(
                    {
                        "audio": str(path),
                        "relacion": record.to_dict(),
                    }
                )
            else:
                missing_in_excel.append(str(path))

        for record in self.records:
            name = Path(record.archivo_original).name.casefold()
            if name not in existing_names:
                missing_audio.append(record.archivo_original)

        return {
            "coincidencias": matched,
            "total_coincidencias": len(matched),
            "audios_sin_relacion": missing_in_excel,
            "registros_sin_audio": missing_audio,
            "puntaje_minimo": self.get_minimum_score(),
        }

    # ------------------------------------------------------------------
    # Clasificación por calidad y estatus
    # ------------------------------------------------------------------
    def is_status_authorized(self, status: str) -> bool:
        """Indica si el estatus del Excel está autorizado para grabación.

        Si no existe catálogo de estatus, no autoriza automáticamente; esto
        evita aprobar una llamada por accidente.
        """
        if not self.authorized_statuses:
            return False

        target = self._normalize_text(status)
        return any(
            self._normalize_text(item) == target
            for item in self.authorized_statuses
        )

    def classify(
        self,
        compliance_score: float | int,
        record: Optional[RelationRecord] = None,
    ) -> Dict[str, Any]:
        """Clasifica una llamada usando el umbral configurable.

        Con Excel, APROBADO exige simultáneamente:
        - score >= puntaje mínimo configurado
        - estatus autorizado para grabación
        - nombre_final disponible

        Sin relación Excel, la clasificación se basa únicamente en score.
        """
        minimum_score = self.get_minimum_score()

        try:
            score = float(compliance_score)
        except (TypeError, ValueError):
            score = 0.0

        passes_score = score >= minimum_score

        if record is None:
            approved = passes_score
            authorized_status = None
            has_output_name = None
            reason = (
                f"Score {score:.1f} >= {minimum_score}"
                if approved
                else f"Score {score:.1f} < {minimum_score}"
            )
        else:
            authorized_status = self.is_status_authorized(
                record.estatus_grabacion
            )
            has_output_name = bool(record.nombre_final.strip())
            approved = passes_score and authorized_status and has_output_name

            reasons: List[str] = []
            if not passes_score:
                reasons.append(f"score menor a {minimum_score}")
            if not authorized_status:
                reasons.append("estatus no autorizado")
            if not has_output_name:
                reasons.append("nombre_final vacío")
            reason = "Cumple reglas de aprobación" if approved else "; ".join(reasons)

        if approved:
            base_classification = "APROBADO"
        elif not passes_score:
            base_classification = "NO_APROBADO_MENOR_80"
        elif record is not None and authorized_status is False:
            # Cumple el score, pero el estatus no está autorizado para reetiquetado.
            base_classification = "APROBADO_SIN_ETIQUETAR"
        else:
            base_classification = "REVISAR"

        return {
            "clasificacion": base_classification,
            "aprobado": approved,
            "score": round(score, 2),
            "puntaje_minimo": minimum_score,
            "cumple_puntaje": passes_score,
            "estatus_autorizado": authorized_status,
            "tiene_nombre_final": has_output_name,
            "motivo": reason,
        }

    # ------------------------------------------------------------------
    # Nombre seguro de salida
    # ------------------------------------------------------------------
    def build_output_filename(
        self,
        original_audio: str | Path,
        record: RelationRecord,
    ) -> str:
        """Genera el nombre final manteniendo la extensión del audio original."""
        original_path = Path(original_audio)
        desired = Path(record.nombre_final.strip())

        stem = desired.stem if desired.suffix else desired.name
        stem = self._sanitize_filename(stem)

        if not stem:
            stem = self._sanitize_filename(original_path.stem)

        extension = original_path.suffix.lower()
        return f"{stem}{extension}"

    def get_safe_output_path(
        self,
        output_directory: str | Path,
        original_audio: str | Path,
        record: RelationRecord,
    ) -> Path:
        """Devuelve una ruta que no sobrescribe archivos existentes."""
        output_dir = Path(output_directory)
        filename = self.build_output_filename(original_audio, record)
        candidate = output_dir / filename

        if not candidate.exists():
            return candidate

        stem = candidate.stem
        suffix = candidate.suffix
        counter = 2

        while True:
            alternate = output_dir / f"{stem}_{counter}{suffix}"
            if not alternate.exists():
                return alternate
            counter += 1

    # ------------------------------------------------------------------
    # Utilidades
    # ------------------------------------------------------------------
    @staticmethod
    def _mora_key(value: Any) -> str:
        token=re.sub(r'[\s_-]','',str(value or '').upper())
        token=re.sub(r'^MORA','',token)
        return token+'PV' if token in {'1','2','3','4','5','6'} else token

    @staticmethod
    def _clean_cell(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _row_is_empty(row: Iterable[Any]) -> bool:
        return all(value is None or str(value).strip() == "" for value in row)

    @classmethod
    def _normalize_header(cls, value: Any) -> str:
        text = cls._normalize_text(cls._clean_cell(value))
        text = re.sub(r"[^a-z0-9]+", "_", text)
        return text.strip("_")

    @staticmethod
    def _normalize_text(value: str) -> str:
        text = unicodedata.normalize("NFKD", str(value))
        text = "".join(char for char in text if not unicodedata.combining(char))
        return re.sub(r"\s+", " ", text).strip().casefold()

    @staticmethod
    def _sanitize_filename(value: str) -> str:
        # Caracteres inválidos en nombres de archivo de Windows.
        cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value).strip()
        cleaned = cleaned.rstrip(". ")
        return cleaned[:180]

    @staticmethod
    def _validate_output_name(value: str) -> Optional[str]:
        name = Path(value).name.strip()
        if not name:
            return "nombre_final inválido"
        if re.search(r'[<>:"/\\|?*\x00-\x1F]', name):
            return "nombre_final contiene caracteres inválidos para Windows"
        return None


relation_service = RelationService()
