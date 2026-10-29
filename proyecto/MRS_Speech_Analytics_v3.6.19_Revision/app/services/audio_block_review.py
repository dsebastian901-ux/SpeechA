"""Lectura ASR adicional en archivos separados; nunca valida voces ni cambia el lote."""
from copy import deepcopy
from dataclasses import asdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from time import perf_counter

from app.services.analytics_service import AnalyticsService
from app.services.checkpoint_output import create_checkpoint_output
from app.services.asr_evidence_quality import block_reading_quality
from app.services.audio_block_evidence import binding

AUDIO_SUFFIXES={'.ogg','.wav','.mp3','.m4a','.flac','.aac','.wma'}


def sha256_file(path):
    digest=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024*1024),b''):
            digest.update(block)
    return digest.hexdigest()


def audio_duration(path):
    import av
    duration=0.0
    with av.open(str(path)) as container:
        for frame in container.decode(audio=0):
            duration+=frame.samples/frame.sample_rate
    if duration<=0:
        raise ValueError('El audio no contiene muestras decodificables.')
    return duration


def select_audio_matches(payload, directory):
    """Cruce exacto sin adivinar por tarjeta, gestor, orden o índice de fila."""
    folder=Path(directory).expanduser().resolve()
    if not folder.is_dir():
        raise ValueError('La carpeta de audios no existe.')
    aliases={}
    for index,item in enumerate(payload['items'],1):
        names=[item.get('file_name'),(item.get('result') or {}).get('archivo')]
        relation=item.get('relation') or {}
        for name in [relation.get('nombre_final'),(item.get('output_name_review') or {}).get('previous_name')]:
            if name:
                names.extend([name+ext for ext in AUDIO_SUFFIXES] if Path(name).suffix.lower() not in AUDIO_SUFFIXES else [name])
        for name in names:
            if name:
                key=str(name).replace('\\','/').rsplit('/',1)[-1].casefold()
                aliases.setdefault(key,set()).add(index)
    matches={};unmatched=[]
    for path in sorted(folder.rglob('*')):
        if not path.is_file() or path.suffix.lower() not in AUDIO_SUFFIXES:
            continue
        indices=aliases.get(path.name.casefold(),set())
        if not indices:
            unmatched.append(path.name);continue
        if len(indices)!=1:
            raise ValueError('Un audio coincide con varias llamadas; revise nombres duplicados.')
        index=next(iter(indices))
        if index in matches:
            raise ValueError('Hay dos archivos para una misma llamada; retire duplicados antes de verificar.')
        matches[index]=path
    if not matches:
        raise ValueError('Ningún audio coincide por nombre original o nombre de salida con el checkpoint.')
    return sorted(matches.items()),unmatched


class CachedWhisperReader:
    """Misma implementación Faster-Whisper, exclusivamente con modelo local."""
    def __init__(self, model_name, device, compute_type):
        from faster_whisper import WhisperModel
        try:
            self.model=WhisperModel(model_name,device=device,compute_type=compute_type,
                                    local_files_only=True,cpu_threads=4)
        except Exception as error:
            raise RuntimeError('No se pudo cargar el modelo Whisper local. Ejecute esta herramienta '
                'en la misma sesión de Windows donde ya funciona el programa. No se descarga otro modelo. '
                f'Detalle: {error}') from error

    def __call__(self,path,duration):
        segments,info=self.model.transcribe(str(path),language='es',task='transcribe',
            beam_size=5,vad_filter=True,vad_parameters={'min_silence_duration_ms':700,'speech_pad_ms':200},
            word_timestamps=True,condition_on_previous_text=False,temperature=0)
        values=[asdict(s) for s in segments]
        # Segunda lectura del cierre, sin VAD, para conservar frases finales
        # que una detección de voz puede omitir. Es otra hipótesis ASR, no prueba humana.
        from faster_whisper.audio import decode_audio
        rate=16000;decoded=decode_audio(str(path),sampling_rate=rate)
        offset=max(0,len(decoded)/rate-20)
        tail,_=self.model.transcribe(decoded[int(offset*rate):],language='es',task='transcribe',
            beam_size=5,vad_filter=False,word_timestamps=True,condition_on_previous_text=False,temperature=0)
        tail_values=[]
        for segment in tail:
            item=asdict(segment);item['start']+=offset;item['end']+=offset
            for word in item.get('words') or []:
                word['start']+=offset;word['end']+=offset
            tail_values.append(item)
        return {'text':' '.join(s['text'].strip() for s in values),'segments':values,
                'tail_text':' '.join(s['text'].strip() for s in tail_values),
                'tail_segments':tail_values,'language_probability':info.language_probability}


def _save_report(payload,path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font
    book=Workbook();sheet=book.active;sheet.title='COMPARACION_BLOQUE'
    sheet.append(['Llamada del lote','Archivo original','Estado de lectura','Duración audio',
                  'Texto original conservado','Nueva lectura ASR','Cierre sin VAD','Motivo previo de revisión','Error',
                  'Alertas técnicas de lectura','Detalle de alertas ASR'])
    for item in payload['items']:
        sheet.append([item['case'],item['file_name'],item['status'],item.get('audio_duration'),
                      item['original_text'][:30000],item.get('text','')[:30000],
                      item.get('tail_text','')[:30000],item['previous_reason'][:30000],item.get('error','')[:30000],
                      'REVISAR LECTURA ASR' if (item.get('quality') or {}).get('degraded') else '',
                      json.dumps(item.get('quality') or {},ensure_ascii=False)[:30000]])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=55
    note=book.create_sheet('ALCANCE');note.append(['Lectura automática con el modelo Whisper ya instalado.'])
    note.append(['La lectura nueva y la del cierre son hipótesis ASR, no escucha humana ni identificación de voces.'])
    note.append(['No se cambian transcripciones, validaciones, dictámenes, nombres ni audios del checkpoint original.'])
    note.append(['Errores de lectura',payload['errors']]);note.append(['Lecturas completas',payload['completed_calls']])
    note.append(['Lecturas con advertencias técnicas',payload.get('degraded_calls',0)])
    note.append(['Las alertas técnicas no son infracciones del gestor ni respuestas confirmadas del interlocutor.'])
    book.save(path);book.close()


def _asr_segments(values):
    if not isinstance(values,list):
        raise ValueError('La lectura no contiene una lista de segmentos ASR.')
    keys={'id','seek','start','end','text','tokens','temperature','avg_logprob',
          'compression_ratio','no_speech_prob','words'}
    result=[]
    for value in values:
        if not isinstance(value,dict) or not isinstance(value.get('text'),str):
            raise ValueError('Segmento ASR no válido.')
        segment={k:deepcopy(v) for k,v in value.items() if k in keys}
        if segment.get('words'):
            segment['words']=[{k:deepcopy(v) for k,v in word.items()
                               if k in {'start','end','word','probability'}} for word in segment['words']]
        result.append(segment)
    return result


def review_audio_block(checkpoint_path,audio_directory,output_directory=None,*,reader_factory=None,
                       duration_reader=None,progress=None):
    source=Path(checkpoint_path).expanduser().resolve();raw=source.read_bytes()
    original=json.loads(raw.decode('utf-8-sig'))
    if not isinstance(original,dict) or not isinstance(original.get('items'),list) or not original['items']:
        raise ValueError('El checkpoint no contiene llamadas.')
    matches,unmatched=select_audio_matches(original,audio_directory)
    if len(matches)>10:
        raise ValueError('La carpeta contiene más de 10 llamadas del checkpoint. Seleccione solo el bloque para revisar.')
    for index,_ in matches:
        data=original['items'][index-1].get('result') or {}
        if not isinstance(data.get('texto'),str) or not data['texto'].strip() or not data.get('duracion_segundos'):
            raise ValueError('Seleccione el checkpoint del lote con sus transcripciones y duraciones, no el JSON de lectura adicional.')
    models={str((original['items'][index-1].get('result') or {}).get('modelo') or 'small') for index,_ in matches}
    if len(models)!=1 or not re.fullmatch(r'[a-zA-Z0-9_.-]+',next(iter(models))):
        raise ValueError('El bloque mezcla modelos o contiene un nombre de modelo no válido; no se elige otro automáticamente.')
    model_name=next(iter(models))
    from app.services.whisper_service import WhisperService
    runtime=WhisperService.get_runtime_configuration()
    # Cargar antes de crear resultados: un modelo ausente no genera una lectura vacía exitosa.
    factory=reader_factory or CachedWhisperReader
    reader=factory(model_name,runtime['device'],runtime['compute_type'])
    target=create_checkpoint_output(output_directory,'Revision_Audios_'+AnalyticsService.ENGINE_REVISION.rsplit('-',1)[-1]+'_')
    payload={'schema':1,'engine_revision':AnalyticsService.ENGINE_REVISION,'source_checkpoint_sha256':hashlib.sha256(raw).hexdigest(),
             'model':model_name,'device':runtime['device'],'compute_type':runtime['compute_type'],
             'beam_size':5,'word_timestamps':True,'tail_vad_filter':False,'matched_calls':len(matches),
             'unmatched_audio_names':unmatched,'human_verified':False,'speaker_roles_confirmed':False,
             'checkpoint_modified':False,'audio_files_copied':False,'whisper_executed':False,
             'completed':False,'completed_calls':0,'errors':0,'items':[]}
    destination=target/'revision_audios_bloque.json'
    for number,(index,path) in enumerate(matches,1):
        original_item=original['items'][index-1];data=original_item.get('result') or {}
        item={'case':index,'file_name':original_item['file_name'],'audio_name':path.name,
              'audio_sha256':sha256_file(path),'original_text':data.get('texto',''),
              'original_segments':deepcopy(data.get('segmentos',[])),
              'original_text_sha256':hashlib.sha256(data.get('texto','').encode()).hexdigest(),
              'previous_reason':(original_item.get('classification') or {}).get('motivo',''),
              'status':'error','human_verified':False,'speaker_roles_confirmed':False}
        item['source_binding']=binding(data,original_item.get('relation') or {})
        started=perf_counter()
        if progress:progress(f'Leyendo audio {number}/{len(matches)}: llamada {index}')
        try:
            duration=(duration_reader or audio_duration)(path);item['audio_duration']=duration
            expected=float(data.get('duracion_segundos') or 0)
            if expected<=0 or abs(duration-expected)>0.25:
                raise ValueError('Duración distinta del checkpoint; no se atribuye la lectura a esta llamada.')
            reading=reader(path,duration)
            if not isinstance(reading,dict) or not isinstance(reading.get('text'),str) or not reading['text'].strip() or not reading.get('segments'):
                raise ValueError('Whisper no produjo segmentos: revisar silencio, archivo o modelo.')
            # Campos controlados del resultado: una implementación alternativa no
            # puede introducir validaciones humanas, aprobaciones o rutas de salida.
            item.update({'text':reading['text'],'tail_text':str(reading.get('tail_text') or ''),
                         'segments':_asr_segments(reading['segments']),
                         'tail_segments':_asr_segments(reading.get('tail_segments') or [])})
            if sha256_file(path)!=item['audio_sha256']:
                raise ValueError('El archivo de audio cambió durante la lectura; no se valida esta evidencia.')
            item['quality']=block_reading_quality(item['segments'],item['tail_segments'])
            item['status']='completado';payload['completed_calls']+=1;payload['whisper_executed']=True
        except Exception as error:
            item['error']=str(error);payload['errors']+=1
        item['elapsed_seconds']=round(perf_counter()-started,3);payload['items'].append(item)
        payload['updated_at']=datetime.now(timezone.utc).isoformat()
        destination.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    payload['completed']=payload['errors']==0
    payload['degraded_calls']=sum(bool((item.get('quality') or {}).get('degraded')) for item in payload['items'])
    destination.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    report=target/'Comparacion_Bloque_Audios.xlsx';_save_report(payload,report)
    assert source.read_bytes()==raw,'El checkpoint original cambió durante la lectura.'
    return {'evidence':str(destination),'report':str(report),'matched_calls':len(matches),
            'completed_calls':payload['completed_calls'],'errors':payload['errors']}
