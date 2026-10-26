"""Datos de origen: la identidad del nombre de salida no prueba una falta en audio."""
import re
from pathlib import PureWindowsPath


def validate_account_source(relation):
    name = PureWindowsPath(str(relation.get("nombre_final") or "")).name
    match = re.match(r"^(0*\d{16})(?=_)", name)
    card = str(relation.get("numero_de_tarjeta") or "").strip()
    if not match:
        return {"review_required": False, "reason": ""}
    encoded = match[1][-16:]
    if not re.fullmatch(r"\d{16}", card):
        return {"review_required": True, "reason": "El nombre final contiene una tarjeta sin referencia válida de origen."}
    mismatch = encoded != card
    return {"review_required": mismatch, "expected_last4": card[-4:], "name_last4": encoded[-4:],
            "reason": (f"Datos de cuenta: nombre_final termina en {encoded[-4:]} y numero_de_tarjeta en {card[-4:]}; "
                       "confirmar el identificador antes de renombrar.") if mismatch else ""}


def corrected_account_name(relation):
    if not validate_account_source(relation)["review_required"]:
        raise ValueError("No hay una diferencia de identificador que corregir.")
    card = str(relation.get("numero_de_tarjeta") or "").strip()
    if not re.fullmatch(r"\d{16}", card):
        raise ValueError("Se requiere numero_de_tarjeta de 16 dígitos.")
    name = str(relation.get("nombre_final") or "")
    match = re.match(r"^(0*\d{16})(?=_)", name)
    if not match:
        raise ValueError("El nombre final no tiene el formato de tarjeta reconocido.")
    prefix = match[1][:-16]
    return prefix+card+name[match.end():]


def write_corrected_relation_excel(source_excel, corrections, destination):
    """Solo filas confirmadas; no se cambia numero_de_tarjeta ni el original."""
    from pathlib import Path
    from openpyxl import load_workbook
    source, target = Path(source_excel), Path(destination)
    if source.resolve() == target.resolve() or target.exists():
        raise ValueError("No se sobrescribe el Excel original ni una copia existente.")
    workbook=load_workbook(source,keep_vba=source.suffix.lower()==".xlsm")
    try:
        sheet=workbook["Relacion_Audios"]
        headers={str(c.value or "").strip().lower():c.column for c in sheet[1]}
        for name in ("archivo_original","nombre_final","numero_de_tarjeta"):
            if name not in headers:raise ValueError(f"Falta la columna {name} en el Excel.")
        applied=set()
        for row in range(2,sheet.max_row+1):
            name=str(sheet.cell(row,headers["archivo_original"]).value or "")
            if name not in corrections:continue
            entry=corrections[name]
            if name in applied:raise ValueError("El Excel contiene duplicados para una cuenta revisada.")
            card=str(sheet.cell(row,headers["numero_de_tarjeta"]).value or "").strip()
            if card != entry["numero_de_tarjeta"]:
                raise ValueError("La tarjeta del Excel cambió respecto al checkpoint; revisar datos de origen.")
            sheet.cell(row,headers["nombre_final"]).value=entry["nombre_final"]
            applied.add(name)
        if applied!=set(corrections):raise ValueError("No se encontraron todas las llamadas revisadas en el Excel.")
        workbook.save(target)
    finally:
        workbook.close()
    return str(target)
