"""Reanálisis de reglas sin retranscripción, sin copiar audios y sin alterar el origen."""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from app.models.transcriptions import TranscriptionModel
from app.services.analytics_service import AnalyticsService
from app.services.batch_report_service import BatchReportService
from app.services.transcription_services import TranscriptionService
from app.services.checkpoint_output import create_checkpoint_output


def reanalyze_checkpoint(checkpoint_path, output_directory=None):
    source = Path(checkpoint_path).expanduser().resolve()
    raw = source.read_bytes()
    payload = json.loads(raw.decode("utf-8-sig"))
    if not isinstance(payload, dict) or not isinstance(payload.get("items"), list) or not payload["items"]:
        raise ValueError("El checkpoint no contiene llamadas para reanalizar.")
    target = create_checkpoint_output(output_directory,
        "Reanalisis_MRS_" + AnalyticsService.ENGINE_REVISION.rsplit("-", 1)[-1] + "_")
    destination = target/"checkpoint_reanalizado.json"
    if destination == source or destination.exists():
        raise ValueError("Elija una carpeta nueva; no se sobrescriben checkpoints.")
    updated = deepcopy(payload)
    outcomes = {}
    for item in updated["items"]:
        try:
            data = item.get("result") or {}
            if not isinstance(data, dict) or not str(data.get("texto") or "").strip():
                raise ValueError("La llamada no contiene transcripción recuperable.")
            transcription = TranscriptionModel.from_dict(data)
            transcription.text = TranscriptionService._normalize_asr_money_display(transcription.text)
            previous_review = (item.get("analytics_metadata") or {}).get("hostile_interlocutor_exception")
            if isinstance(previous_review, dict):
                transcription.metadata["hostile_interlocutor_exception"] = deepcopy(previous_review)
            amount_review = (item.get("amount_audio_review") or
                             ((item.get("analytics_metadata") or {}).get("amount_recovery_validation") or {}).get("human_review_record"))
            if isinstance(amount_review, dict):
                transcription.metadata["amount_audio_review"] = deepcopy(amount_review)
            financial_review = (item.get("financial_audio_review") if "financial_audio_review" in item else
                ((item.get("analytics_metadata") or {}).get("financial_review_validation") or {}).get("human_review_record"))
            transcription.metadata["financial_audio_review"] = deepcopy(financial_review)
            relation = item.get("relation") or {}
            analysis = AnalyticsService.analyze_transcription(transcription, relation)
            old = item.get("classification") or {}
            base = {"estatus_autorizado": old.get("estatus_autorizado"),
                    "tiene_nombre_final": bool(relation.get("nombre_final"))}
            decision = AnalyticsService.shared_quality_decision(
                analysis.metadata, analysis.compliance_score, old.get("puntaje_minimo") or 80,
                transcription.duration_seconds, base)
            analysis.metadata["final_result"] = decision["clasificacion"]
            analysis.metadata["shared_classification"] = decision
            item["analytics"] = analysis.to_dict()
            item["analytics_metadata"] = analysis.metadata
            item["classification"] = {**old, **decision}
            item["compliance_score"] = analysis.compliance_score
            item["status"] = "completado"
            item["message"] = decision["motivo"]
            category = decision["clasificacion"]
        except Exception as error:
            category = "ERROR"
            item["status"] = "error"
            item["message"] = str(error)
            item["classification"] = {"clasificacion": "ERROR", "aprobado": False, "motivo": str(error)}
            item["analytics"] = {}
            item["analytics_metadata"] = {"final_result": "ERROR"}
            item["compliance_score"] = None
        item["classified_audio_path"] = None
        item["copy_status"] = "NO_COPIADO_REANALISIS"
        item["transcription_recovered"] = category != "ERROR"
        outcomes[category] = outcomes.get(category, 0)+1
    report = BatchReportService.generate_report(updated["items"], target/"02_Audios"/"Transcripciones",
                                                report_filename="Reporte_Reanalisis_Reglas.xlsx")
    # El encabezado histórico de clasificación física se conserva para compatibilidad.
    # Esta hoja hace explícito que el reanálisis no mueve ni reclasifica archivos de audio.
    from openpyxl import load_workbook
    workbook = load_workbook(report)
    note = workbook.create_sheet("NOTA_REANALISIS")
    note.append(["Reanálisis de reglas sobre transcripciones existentes"])
    note.append(["No se ejecutó Whisper ni se copiaron/movieron audios."])
    note.append(["La columna 'Clasificación física' muestra la clasificación calculada; no acredita una copia nueva."])
    note.append(["El checkpoint y el reporte originales permanecen sin cambios."])
    note.append(["Revisión del motor", AnalyticsService.ENGINE_REVISION])
    workbook.save(report)
    workbook.close()
    updated["completed"] = not outcomes.get("ERROR")
    updated["updated_at"] = datetime.now(timezone.utc).isoformat()
    updated["last_report_path"] = report
    updated["reanalysis"] = {"engine_revision": AnalyticsService.ENGINE_REVISION,
                              "source_sha256": hashlib.sha256(raw).hexdigest(),
                              "whisper_executed": False, "audio_files_copied": False, "outcomes": outcomes}
    destination.write_text(json.dumps(updated, ensure_ascii=False, indent=2), encoding="utf-8")
    assert source.read_bytes() == raw, "El origen cambió durante el reanálisis"
    return {"report": report, "checkpoint": str(destination), "outcomes": outcomes,
            "evaluations": len(updated["items"]), "engine_revision": AnalyticsService.ENGINE_REVISION}
