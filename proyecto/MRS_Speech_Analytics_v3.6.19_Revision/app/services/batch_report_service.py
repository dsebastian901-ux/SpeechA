"""
MRS Speech Analytics Professional
Servicio: Reporte Excel del procesamiento por lotes

Versión con observaciones ejecutivas:
- Fortalezas
- Áreas de oportunidad
- Acción recomendada
- Reglas aplicadas
- Nivel de atención
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, date
from pathlib import Path
import json
import re
from typing import Any, Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class BatchReportService:
    """Genera el reporte operativo posterior a Analytics y clasificación."""

    HEADERS = [
        "Archivo original",
        "Nombre final",
        "Gestor",
        "Supervisor",
        "Estatus Excel",
        "Mora",
        "Fecha llamada",
        "Fecha límite",
        "Vencimiento",
        "Score",
        "Resultado IA",
        "Clasificación física",
        "Nivel de atención",
        "MRS detectado",
        "Estatus detectado IA",
        "Fecha textual escuchada",
        "Fecha de pago detectada",
        "Evidencia de fecha",
        "Regla fecha",
        "Efecto fecha",
        "Monto confirmado",
        "Compromiso confirmado",
        "Liquidación detectada",
        "Atribución liquidación",
        "Conflicto negociación",
        "Penalizaciones",
        "Fallas críticas",
        "Reglas aplicadas",
        "Fortalezas",
        "Áreas de oportunidad",
        "Acción recomendada",
        "Observaciones IA",
        "Motivo de revisión",
        "Resultado Calidad",
        "Validación Dictamen",
        "Motivo Dictamen",
        "Resultado Final Operativo",
        "Validación Negociación",
        "Motivo Negociación",
        "Validación Monto",
        "Motivo Validación Monto",
        "Monto Detectado",
        "Tipo Monto Detectado",
        "Evidencia Monto",
        "Pago Mínimo Esperado",
        "Saldo Total Esperado",
        "Confrontación",
        "Evidencia Confrontación",
        "Interrupciones",
        "Evidencia Interrupciones",
        "Conciliación",
        "Evidencia Conciliación",
        "Finalización Llamada",
        "Abandono/Silencio",
        "Comprobante WhatsApp",
        "Evidencia Comprobante WhatsApp",
        "Lugares Pago Autorizados",
        "Cantidad Lugares Pago",
        "Cumple Mínimo 2 Lugares",
        "Lugar/Método No Autorizado",
        "Alerta Posible Fraude",
        "Evidencia Posible Fraude",
        "Transferencia Detectada",
        "Cuenta Santander Autorizada",
        "Revisión Datos Transferencia",
        "Duración Segundos",
        "Duración Mínima",
        "Duración Máxima",
        "Validación Duración",
        "Transcripción Completa",
        "Transcripción Depurada",
        "Ruta audio clasificado",
        "Interlocutor detectado",
        "Evidencia estatus",
        "Método detección estatus",
        "Recuperación importe audio",
        "Evidencia recuperación importe",
        "Efecto pago mínimo",
        "Política de plazo",
        "Fecha máxima de negociación",
        "Datos de cuenta",
        "Motivo datos de cuenta",
        "Confirmación financiera",
        "Saldo ASR original",
        "Saldo confirmado por usuario",
        "Parcialidades finales confirmadas",
        "Historial de propuestas de parcialidades",
    ]

    MODE_B_HEADERS = [
        "Archivo original",
        "Score",
        "Resultado IA",
        "Clasificación física",
        "Nivel de atención",
        "MRS detectado",
        "Estatus detectado IA",
        "Fecha textual escuchada",
        "Fecha de pago detectada",
        "Evidencia de fecha",
        "Regla fecha",
        "Efecto fecha",
        "Monto confirmado",
        "Compromiso confirmado",
        "Liquidación detectada",
        "Atribución liquidación",
        "Conflicto negociación",
        "Penalizaciones",
        "Fallas críticas",
        "Reglas aplicadas",
        "Fortalezas",
        "Áreas de oportunidad",
        "Acción recomendada",
        "Observaciones IA",
        "Motivo de revisión",
        "Resultado Calidad",
        "Validación Dictamen",
        "Motivo Dictamen",
        "Resultado Final Operativo",
        "Validación Negociación",
        "Motivo Negociación",
        "Validación Monto",
        "Motivo Validación Monto",
        "Monto Detectado",
        "Tipo Monto Detectado",
        "Evidencia Monto",
        "Pago Mínimo Esperado",
        "Saldo Total Esperado",
        "Confrontación",
        "Evidencia Confrontación",
        "Interrupciones",
        "Evidencia Interrupciones",
        "Conciliación",
        "Evidencia Conciliación",
        "Finalización Llamada",
        "Abandono/Silencio",
        "Comprobante WhatsApp",
        "Evidencia Comprobante WhatsApp",
        "Lugares Pago Autorizados",
        "Cantidad Lugares Pago",
        "Cumple Mínimo 2 Lugares",
        "Lugar/Método No Autorizado",
        "Alerta Posible Fraude",
        "Evidencia Posible Fraude",
        "Transferencia Detectada",
        "Cuenta Santander Autorizada",
        "Revisión Datos Transferencia",
        "Duración Segundos",
        "Duración Mínima",
        "Duración Máxima",
        "Validación Duración",
        "Transcripción Completa",
        "Transcripción Depurada",
        "Ruta audio clasificado",
        "Interlocutor detectado",
        "Evidencia estatus",
        "Método detección estatus",
        "Recuperación importe audio",
        "Evidencia recuperación importe",
        "Efecto pago mínimo",
        "Política de plazo",
        "Fecha máxima de negociación",
        "Datos de cuenta",
        "Motivo datos de cuenta",
    ]

    MODE_B_HEADERS += HEADERS[-5:]


    @staticmethod
    def _yes_no(value) -> str:
        if value is None:
            return ""
        return "SI" if bool(value) else "NO"

    @staticmethod
    def _clean_transcription(value: str) -> str:
        """Mejora legibilidad sin inventar ni cambiar montos/fechas."""
        text = str(value or "").strip()
        if not text:
            return ""
        text = re.sub(r"\s+", " ", text)
        # Normalizaciones seguras de marcas recurrentes del dominio.
        text = re.sub(r"\bM\s*[.\-]?\s*R\s*[.\-]?\s*S\b", "MRS", text, flags=re.I)
        text = re.sub(r"\bBradescar(?:t|d)?\b", "Bradescard", text, flags=re.I)
        text = re.sub(r"\s+([,.;:!?])", r"\1", text)
        text = re.sub(r"([,;:])(?=\S)", r"\1 ", text)
        text = re.sub(r"([.!?])\s+", r"\1\n", text)
        # Evita bloques excesivos: conserva contenido y solo mejora lectura.
        return text.strip()

    @staticmethod
    def _fmt(value):
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y")
        if isinstance(value, date):
            return value.strftime("%d/%m/%Y")
        return str(value)

    @staticmethod
    def _unique_text(values: Iterable[Any]) -> List[str]:
        result = []
        for value in values or []:
            text = str(value or "").strip()
            if text and text not in result:
                result.append(text)
        return result

    @classmethod
    def _join_reasons(cls, values: Iterable[Any]) -> str:
        result = []
        for value in values or []:
            if isinstance(value, dict):
                text = str(value.get("reason") or value.get("id") or "").strip()
            else:
                text = str(value or "").strip()
            if text and text not in result:
                result.append(text)
        return " | ".join(result)

    @staticmethod
    def _penalty_codes(penalties) -> List[str]:
        result = []
        for penalty in penalties or []:
            if not isinstance(penalty, dict):
                continue
            code = str(penalty.get("id") or "").strip().upper()
            if code and code not in result:
                result.append(code)
        return result

    @classmethod
    def _rules_applied(cls, metadata: Dict[str, Any]) -> str:
        negotiation = metadata.get("negotiation_validation") or {}
        penalties = metadata.get("penalties") or []
        failures = metadata.get("critical_failures") or []
        manual_rules = metadata.get("manual_rules") or []

        rules = []

        date_rule = str(negotiation.get("date_rule") or "").strip().upper()
        if date_rule:
            rules.append(date_rule)

        rules.extend(cls._penalty_codes(penalties))

        for failure in failures:
            if isinstance(failure, dict):
                code = str(
                    failure.get("id")
                    or failure.get("rule_id")
                    or failure.get("code")
                    or ""
                ).strip().upper()
                if code:
                    rules.append(code)

        if isinstance(manual_rules, dict):
            manual_rules = manual_rules.values()

        for rule in manual_rules or []:
            if isinstance(rule, dict):
                code = str(
                    rule.get("id")
                    or rule.get("rule_id")
                    or rule.get("code")
                    or ""
                ).strip().upper()
            else:
                code = str(rule or "").strip().upper()

            if code and len(code) <= 30:
                rules.append(code)

        return ", ".join(cls._unique_text(rules))

    @classmethod
    def _fortalezas(
        cls,
        classification: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        checks = metadata.get("quality_checks") or {}
        negotiation = metadata.get("negotiation_validation") or {}
        status = metadata.get("status_validation") or {}
        penalties = metadata.get("penalties") or []
        failures = metadata.get("critical_failures") or []

        strengths = []

        if checks.get("identifica_mrs") is True:
            strengths.append("Se detectó identificación como MRS.")

        if status and status.get("valid", True):
            strengths.append("El estatus evaluado no presenta inconsistencias detectadas.")

        if negotiation and negotiation.get("valid", True):
            date_rule = str(negotiation.get("date_rule") or "").strip()
            if date_rule:
                strengths.append(f"La validación de negociación cumple la regla {date_rule}.")
            else:
                strengths.append("La negociación no presenta inconsistencias detectadas.")

        if negotiation.get("amount_confirmed") is True:
            strengths.append("Se detectó monto de pago confirmado.")

        if negotiation.get("commitment_confirmed") is True:
            strengths.append("Se detectó compromiso de pago confirmado.")

        if negotiation.get("promise_date"):
            strengths.append(
                f"Se detectó fecha de pago: {negotiation.get('promise_date')}."
            )

        if classification.get("aprobado") and not penalties and not failures:
            strengths.append("Sin penalizaciones ni fallas críticas detectadas.")

        if not strengths and classification.get("aprobado"):
            strengths.append("La llamada cumple el umbral y las reglas activas de clasificación.")

        return " ".join(cls._unique_text(strengths))

    @classmethod
    def _areas_oportunidad(
        cls,
        classification: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        checks = metadata.get("quality_checks") or {}
        negotiation = metadata.get("negotiation_validation") or {}
        status = metadata.get("status_validation") or {}
        penalties = metadata.get("penalties") or []
        failures = metadata.get("critical_failures") or []

        areas = []

        if checks.get("identifica_mrs") is False:
            areas.append("No se detectó identificación como MRS.")

        areas.extend(status.get("reasons") or [])
        areas.extend(negotiation.get("reasons") or [])

        for penalty in penalties:
            if isinstance(penalty, dict):
                reason = str(penalty.get("reason") or "").strip()
                if reason:
                    areas.append(reason)

        for failure in failures:
            if isinstance(failure, dict):
                reason = str(failure.get("reason") or "").strip()
                if reason:
                    areas.append(reason)

        if negotiation.get("liquidation_regularization_conflict") is True:
            areas.append(
                "Se detectó contradicción entre liquidación y puesta al corriente/regularización."
            )

        motive = str(classification.get("motivo") or "").strip()
        if motive and not classification.get("aprobado"):
            areas.append(motive)

        areas = cls._unique_text(areas)

        if not areas:
            return "Sin áreas de oportunidad detectadas por las reglas activas."

        return " | ".join(areas)

    @classmethod
    def _accion_recomendada(
        cls,
        classification: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        checks = metadata.get("quality_checks") or {}
        negotiation = metadata.get("negotiation_validation") or {}
        status = metadata.get("status_validation") or {}
        failures = metadata.get("critical_failures") or []
        penalties = metadata.get("penalties") or []

        actions = []

        if checks.get("identifica_mrs") is False:
            actions.append("Reforzar la presentación e identificación clara como MRS al inicio de la llamada.")

        if negotiation.get("liquidation_regularization_conflict") is True:
            actions.append(
                "Unificar la explicación del efecto del pago y evitar presentar el mismo pago como liquidación y regularización."
            )

        if negotiation and not negotiation.get("valid", True):
            reasons_text = " ".join(
                str(x or "").lower()
                for x in negotiation.get("reasons", [])
            )
            if "fecha" in reasons_text:
                actions.append("Confirmar expresamente la fecha definitiva del compromiso de pago.")
            if "monto" in reasons_text:
                actions.append("Confirmar expresamente el monto definitivo acordado.")
            if not actions:
                actions.append("Revisar la negociación contra la regla operativa aplicable antes de aprobar.")

        if status and not status.get("valid", True):
            actions.append("Alinear el estatus registrado en Excel con el tipo de contacto detectado en la llamada.")

        if failures:
            actions.append("Escalar a revisión de calidad por existir una falla crítica.")

        if penalties and not actions:
            actions.append("Retroalimentar al gestor sobre las penalizaciones detectadas y reforzar el script aplicable.")

        if classification.get("aprobado") and not actions:
            actions.append("Mantener el estándar observado y replicar las buenas prácticas en llamadas similares.")

        if not classification.get("aprobado") and not actions:
            actions.append("Realizar revisión manual de la llamada antes de liberar el resultado.")

        return " ".join(cls._unique_text(actions))

    @staticmethod
    def _nivel_atencion(
        classification: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        final_result = str(metadata.get("final_result") or "").strip().upper()
        failures = metadata.get("critical_failures") or []
        penalties = metadata.get("penalties") or []

        if final_result in {"INCORRECTA", "FALLA_GRAVE"} or failures:
            return "CRÍTICA"
        if not classification.get("aprobado"):
            return "ALTA"
        if penalties:
            return "MEDIA"
        return "SIN INCIDENCIAS"

    @classmethod
    def _quality_observation(
        cls,
        classification: Dict[str, Any],
        metadata: Dict[str, Any],
    ) -> str:
        strengths = cls._fortalezas(classification, metadata)
        areas = cls._areas_oportunidad(classification, metadata)
        action = cls._accion_recomendada(classification, metadata)

        if classification.get("aprobado"):
            return (
                f"Resultado aprobado. Fortalezas: {strengths} "
                f"Seguimiento: {action}"
            )

        return (
            f"Requiere revisión. Áreas de oportunidad: {areas} "
            f"Acción: {action}"
        )

    @classmethod
    def _row_from_item(cls, item: Dict[str, Any]) -> Dict[str, Any]:
        relation = item.get("relation") or {}
        analytics = item.get("analytics") or {}
        metadata = item.get("analytics_metadata") or analytics.get("metadata") or {}
        classification = item.get("classification") or {}
        checks = metadata.get("quality_checks") or {}
        negotiation = metadata.get("negotiation_validation") or {}
        status = metadata.get("status_validation") or {}
        penalties = metadata.get("penalties") or []
        failures = metadata.get("critical_failures") or []
        treatment = metadata.get("treatment_validation") or {}
        call_end = metadata.get("call_end_validation") or {}
        convenio_payment = metadata.get("convenio_payment_validation") or {}
        card_validation = metadata.get("card_validation") or {}
        duration_validation = metadata.get("duration_validation") or {}
        financial = metadata.get("financial_review_validation") or {}
        reviewed_balance = ((financial.get("human_review_record") or {}).get("balance") or {})

        result_payload = item.get("result") or {}
        transcription = str(
            result_payload.get("texto")
            or item.get("transcription_text")
            or item.get("transcription")
            or ""
        ).strip()

        # v3.5.2: el checkpoint de lotes recuperados guarda la RUTA del JSON,
        # no necesariamente el texto. Leer el JSON individual para que el
        # reporte siempre incluya la transcripción ya existente, sin volver
        # a ejecutar Whisper.
        if not transcription:
            transcription_path = (
                result_payload.get("json_path")
                or result_payload.get("ruta_json")
                or analytics.get("transcription_file")
                or ""
            )
            if transcription_path:
                try:
                    tpath = Path(str(transcription_path))
                    if tpath.exists() and tpath.is_file():
                        payload = json.loads(tpath.read_text(encoding="utf-8-sig"))
                        transcription = str(payload.get("texto") or "").strip()
                        if not transcription:
                            segments = payload.get("segmentos") or []
                            transcription = " ".join(
                                str(seg.get("texto") or "").strip()
                                for seg in segments
                                if isinstance(seg, dict) and str(seg.get("texto") or "").strip()
                            ).strip()
                except Exception:
                    # El reporte no debe fallar completo por un JSON individual.
                    transcription = transcription or ""

        transcription_clean = cls._clean_transcription(transcription)

        penalty_text = []
        for p in penalties:
            if not isinstance(p, dict):
                continue
            pid = p.get("id") or ""
            points = p.get("points")
            reason = p.get("reason") or ""
            label = pid
            if points not in (None, ""):
                try:
                    label += f" (-{abs(int(points))})"
                except (TypeError, ValueError):
                    label += f" (-{points})"
            if reason:
                label += f": {reason}"
            if label:
                penalty_text.append(label)

        result_ia = str(metadata.get("final_result") or "").strip().upper()
        physical = str(classification.get("clasificacion") or "").strip().upper()
        motive = str(classification.get("motivo") or "").strip()

        if item.get("status") == "error":
            result_ia = "ERROR"
            physical = "REVISAR"
            motive = str(item.get("message") or "Error durante el procesamiento.")

        compliance_status = str(
            analytics.get("compliance_status")
            or metadata.get("compliance_status")
            or ("CUMPLE" if classification.get("cumple_puntaje") else "NO CUMPLE")
        ).strip().upper()

        if status and status.get("valid") is False and status.get("review_required"):
            dictamen = "REVISAR"
            motivo_dictamen = cls._join_reasons(status.get("reasons") or [])
        elif status and status.get("valid") is False:
            dictamen = "MAL DICTAMINADO"
            motivo_dictamen = cls._join_reasons(status.get("reasons") or [])
        elif status and status.get("valid") is True:
            dictamen = "CORRECTO"
            motivo_dictamen = ""
        else:
            dictamen = "NO DETERMINADO"
            motivo_dictamen = ""

        invalidating_failure = any(
            isinstance(f, dict) and f.get("invalidates") is True
            for f in failures
        )
        confrontation = str(treatment.get("confrontation") or "").strip()
        confrontation_severe = "CONFRONTACION DEL GESTOR" in confrontation.upper()

        # La clasificación física ya incorpora la prioridad operativa v3.6.1.
        final_operational_map = {
            "APROBADO": "APROBADO",
            "APROBADOS": "APROBADO",
            "APROBADO_SIN_ETIQUETAR": "APROBADO SIN ETIQUETAR",
            "APROBADO_CON_EXCEPCION": "APROBADO CON EXCEPCIÓN",
            "NO_APROBADO_MENOR_80": "NO APROBADO - MENOR A 80%",
            "MAL_DICTAMINADO": "NO APROBADO - MAL DICTAMINADO",
            "NEGOCIACION_INVALIDA": "NO APROBADO - NEGOCIACION INVALIDA",
            "FALLA_GRAVE": "NO APROBADO - FALLA GRAVE",
            "REVISAR": "REVISAR MANUALMENTE",
        }
        final_operational = final_operational_map.get(physical, "REVISAR MANUALMENTE")

        return {
            "Archivo original": item.get("file_name") or Path(str(item.get("path") or "")).name,
            "Nombre final": relation.get("nombre_final") or "",
            "Gestor": relation.get("gestor") or "",
            "Supervisor": relation.get("supervisor") or "",
            "Estatus Excel": relation.get("estatus_grabacion") or "",
            "Mora": relation.get("asignacion") or "",
            "Campaña": relation.get("campana") or "",
            "Corte aplicado": negotiation.get("cut_applied") if negotiation.get("cut_applied") is not None else "",
            "Fecha de pago calculada": negotiation.get("calculated_payment_date") or "",
            "Fecha llamada": cls._fmt(relation.get("fecha")),
            "Corte": relation.get("corte") or "",
            "Vencimiento": cls._fmt(relation.get("vencimiento")),
            "Score": item.get("compliance_score") if item.get("compliance_score") is not None else "",
            "Resultado IA": result_ia,
            "Clasificación física": physical,
            "Nivel de atención": cls._nivel_atencion(classification, metadata),
            "MRS detectado": cls._yes_no(checks.get("identifica_mrs")),
            "Estatus detectado IA": status.get("detected_status") or "",
            "Fecha de pago detectada": negotiation.get("promise_date") or "",
            "Fecha textual escuchada": negotiation.get("promise_date_text") or "",
            "Evidencia de fecha": negotiation.get("promise_date_evidence") or "",
            "Regla fecha": negotiation.get("date_rule") or "",
            "Efecto fecha": negotiation.get("date_effect") or "",
            "Monto confirmado": cls._yes_no(negotiation.get("amount_confirmed")),
            "Pago anterior detectado": cls._yes_no(negotiation.get("payment_already_made")),
            "Monto pago anterior": negotiation.get("prior_payment_total") if negotiation.get("prior_payment_total") is not None else "",
            "Restante requerido": negotiation.get("remaining_minimum_required") if negotiation.get("remaining_minimum_required") is not None else "",
            "Intento de negociación": cls._yes_no(negotiation.get("negotiation_attempt_detected")),
            "Penalización sin intento": cls._yes_no(negotiation.get("missing_collection_attempt_penalty")),
            "Gestión de visita detectada": cls._yes_no(negotiation.get("visit_management_detected")),
            "Lenguaje prohibido visita/bienes": cls._yes_no(negotiation.get("prohibited_visit_wording")),
            "Compromiso confirmado": cls._yes_no(negotiation.get("commitment_confirmed")),
            "Liquidación detectada": cls._yes_no(negotiation.get("liquidation_detected")),
            "Atribución liquidación": str(negotiation.get("liquidation_attribution") or ""),
            "Conflicto negociación": cls._yes_no(negotiation.get("liquidation_regularization_conflict")),
            "Penalizaciones": " | ".join(penalty_text),
            "Fallas críticas": cls._join_reasons(failures),
            "Reglas aplicadas": cls._rules_applied(metadata),
            "Fortalezas": cls._fortalezas(classification, metadata),
            "Áreas de oportunidad": cls._areas_oportunidad(classification, metadata),
            "Acción recomendada": cls._accion_recomendada(classification, metadata),
            "Observaciones IA": cls._quality_observation(classification, metadata),
            "Motivo de revisión": motive if not classification.get("aprobado") else "",
            "Resultado Calidad": compliance_status,
            "Validación Dictamen": dictamen,
            "Motivo Dictamen": motivo_dictamen,
            "Resultado Final Operativo": final_operational,
            "Validación Negociación": (
                "VÁLIDA" if negotiation.get("valid") is True
                else "INVÁLIDA" if negotiation.get("hard_invalid") is True
                else "REVISAR" if negotiation.get("valid") is False or negotiation.get("review_required")
                else "NO DETERMINADA"
            ),
            "Motivo Negociación": cls._join_reasons(negotiation.get("reasons") or []),
            "Validación Monto": str(negotiation.get("validacion_monto") or ""),
            "Motivo Validación Monto": str(negotiation.get("motivo_validacion_monto") or ""),
            "Monto Transcrito": negotiation.get("monto_transcrito") if negotiation.get("monto_transcrito") is not None else "",
            "Monto Detectado": negotiation.get("monto_detectado") if negotiation.get("monto_detectado") is not None else "",
            "Multipago Detectado": cls._yes_no(negotiation.get("multipago_detectado")),
            "Montos Multipago": ", ".join(str(x) for x in (negotiation.get("montos_multipago") or [])),
            "Total Multipago": negotiation.get("total_multipago") if negotiation.get("total_multipago") is not None else "",
            "Tarjeta Excel (últimos 4)": str(card_validation.get("expected_card_last4") or ""),
            "Tarjeta Detectada": str(card_validation.get("detected_card_digits") or ""),
            "Validación Tarjeta": str(card_validation.get("result") or ""),
            "Evidencia Tarjeta": str(card_validation.get("evidence") or ""),
            "Motivo Validación Tarjeta": str(card_validation.get("reason") or ""),
            "Normalizado ASR x1000": "SI" if negotiation.get("monto_normalizado_asr") else "NO",
            "Tipo Monto Detectado": str(negotiation.get("tipo_monto_detectado") or ""),
            "Evidencia Monto": str(negotiation.get("evidencia_monto") or ""),
            "Saldo Informado Detectado": negotiation.get("saldo_informado_detectado") if negotiation.get("saldo_informado_detectado") is not None else "",
            "Evidencia Saldo Informado": str(negotiation.get("evidencia_saldo_informado") or ""),
            "Validación Saldo Informado": str(negotiation.get("validacion_saldo_informado") or ""),
            "Motivo Validación Saldo Informado": str(negotiation.get("motivo_validacion_saldo_informado") or ""),
            "Pago Mínimo Esperado": negotiation.get("pago_minimo_esperado") if negotiation.get("pago_minimo_esperado") is not None else "",
            "Saldo Total Esperado": negotiation.get("saldo_total_esperado") if negotiation.get("saldo_total_esperado") is not None else "",
            "Confrontación": confrontation,
            "Evidencia Confrontación": cls._join_reasons(treatment.get("confrontation_evidence") or []),
            "Interrupciones": str(treatment.get("interruptions") or ""),
            "Evidencia Interrupciones": cls._join_reasons(treatment.get("interruption_evidence") or []),
            "Conciliación": str(treatment.get("conciliation") or ""),
            "Evidencia Conciliación": cls._join_reasons(treatment.get("conciliation_evidence") or []),
            "Finalización Llamada": str(call_end.get("result") or ""),
            "Abandono/Silencio": cls._yes_no(call_end.get("abandonment_review")),
            "Comprobante WhatsApp": cls._yes_no(convenio_payment.get("requests_receipt_whatsapp")) if convenio_payment.get("applies") else "N/A",
            "Evidencia Comprobante WhatsApp": cls._join_reasons(convenio_payment.get("whatsapp_evidence") or []),
            "Lugares Pago Autorizados": ", ".join(convenio_payment.get("authorized_places_detected") or []),
            "Cantidad Lugares Pago": convenio_payment.get("authorized_places_count") if convenio_payment.get("applies") else "",
            "Cumple Mínimo 2 Lugares": cls._yes_no(convenio_payment.get("minimum_places_met")) if convenio_payment.get("applies") else "N/A",
            "Lugar/Método No Autorizado": cls._yes_no(convenio_payment.get("unauthorized_payment_method")) if convenio_payment.get("applies") else "N/A",
            "Alerta Posible Fraude": "GRAVE - POSIBLE FRAUDE" if convenio_payment.get("possible_fraud") else "",
            "Evidencia Posible Fraude": cls._join_reasons(convenio_payment.get("unauthorized_evidence") or []),
            "Transferencia Detectada": cls._yes_no(convenio_payment.get("transfer_detected")) if convenio_payment.get("applies") else "N/A",
            "Cuenta Santander Autorizada": cls._yes_no(convenio_payment.get("authorized_transfer_account_detected")) if convenio_payment.get("transfer_detected") else "N/A",
            "Revisión Datos Transferencia": cls._yes_no(convenio_payment.get("transfer_data_review")) if convenio_payment.get("applies") else "N/A",
            "Duración Segundos": duration_validation.get("duration_seconds", ""),
            "Duración Mínima": duration_validation.get("minimum_seconds", ""),
            "Duración Máxima": (
                duration_validation.get("maximum_seconds")
                if duration_validation.get("maximum_seconds") is not None
                else "SIN LIMITE"
            ),
            "Validación Duración": duration_validation.get("result", ""),
            "Transcripción Completa": transcription,
            "Transcripción Depurada": transcription_clean,
            "Ruta audio clasificado": item.get("classified_audio_path") or "",
            "Interlocutor detectado": metadata.get("quality_contact_type") or "",
            "Evidencia estatus": status.get("evidence") or "",
            "Método detección estatus": status.get("evidence_method") or "",
            "Recuperación importe audio": ("CONFIRMADO POR USUARIO" if (metadata.get("amount_recovery_validation") or {}).get("human_verified") else
                                           "REVISAR" if (metadata.get("amount_recovery_validation") or {}).get("review_required") else ""),
            "Evidencia recuperación importe": negotiation.get("evidencia_monto") if (metadata.get("amount_recovery_validation") or {}).get("entries") else "",
            "Efecto pago mínimo": (negotiation.get("mora_payment_policy") or {}).get("minimum_effect") or "",
            "Política de plazo": (negotiation.get("mora_payment_policy") or {}).get("mode") or "",
            "Fecha máxima de negociación": (negotiation.get("mora_payment_policy") or {}).get("window_end") or "",
            "Datos de cuenta": "NOMBRE DE SALIDA POR CORREGIR" if (metadata.get("account_source_validation") or {}).get("review_required") else "",
            "Motivo datos de cuenta": (metadata.get("account_source_validation") or {}).get("reason") or "",
            "Confirmación financiera": ("CONFIRMADO POR USUARIO" if financial.get("human_verified") else
                "REVISAR CONFIRMACIÓN" if financial.get("review_required") else ""),
            "Saldo ASR original": reviewed_balance.get("original_amount", ""),
            "Saldo confirmado por usuario": reviewed_balance.get("confirmed_amount", ""),
            "Parcialidades finales confirmadas": (" + ".join(f"{v:g}" for v in (negotiation.get("installment_validation") or {}).get("amounts", []))
                if financial.get("installments_applied") else ""),
            "Historial de propuestas de parcialidades": (json.dumps((negotiation.get("installment_validation") or {}).get("evidence", []), ensure_ascii=False)
                if financial.get("installments_applied") else ""),
        }

    @staticmethod
    def _style_sheet(ws, freeze="A2"):
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions

        header_fill = PatternFill("solid", fgColor="E9E2F3")
        approved_fill = PatternFill("solid", fgColor="EAF7EE")
        review_fill = PatternFill("solid", fgColor="FFF4E5")
        critical_fill = PatternFill("solid", fgColor="FDECEC")

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Localiza "Nivel de atención" por encabezado para soportar
            # tanto Modo A como Modo B.
            header_map = {
                str(cell.value or "").strip(): idx
                for idx, cell in enumerate(ws[1])
            }
            attention_idx = header_map.get("Nivel de atención")
            attention = (
                str(row[attention_idx].value or "").strip().upper()
                if attention_idx is not None and attention_idx < len(row)
                else ""
            )
            if attention == "CRÍTICA":
                fill = critical_fill
            elif attention == "ALTA":
                fill = review_fill
            elif attention == "SIN INCIDENCIAS":
                fill = approved_fill
            else:
                fill = None

            if fill:
                for cell in row:
                    cell.fill = fill

        widths = {
            1: 42, 2: 48, 3: 24, 4: 24, 5: 22, 6: 12,
            7: 16, 8: 16, 9: 16, 10: 12, 11: 18, 12: 20,
            13: 18, 14: 16, 15: 24, 16: 20, 17: 14, 18: 24,
            19: 18, 20: 22, 21: 20, 22: 22, 23: 55, 24: 55,
            25: 30, 26: 62, 27: 70, 28: 70, 29: 78, 30: 70, 31: 70,
        }

        for idx, width in widths.items():
            if idx <= ws.max_column:
                ws.column_dimensions[get_column_letter(idx)].width = width

        # v3.5.0: anchos por nombre para campos de auditoría y transcripción.
        audit_widths = {
            "Resultado Calidad": 20, "Validación Dictamen": 22,
            "Motivo Dictamen": 55, "Resultado Final Operativo": 32,
            "Validación Negociación": 24, "Motivo Negociación": 65,
            "Confrontación": 30, "Evidencia Confrontación": 65,
            "Interrupciones": 28, "Evidencia Interrupciones": 65,
            "Finalización Llamada": 38, "Abandono/Silencio": 20,
            "Transcripción Completa": 90,
            "Transcripción Depurada": 90,
            "Atribución liquidación": 22,
            "Comprobante WhatsApp": 22,
            "Evidencia Comprobante WhatsApp": 42,
            "Lugares Pago Autorizados": 55,
            "Cantidad Lugares Pago": 20,
            "Cumple Mínimo 2 Lugares": 24,
            "Lugar/Método No Autorizado": 28,
            "Alerta Posible Fraude": 28,
            "Evidencia Posible Fraude": 55,
            "Transferencia Detectada": 23,
            "Cuenta Santander Autorizada": 28,
            "Revisión Datos Transferencia": 28,
            "Interlocutor detectado": 28,
            "Evidencia estatus": 70,
            "Método detección estatus": 48,
        }
        header_map = {str(c.value or "").strip(): c.column for c in ws[1]}
        for header, width in audit_widths.items():
            col_idx = header_map.get(header)
            if col_idx:
                ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[1].height = 34

    @classmethod
    def generate_report(
        cls,
        items: List[Dict[str, Any]],
        output_directory,
        minimum_score: float = 80.0,
        report_filename: str = None,
        processed_only: bool = False,
    ) -> str:
        if output_directory:
            audio_root = Path(output_directory).expanduser().resolve().parent
            project_root = audio_root.parent
        elif items:
            audio_root = Path(items[0]["path"]).expanduser().resolve().parent
            project_root = audio_root.parent
        else:
            raise ValueError("No hay ubicación disponible para generar el reporte.")

        reports_dir = project_root / "04_Reportes"
        reports_dir.mkdir(parents=True, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if report_filename:
            safe_name = Path(str(report_filename)).name
            if not safe_name.lower().endswith(".xlsx"):
                safe_name += ".xlsx"
            report_path = reports_dir / safe_name
        else:
            report_path = reports_dir / f"Reporte_Analisis_Audios_{timestamp}.xlsx"

        source_items = items
        if processed_only:
            source_items = [
                item for item in items
                if item.get("status") in ("completado", "error")
            ]

        rows = [cls._row_from_item(item) for item in source_items]

        # Modo B se identifica porque el Batch clasifica explícitamente como
        # SOLO_AUDIOS. Si no existe esa marca, se conserva el reporte Modo A.
        mode_b = bool(source_items) and all(
            str((item.get("classification") or {}).get("modo") or "").strip().upper()
            == "SOLO_AUDIOS"
            for item in source_items
            if item.get("classification")
        )
        headers = cls.MODE_B_HEADERS if mode_b else cls.HEADERS

        wb = Workbook()

        # =====================================================
        # RESULTADOS
        # =====================================================
        ws = wb.active
        ws.title = "RESULTADOS"
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        cls._style_sheet(ws)

        # =====================================================
        # CLASIFICACIONES OPERATIVAS v3.6.1
        # =====================================================
        category_sheets = [
            ("APROBADO_SIN_ETIQUETAR", "APROBADO_SIN_ETIQUETAR"),
            ("APROBADO_CON_EXCEPCION", "APROBADO_CON_EXCEPCION"),
            ("NO_APROBADOS_MENOR_80", "NO_APROBADO_MENOR_80"),
            ("MAL_DICTAMINADO", "MAL_DICTAMINADO"),
            ("NEGOCIACION_INVALIDA", "NEGOCIACION_INVALIDA"),
            ("FALLA_GRAVE", "FALLA_GRAVE"),
            ("REVISAR", "REVISAR"),
        ]
        for sheet_name, category in category_sheets:
            cws = wb.create_sheet(sheet_name[:31])
            cws.append(headers)
            for row in rows:
                if str(row.get("Clasificación física") or "").upper() == category:
                    cws.append([row.get(h, "") for h in headers])
            cls._style_sheet(cws)

        # =====================================================
        # AUDITORÍAS ESPECÍFICAS v3.5.0
        # =====================================================
        audit_headers = [
            "Archivo original", "Gestor", "Supervisor", "Estatus Excel", "Mora",
            "Score", "Resultado Calidad", "Validación Dictamen", "Motivo Dictamen",
            "Validación Negociación", "Motivo Negociación", "Fecha textual escuchada",
            "Fecha de pago detectada", "Evidencia de fecha", "Atribución liquidación",
            "Regla fecha", "Efecto fecha", "Confrontación", "Evidencia Confrontación",
            "Interrupciones", "Evidencia Interrupciones", "Finalización Llamada",
            "Abandono/Silencio", "Resultado Final Operativo", "Transcripción Completa",
            "Transcripción Depurada",
        ]

        def add_audit_sheet(title, selected_rows):
            aws = wb.create_sheet(title)
            aws.append(audit_headers)
            for r in selected_rows:
                aws.append([r.get(h, "") for h in audit_headers])
            cls._style_sheet(aws)
            return aws

        add_audit_sheet(
            "NEGOCIACIONES_INVALIDAS",
            [r for r in rows if r.get("Validación Negociación") == "INVÁLIDA" and str(r.get("Estatus Excel") or "").upper() != "CONVENIO"],
        )
        add_audit_sheet(
            "CONVENIOS_INVALIDOS",
            [r for r in rows if r.get("Validación Negociación") == "INVÁLIDA" and str(r.get("Estatus Excel") or "").upper() == "CONVENIO"],
        )
        add_audit_sheet(
            "AUDIT_MAL_DICTAMINADO",
            [r for r in rows if r.get("Validación Dictamen") == "MAL DICTAMINADO"],
        )
        add_audit_sheet(
            "CONFRONTACIONES",
            [r for r in rows if "CONFRONTACION DEL GESTOR" in str(r.get("Confrontación") or "").upper()],
        )

        # =====================================================
        # RESUMEN
        # =====================================================
        summary_ws = wb.create_sheet("RESUMEN")
        total = len(rows)
        approved = sum(
            1 for r in rows
            if r.get("Clasificación física") in {"APROBADO", "APROBADOS"}
        )
        review = sum(1 for r in rows if r.get("Clasificación física") == "REVISAR")
        approved_untagged = sum(1 for r in rows if r.get("Clasificación física") == "APROBADO_SIN_ETIQUETAR")
        under_80 = sum(1 for r in rows if r.get("Clasificación física") == "NO_APROBADO_MENOR_80")
        bad_dictation = sum(1 for r in rows if r.get("Clasificación física") == "MAL_DICTAMINADO")
        invalid_negotiation = sum(1 for r in rows if r.get("Clasificación física") == "NEGOCIACION_INVALIDA")
        severe_failure = sum(1 for r in rows if r.get("Clasificación física") == "FALLA_GRAVE")
        scores = [
            float(r["Score"])
            for r in rows
            if r.get("Score") not in ("", None)
        ]
        incorrect = sum(
            1 for r in rows
            if r.get("Resultado IA") in {"INCORRECTA", "FALLA_GRAVE"}
        )
        errors = sum(
            1 for r in rows
            if r.get("Resultado IA") == "ERROR"
        )
        critical = sum(
            1 for r in rows
            if r.get("Nivel de atención") == "CRÍTICA"
        )
        high_attention = sum(
            1 for r in rows
            if r.get("Nivel de atención") == "ALTA"
        )
        avg_score = round(sum(scores) / len(scores), 2) if scores else 0.0
        approval_pct = round((approved / total) * 100, 2) if total else 0.0

        summary_ws.append(["INDICADOR", "VALOR"])
        summary_ws.append([
            "Modo de procesamiento",
            "Modo B · Solo audios" if mode_b else "Modo A · Audios + Excel",
        ])
        summary_ws.append(["Total audios", total])
        summary_ws.append(["Aprobados", approved])
        summary_ws.append(["Aprobados sin etiquetar", approved_untagged])
        summary_ws.append(["No aprobados menor a 80%", under_80])
        summary_ws.append(["Mal dictaminados", bad_dictation])
        summary_ws.append(["Negociaciones inválidas", invalid_negotiation])
        summary_ws.append(["Fallas graves", severe_failure])
        summary_ws.append(["Revisión manual", review])
        summary_ws.append(["Incorrectos IA", incorrect])
        summary_ws.append(["Errores", errors])
        summary_ws.append(["Atención crítica", critical])
        summary_ws.append(["Atención alta", high_attention])
        summary_ws.append(["Score promedio", avg_score])
        summary_ws.append(["% aprobación", approval_pct])
        summary_ws.append(["Umbral mínimo", minimum_score])
        summary_ws.append(["Mal dictaminado", sum(1 for r in rows if r.get("Validación Dictamen") == "MAL DICTAMINADO")])
        summary_ws.append(["Negociaciones inválidas", sum(1 for r in rows if r.get("Validación Negociación") == "INVÁLIDA")])
        summary_ws.append(["Confrontaciones del gestor", sum(1 for r in rows if "CONFRONTACION DEL GESTOR" in str(r.get("Confrontación") or "").upper())])

        summary_ws.column_dimensions["A"].width = 30
        summary_ws.column_dimensions["B"].width = 18
        cls._style_sheet(summary_ws)

        # =====================================================
        # RESUMEN ESPECÍFICO SEGÚN MODO
        # =====================================================
        if mode_b:
            archivos_ws = wb.create_sheet("RESUMEN_ARCHIVOS")
            archivos_ws.append([
                "Archivo",
                "Score",
                "Resultado IA",
                "Clasificación",
                "Nivel de atención",
                "Principales incidencias",
                "Acción recomendada",
            ])

            for row in rows:
                incidencias = str(row.get("Áreas de oportunidad") or "").strip()
                if incidencias.startswith("Sin áreas"):
                    incidencias = ""

                archivos_ws.append([
                    row.get("Archivo original") or "",
                    row.get("Score") or "",
                    row.get("Resultado IA") or "",
                    row.get("Clasificación física") or "",
                    row.get("Nivel de atención") or "",
                    incidencias,
                    row.get("Acción recomendada") or "",
                ])

            archivos_ws.column_dimensions["A"].width = 44
            archivos_ws.column_dimensions["B"].width = 12
            archivos_ws.column_dimensions["C"].width = 18
            archivos_ws.column_dimensions["D"].width = 20
            archivos_ws.column_dimensions["E"].width = 18
            archivos_ws.column_dimensions["F"].width = 70
            archivos_ws.column_dimensions["G"].width = 70
            cls._style_sheet(archivos_ws)
        else:
            # =====================================================
            # RESUMEN POR GESTOR
            # =====================================================
            gestores_ws = wb.create_sheet("RESUMEN_GESTORES")
            gestores_ws.append([
                "Gestor",
                "Total",
                "Aprobados",
                "Revisar",
                "Incorrectas IA",
                "Atención crítica",
                "Score promedio",
                "% aprobación",
                "Principales penalizaciones",
                "Principales áreas de oportunidad",
            ])

            grouped = defaultdict(list)
            for row in rows:
                grouped[row.get("Gestor") or "SIN GESTOR"].append(row)

            for gestor, group in sorted(grouped.items()):
                gtotal = len(group)
                gapproved = sum(
                    1 for r in group
                    if r.get("Clasificación física") in {"APROBADO", "APROBADOS"}
                )
                greview = gtotal - gapproved
                gincorrect = sum(
                    1 for r in group
                    if r.get("Resultado IA") in {"INCORRECTA", "FALLA_GRAVE"}
                )
                gcritical = sum(
                    1 for r in group
                    if r.get("Nivel de atención") == "CRÍTICA"
                )
                gscores = [
                    float(r["Score"])
                    for r in group
                    if r.get("Score") not in ("", None)
                ]
                gavg = round(sum(gscores) / len(gscores), 2) if gscores else 0.0
                gpct = round((gapproved / gtotal) * 100, 2) if gtotal else 0.0

                codes = []
                areas = []

                for r in group:
                    for part in str(r.get("Penalizaciones") or "").split(" | "):
                        code = part.split(" ", 1)[0].strip()
                        if code.startswith("PEN"):
                            codes.append(code)

                    area = str(r.get("Áreas de oportunidad") or "").strip()
                    if area and not area.startswith("Sin áreas"):
                        for part in area.split(" | "):
                            part = part.strip()
                            if part:
                                areas.append(part)

                common_codes = ", ".join(
                    f"{k} ({v})"
                    for k, v in Counter(codes).most_common(5)
                )
                common_areas = " | ".join(
                    f"{k} ({v})"
                    for k, v in Counter(areas).most_common(3)
                )

                gestores_ws.append([
                    gestor,
                    gtotal,
                    gapproved,
                    greview,
                    gincorrect,
                    gcritical,
                    gavg,
                    gpct,
                    common_codes,
                    common_areas,
                ])

            gestor_widths = {
                "A": 28, "B": 12, "C": 12, "D": 12, "E": 16,
                "F": 16, "G": 16, "H": 16, "I": 42, "J": 70,
            }
            for col, width in gestor_widths.items():
                gestores_ws.column_dimensions[col].width = width
            cls._style_sheet(gestores_ws)

        if report_filename:
            temp_path = report_path.with_name(report_path.stem + "_nuevo.xlsx")
            wb.save(temp_path)
            temp_path.replace(report_path)
        else:
            wb.save(report_path)
        return str(report_path)

    @classmethod
    def generate_partial_report(
        cls,
        items: List[Dict[str, Any]],
        output_directory,
        minimum_score: float = 80.0,
    ) -> str:
        """Genera/actualiza un único Excel con los audios ya procesados."""
        processed = [
            item for item in items
            if item.get("status") in ("completado", "error")
        ]
        if not processed:
            raise ValueError("Todavía no hay audios procesados para generar el reporte parcial.")

        return cls.generate_report(
            items=processed,
            output_directory=output_directory,
            minimum_score=minimum_score,
            report_filename="Reporte_Parcial_Lote.xlsx",
            processed_only=True,
        )
