"""Confirmación humana con alcance limitado y cotejo independiente de cuenta."""
from copy import deepcopy
import json
from pathlib import Path

import pytest
from openpyxl import Workbook,load_workbook

from app.models.transcriptions import TranscriptionModel
from app.services.analytics_service import AnalyticsService as A
from app.services.asr_amount_recovery import recover_amounts
from app.services.amount_review_service import AmountReviewService,make_amount_review,amount_review_status
from app.services.account_source_validation import validate_account_source,corrected_account_name,write_corrected_relation_excel
from app.services.checkpoint_human_review import review_checkpoint
from app.services.checkpoint_reanalysis import reanalyze_checkpoint
from app.services.transcription_services import TranscriptionService
from app.services.review_service import ReviewService

CARD="4800000000001234"
OTHER="4800000000005678"
TEXT=("Habla MRS. Soy yo. Le ofrecemos un descuento para finiquitar la deuda, le están pidiendo la cuenta. "
      "Me interesa, pero yo creo a principios de noviembre.")
SEGMENTS=[{"inicio":20,"fin":28,"texto":TEXT}]
RELATION={"asignacion":"C4","fecha":"2026-10-14","vencimiento":"2026-10-31",
          "pago_minimo":"2000","saldo_total":"50000","estatus_grabacion":"NO DEFINE FECHA DE PAGO",
          "numero_de_tarjeta":CARD,"nombre_final":"000"+CARD+"_C4_synthetic"}


def data():
    readings=iter(["Le están pidiendo dos mil cuatrocientos cincuenta pesos para liquidar la cuenta.",
                   "Únicamente 2450 pesos para liquidar la cuenta."])
    return {"archivo":"synthetic.wav","texto":TEXT,"segmentos":deepcopy(SEGMENTS),"duracion_segundos":90,
            "recuperacion_importes":recover_amounts(TEXT,SEGMENTS,lambda *args:next(readings))}


def review(raw):
    anchors=[e["anchor_start"] for e in raw["recuperacion_importes"]["entries"]]
    return make_amount_review(raw,anchors,audio_verified=True)


def analyze(raw,record=None,relation=None):
    t=TranscriptionModel.from_dict(raw)
    if record is not None:t.metadata["amount_audio_review"]=record
    return A.analyze_transcription(t,RELATION if relation is None else relation)


def test_unreviewed_and_reviewed_amount_produce_distinct_decisions():
    raw=data();original=deepcopy(raw)
    assert analyze(raw).metadata["shared_classification"]["clasificacion"]=="REVISAR"
    a=analyze(raw,review(raw))
    assert a.metadata["amount_recovery_validation"]["human_verified"] is True
    assert a.metadata["shared_classification"]["clasificacion"]=="APROBADO"
    assert raw==original
    assert raw["recuperacion_importes"]["human_verified"] is False


def test_explicit_audio_confirmation_is_required():
    raw=data();anchor=raw["recuperacion_importes"]["entries"][0]["anchor_start"]
    with pytest.raises(ValueError,match="explícita"):
        make_amount_review(raw,[anchor],audio_verified=False)


@pytest.mark.parametrize("change",["text","segments","amount","other_call","scope","flag","source_hash"])
def test_changed_evidence_or_invalid_claim_does_not_clear_review(change):
    raw=data();record=review(raw)
    if change=="text":raw["texto"]+=" Cambio."
    if change=="segments":raw["segmentos"][0]["fin"]+=1
    if change=="amount":record["confirmed_entries"][0]["amount"]=2500
    if change=="other_call":raw["archivo"]="other.wav"
    if change=="scope":record["scope"]="ALL_APPROVED"
    if change=="flag":record["audio_verified"]="true"
    if change=="source_hash":record["source_sha256"]="invalid"
    assert amount_review_status(raw,record)["human_verified"] is False


def test_unknown_anchor_does_not_create_a_confirmation():
    with pytest.raises(ValueError,match="evidencia"):
        make_amount_review(data(),[999999],audio_verified=True)


def test_recovery_error_cannot_be_confirmed_as_a_success():
    raw=data();record=review(raw)
    raw["recuperacion_importes"]["entries"].append({"status":"REQUIERE_REVISION","amount":None})
    assert not amount_review_status(raw,record)["human_verified"]


def test_name_mismatch_is_not_an_audio_card_mismatch():
    raw=data();relation=dict(RELATION,nombre_final="000"+OTHER+"_C4_synthetic")
    a=analyze(raw,review(raw),relation)
    assert a.metadata["account_source_validation"]["review_required"]
    assert a.metadata["shared_classification"]["clasificacion"]=="REVISAR"
    assert "PEN003" not in {p["id"] for p in a.metadata["penalties"]}
    assert "1234" in a.metadata["account_source_validation"]["reason"]
    assert a.metadata["card_validation"]["expected_card_last4"]=="1234"


def test_actual_audio_reference_is_compared_only_with_excel_card():
    raw=data();raw["texto"]+=" Referencia de pago "+OTHER+"."
    # Una revisión del importe de otra versión del texto es inválida también.
    a=analyze(raw,None,dict(RELATION,nombre_final="000"+OTHER+"_C4_synthetic"))
    assert a.metadata["card_validation"]["result"].startswith("REVISAR")
    assert not a.metadata["shared_classification"]["aprobado"]


def test_scoped_amount_confirmation_does_not_waive_critical_failures():
    raw=data();raw["texto"]="Familiar: Soy su madre.\nGestor: Tiene una deuda de 50000 pesos."
    a=analyze(raw,review(data()),dict(RELATION,estatus_grabacion="MENSAJE CON FAMILIAR"))
    assert a.metadata["shared_classification"]["clasificacion"]=="FALLA_GRAVE"


def test_sidecar_persists_and_can_be_revoked_without_touching_transcript(tmp_path):
    raw=data();p=tmp_path/"synthetic.json";p.write_text(json.dumps(raw));before=p.read_bytes()
    anchor=raw["recuperacion_importes"]["entries"][0]["anchor_start"]
    AmountReviewService.save(p,[anchor],audio_verified=True)
    ReviewService.save(p,dict.fromkeys(ReviewService.FLAGS,True))
    t=TranscriptionService.load_from_json(p)
    assert t.metadata["amount_audio_review"]["audio_verified"] is True
    assert t.metadata["hostile_interlocutor_exception"]["audio_verified"] is True
    assert p.read_bytes()==before
    AmountReviewService.save(p,[],audio_verified=False)
    assert AmountReviewService.load(p) is None
    assert ReviewService.load(p)["audio_verified"] is True


def checkpoint(tmp_path):
    raw=data();relation=dict(RELATION,nombre_final="000"+OTHER+"_C4_synthetic")
    p=tmp_path/"batch.json";p.write_text(json.dumps({"items":[{"file_name":"synthetic.wav","result":raw,"relation":relation}]}))
    return p,raw,relation


def test_checkpoint_confirmation_and_name_correction_persist_through_reanalysis(tmp_path):
    p,raw,relation=checkpoint(tmp_path);original=p.read_bytes()
    entry=raw["recuperacion_importes"]["entries"][0]
    args={"synthetic.wav":{"anchors":[entry["anchor_start"]],"audio_verified":True}}
    result=review_checkpoint(p,args,["synthetic.wav"],tmp_path/"new")
    assert result["outcomes"]=={"APROBADO":1}
    saved=json.loads(Path(result["checkpoint"]).read_text());item=saved["items"][0]
    assert item["result"]==raw
    assert item["relation"]["numero_de_tarjeta"]==CARD
    assert item["relation"]["nombre_final"]=="000"+CARD+"_C4_synthetic"
    assert item["output_name_review"]["previous_name"]==relation["nombre_final"]
    assert p.read_bytes()==original
    again=reanalyze_checkpoint(result["checkpoint"],tmp_path/"again")
    assert again["outcomes"]=={"APROBADO":1}
    book=load_workbook(result["report"],read_only=True,data_only=True)
    rows=list(book["RESULTADOS"].values);head=rows[0]
    assert rows[1][head.index("Recuperación importe audio")]=="CONFIRMADO POR USUARIO"
    book.close()


def test_amount_only_confirmation_leaves_the_name_block(tmp_path):
    p,raw,_=checkpoint(tmp_path)
    result=review_checkpoint(p,{"synthetic.wav":{"anchors":[raw["recuperacion_importes"]["entries"][0]["anchor_start"]],"audio_verified":True}},output_directory=tmp_path/"new")
    assert result["outcomes"]=={"REVISAR":1}


def test_confirming_one_call_cannot_affect_another_checkpoint(tmp_path):
    p,_,_=checkpoint(tmp_path)
    with pytest.raises(ValueError,match="ajena"):
        review_checkpoint(p,{"other.wav":{}},output_directory=tmp_path/"new")
    assert not (tmp_path/"new").exists()


def test_excel_copy_corrects_only_output_name_and_preserves_other_data(tmp_path):
    source=tmp_path/"original.xlsx";book=Workbook();sheet=book.active;sheet.title="Relacion_Audios"
    sheet.append(["archivo_original","numero_de_tarjeta","nombre_final","saldo_total"])
    sheet.append(["synthetic.wav",CARD,"000"+OTHER+"_C4_synthetic",50000]);sheet.append(["other.wav",OTHER,"keep",123])
    book.save(source);before=source.read_bytes()
    updates={"synthetic.wav":{"numero_de_tarjeta":CARD,"nombre_final":"000"+CARD+"_C4_synthetic"}}
    target=tmp_path/"corrected.xlsx";write_corrected_relation_excel(source,updates,target)
    copied=load_workbook(target);rows=list(copied.active.values)
    assert rows[1][1]==CARD and rows[1][2]==updates["synthetic.wav"]["nombre_final"] and rows[1][3]==50000
    assert rows[2]==("other.wav",OTHER,"keep",123)
    copied.close();assert source.read_bytes()==before


@pytest.mark.parametrize("relation",[dict(RELATION,nombre_final="result.wav"),dict(RELATION,nombre_final="000"+CARD+"_C4")])
def test_non_card_output_names_and_matching_names_do_not_add_review(relation):
    assert not validate_account_source(relation)["review_required"]


def test_correction_cannot_guess_an_invalid_card():
    with pytest.raises(ValueError,match="16 dígitos"):
        corrected_account_name(dict(RELATION,numero_de_tarjeta="1234"))


def test_checkpoint_confirmation_can_be_revoked(tmp_path):
    p,raw,_=checkpoint(tmp_path);anchor=raw["recuperacion_importes"]["entries"][0]["anchor_start"]
    first=review_checkpoint(p,{"synthetic.wav":{"anchors":[anchor],"audio_verified":True}},["synthetic.wav"],tmp_path/"first")
    assert first["outcomes"]=={"APROBADO":1}
    second=review_checkpoint(first["checkpoint"],{"synthetic.wav":{"anchors":[],"audio_verified":False}},output_directory=tmp_path/"second")
    assert second["outcomes"]=={"REVISAR":1}


def test_confirmation_does_not_lower_saved_score_threshold(tmp_path):
    p,raw,_=checkpoint(tmp_path);payload=json.loads(p.read_text())
    raw["texto"]=TEXT.replace("Habla MRS.","Buenos días.")
    raw["segmentos"][0]["texto"]=raw["texto"]
    raw["recuperacion_importes"]=recover_amounts(raw["texto"],raw["segmentos"],lambda *args:"2450 pesos para liquidar la cuenta.")
    payload["items"][0]["result"]=raw
    payload["items"][0]["classification"]={"puntaje_minimo":90}
    p.write_text(json.dumps(payload))
    result=review_checkpoint(p,{"synthetic.wav":{"anchors":[raw["recuperacion_importes"]["entries"][0]["anchor_start"]],"audio_verified":True}},["synthetic.wav"],tmp_path/"new")
    assert "APROBADO" not in result["outcomes"]
    item=json.loads(Path(result["checkpoint"]).read_text())["items"][0]
    assert item["classification"]["score"]==80 and item["classification"]["puntaje_minimo"]==90


def test_excel_card_change_cannot_be_silently_accepted(tmp_path):
    source=tmp_path/"source.xlsx";book=Workbook();sheet=book.active;sheet.title="Relacion_Audios"
    sheet.append(["archivo_original","numero_de_tarjeta","nombre_final"])
    sheet.append(["synthetic.wav",OTHER,"old"]);book.save(source)
    target=tmp_path/"new.xlsx"
    with pytest.raises(ValueError,match="cambió"):
        write_corrected_relation_excel(source,{"synthetic.wav":{"numero_de_tarjeta":CARD,"nombre_final":"corrected"}},target)
    assert not target.exists()


def test_copy_operation_does_not_overwrite_original_excel(tmp_path):
    source=tmp_path/"source.xlsx";source.write_bytes(b"untouched")
    with pytest.raises(ValueError,match="sobrescribe"):
        write_corrected_relation_excel(source,{},source)
    assert source.read_bytes()==b"untouched"
