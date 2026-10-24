import json
import wave
from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook

from app.services.batch_service import BatchService
from app.services.whisper_service import whisper_service


@pytest.mark.parametrize("text,expected", [
    ("Habla MRS. Soy su madre. Dígale que se comunique al número de contacto. Le paso el recado.", "APROBADO"),
    ("Habla MRS. Soy su madre. Tiene una deuda de 1500 pesos.", "REVISAR"),
    ("Familiar: Soy su madre.\nGestor: Tiene una deuda de 1500 pesos.", "FALLA_GRAVE"),
    ("Habla MRS. Referencia de pago 1234567890123456.", "REVISAR"),
    ("Habla MRS. Gracias por su tiempo. Mierda.", "REVISAR"),
])
def test_reanalysis_reuses_json_and_preserves_report_and_folders(tmp_path, monkeypatch, text, expected):
    source = tmp_path / "input" / "synthetic.wav"
    source.parent.mkdir()
    with wave.open(str(source), "wb") as audio:
        audio.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        audio.writeframes(b"\0\0" * 160)
    output = tmp_path / "project" / "02_Audios" / "Transcripciones"
    output.mkdir(parents=True)
    payload = {"archivo": source.name, "ruta_audio": str(source), "texto": text,
               "duracion_segundos": 90, "segmentos": []}
    transcript = output / "synthetic.json"
    transcript.write_text(json.dumps(payload), encoding="utf-8")

    def must_not_transcribe(*args, **kwargs):
        pytest.fail("Una transcripción existente no debe volver a pasar por Whisper")
    monkeypatch.setattr(whisper_service, "transcribe_audio", must_not_transcribe)
    batch = BatchService()
    batch.set_files([str(source)])
    results = batch.process_all(output_directory=str(output))
    assert len(results) == 1, batch.items
    item = batch.items[0]
    assert item["status"] == "completado", item.get("message")
    assert item["transcription_recovered"] is True
    assert item["classification"]["clasificacion"] == expected
    assert Path(item["classified_audio_path"]).exists()
    assert Path(item["classified_audio_path"]).parent.name == ("APROBADOS" if expected == "APROBADO" else expected)
    assert source.read_bytes() == Path(item["classified_audio_path"]).read_bytes()
    assert json.loads(transcript.read_text(encoding="utf-8")) == payload
    checkpoint = json.loads(Path(batch.checkpoint_path).read_text(encoding="utf-8"))
    assert checkpoint["completed"] is True
    assert checkpoint["items"][0]["classification"]["clasificacion"] == expected
    assert batch.last_report_path, item.get("report_error")
    book = load_workbook(batch.last_report_path, read_only=True, data_only=True)
    try:
        found = False
        for sheet in book:
            rows = list(sheet.values)
            for index, row in enumerate(rows):
                if "Clasificación física" in row:
                    column = row.index("Clasificación física")
                    found = any(r[column] == expected for r in rows[index+1:] if len(r)>column)
                    if found:
                        break
            if found:
                break
        assert found, "El Excel debe contener el dictamen físico correcto"
    finally:
        book.close()
    # Un lote ya terminado no reprocesa sus llamadas.
    assert batch.process_all(output_directory=str(output)) == []


def test_excel_recado_with_uncertain_disclosure_cannot_be_approved(tmp_path, monkeypatch):
    source = tmp_path / "synthetic.wav"
    source.write_bytes(b"synthetic copy fixture")
    output = tmp_path / "project" / "02_Audios" / "Transcripciones"
    output.mkdir(parents=True)
    text = ("Habla MRS. Soy su madre. Dígale que se comunique al número de contacto. "
            "Le paso el recado. Tiene una deuda de 1500 pesos.")
    (output/"synthetic.json").write_text(json.dumps({"archivo": source.name, "texto": text,
        "duracion_segundos": 90, "segmentos": []}), encoding="utf-8")
    excel = tmp_path/"relation.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Relacion_Audios"
    sheet.append(["archivo_original", "gestor", "estatus_grabacion", "nombre_final"])
    sheet.append([source.name, "Gestor sintético", "RECADO REF1", "resultado.wav"])
    book.save(excel)
    batch = BatchService()
    batch.set_files([str(source)])
    batch.set_relation_excel(str(excel))
    monkeypatch.setattr(whisper_service, "transcribe_audio", lambda **kw: pytest.fail("No retranscribir"))
    assert len(batch.process_all(output_directory=str(output))) == 1
    item = batch.items[0]
    m = item["analytics_metadata"]
    assert m["status_validation"]["matches"] is True
    assert m["privacy_validation"]["review_required"] is True
    assert item["compliance_score"] >= 80
    assert item["classification"]["clasificacion"] == "REVISAR"
    assert Path(item["classified_audio_path"]).parent.name == "REVISAR"
