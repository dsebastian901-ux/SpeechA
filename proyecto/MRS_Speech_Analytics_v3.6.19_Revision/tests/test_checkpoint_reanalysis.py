import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app.services.checkpoint_reanalysis import reanalyze_checkpoint


def test_reanalysis_generates_excel_without_original_audio_or_whisper(tmp_path):
    source = tmp_path/"batch_checkpoint.json"
    payload = {"items": [{"path": "C:/missing/synthetic.wav", "file_name": "synthetic.wav",
        "result": {"archivo": "synthetic.wav", "duracion_segundos": 90,
                   "texto": "Habla MRS. Soy su madre. Dígale que se comunique. Le paso el recado."},
        "classification": {}, "status": "completado"}]}
    source.write_text(json.dumps(payload), encoding="utf-8")
    original = source.read_bytes()
    result = reanalyze_checkpoint(source, tmp_path/"new-results")
    assert result["evaluations"] == 1 and result["outcomes"] == {"APROBADO": 1}
    assert source.read_bytes() == original
    saved = json.loads(Path(result["checkpoint"]).read_text(encoding="utf-8"))
    assert saved["reanalysis"]["whisper_executed"] is False
    assert saved["reanalysis"]["engine_revision"] == result["engine_revision"] == "3.6.19-r17"
    assert saved["items"][0]["analytics_metadata"]["engine_revision"] == result["engine_revision"]
    assert saved["items"][0]["classified_audio_path"] is None
    book = load_workbook(result["report"], read_only=True, data_only=True)
    try:
        assert "NOTA_REANALISIS" in book.sheetnames
        assert book["NOTA_REANALISIS"]["B5"].value == result["engine_revision"]
        rows = list(book["RESULTADOS"].values)
        index = rows[0].index("Resultado IA")
        assert rows[1][index] == "APROBADO"
    finally:
        book.close()


def test_empty_checkpoint_is_not_a_successful_reanalysis(tmp_path):
    source = tmp_path/"empty.json"
    source.write_text('{"items": []}')
    with pytest.raises(ValueError, match="no contiene llamadas"):
        reanalyze_checkpoint(source, tmp_path/"output")


@pytest.mark.parametrize("complete,expected", [(True, "APROBADO_CON_EXCEPCION"), (False, "REVISAR")])
def test_existing_human_treatment_review_is_preserved(tmp_path, complete, expected):
    source = tmp_path/"checkpoint.json"
    flags = {"audio_verified": True, "interlocutor_hostile_or_lascivious": complete,
             "prevented_script": complete, "agent_professional": complete}
    source.write_text(json.dumps({"items": [{"path": "missing.wav", "file_name": "missing.wav",
        "result": {"archivo": "missing.wav", "duracion_segundos": 90,
                   "texto": "Habla MRS. Soy su madre. Dígale que se comunique. Le paso el recado. Mierda."},
        "analytics_metadata": {"hostile_interlocutor_exception": flags}}]}), encoding="utf-8")
    result = reanalyze_checkpoint(source, tmp_path/"output")
    assert result["outcomes"] == {expected: 1}
