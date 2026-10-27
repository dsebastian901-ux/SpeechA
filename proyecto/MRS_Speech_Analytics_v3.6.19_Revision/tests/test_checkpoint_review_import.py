"""Incorporación de revisiones sin perder llamadas ni reutilizar evidencia incompatible."""
from copy import deepcopy
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.checkpoint_review_import import import_checkpoint_reviews
from app.services.checkpoint_reanalysis import reanalyze_checkpoint
from test_financial_review import item, make_review
from test_human_amount_review import data, review, RELATION, CARD, OTHER


def batches():
    master = [item(), item("plan")]
    donors = deepcopy(master)
    for i in donors:
        i["financial_audio_review"] = make_review(i)
        i["classification"] = {"clasificacion":"APROBADO", "puntaje_minimo":1}
    amount = {"file_name":"synthetic.wav", "result":data(),
        "relation":dict(RELATION, nombre_final="000"+OTHER+"_C4_synthetic")}
    validated = deepcopy(amount)
    validated["amount_audio_review"] = review(validated["result"])
    validated["relation"]["nombre_final"] = "000"+CARD+"_C4_synthetic"
    validated["output_name_review"] = {"source":"EXPLICIT_USER_CHOICE_NUMERO_DE_TARJETA",
        "previous_name":amount["relation"]["nombre_final"], "corrected_name":validated["relation"]["nombre_final"]}
    amount["result"].pop("recuperacion_importes")
    amount["result"]["ruta_audio"] = "C:/original/synthetic.wav"
    validated["result"]["ruta_audio"] = "D:/new/synthetic.wav"
    master.append(amount); donors.append(validated)
    master.append({"file_name":"untouched.wav", "result":{"archivo":"untouched.wav", "duracion_segundos":90,
        "texto":"Habla MRS. Soy su madre. Dígale que se comunique. Le paso el recado."}})
    return {"items":master}, {"items":donors}


def write(tmp_path, master, donors):
    batch, reviewed = tmp_path/"full.json", tmp_path/"reviewed.json"
    batch.write_text(json.dumps(master), encoding="utf-8")
    reviewed.write_text(json.dumps(donors), encoding="utf-8")
    return batch, reviewed


def test_import_keeps_all_calls_original_transcripts_runtime_paths_and_validations(tmp_path):
    master, donors = batches(); batch, reviewed = write(tmp_path, master, donors)
    before = batch.read_bytes(), reviewed.read_bytes()
    out = import_checkpoint_reviews(batch, reviewed, tmp_path/"merged")
    assert out["evaluations"] == 4 and out["imported_calls"] == 3
    assert out["outcomes"] == {"APROBADO":4}
    merged = json.loads(Path(out["checkpoint"]).read_text())
    assert [i["file_name"] for i in merged["items"]] == [i["file_name"] for i in master["items"]]
    for old, new in zip(master["items"], merged["items"]):
        for key, value in old["result"].items():
            assert new["result"][key] == value
    assert merged["items"][2]["result"]["ruta_audio"] == "C:/original/synthetic.wav"
    assert merged["items"][2]["result"]["recuperacion_importes"] == donors["items"][2]["result"]["recuperacion_importes"]
    assert merged["items"][2]["relation"]["numero_de_tarjeta"] == CARD
    assert merged["items"][2]["relation"]["nombre_final"] == "000"+CARD+"_C4_synthetic"
    assert "review_import" not in merged["items"][3]
    assert merged["review_import"]["imported_calls"] == 3
    assert merged["review_import"]["whisper_executed"] is False
    again = reanalyze_checkpoint(out["checkpoint"], tmp_path/"again")
    assert again["outcomes"] == out["outcomes"]
    repeat = import_checkpoint_reviews(out["checkpoint"], reviewed, tmp_path/"repeated-import")
    assert repeat["outcomes"] == out["outcomes"]
    assert (batch.read_bytes(), reviewed.read_bytes()) == before
    book = load_workbook(out["report"], read_only=True, data_only=True)
    rows = list(book["RESULTADOS"].values)
    assert len(rows) == 5
    assert all(r[rows[0].index("Resultado IA")] == "APROBADO" for r in rows[1:])
    book.close()


@pytest.mark.parametrize("problem", ["unknown_call", "internal_name", "text", "segments", "duration", "card",
    "minimum", "status", "date", "name_without_review", "wrong_name_review", "recovery", "amount_record", "financial_record"])
def test_incompatible_evidence_or_context_rejects_whole_import_before_writing(tmp_path, problem):
    master, donors = batches(); target = donors["items"][1]
    if problem == "unknown_call": target["file_name"] = "other.wav"
    if problem == "internal_name": target["result"]["archivo"] = "other.wav"
    if problem == "text": target["result"]["texto"] += " Cambio."
    if problem == "segments": target["result"]["segmentos"][0]["fin"] += 1
    if problem == "duration": target["result"]["duracion_segundos"] += 1
    if problem == "card": target["relation"]["numero_de_tarjeta"] = OTHER
    if problem == "minimum": target["relation"]["pago_minimo"] = 2600
    if problem == "status": target["relation"]["estatus_grabacion"] = "OTRO"
    if problem == "date": target["relation"]["fecha"] = "2026-10-15"
    if problem == "name_without_review": target["relation"]["nombre_final"] = "changed"
    if problem == "wrong_name_review": donors["items"][2]["output_name_review"]["previous_name"] = "other"
    if problem == "recovery": master["items"][2]["result"]["recuperacion_importes"] = {"entries":[]}
    if problem == "amount_record": donors["items"][2]["amount_audio_review"]["source_sha256"] = "invalid"
    if problem == "financial_record": target["financial_audio_review"]["source_sha256"] = "invalid"
    batch, reviewed = write(tmp_path, master, donors); before = batch.read_bytes(), reviewed.read_bytes()
    with pytest.raises(ValueError):
        import_checkpoint_reviews(batch, reviewed, tmp_path/"new")
    assert not (tmp_path/"new").exists()
    assert (batch.read_bytes(), reviewed.read_bytes()) == before


@pytest.mark.parametrize("which", ["batch", "reviewed"])
def test_duplicate_names_are_never_matched_by_row_position(tmp_path, which):
    master, donors = batches()
    data = master if which == "batch" else donors
    data["items"].append(deepcopy(data["items"][0]))
    batch, reviewed = write(tmp_path, master, donors)
    with pytest.raises(ValueError, match="duplicados"):
        import_checkpoint_reviews(batch, reviewed, tmp_path/"new")
    assert not (tmp_path/"new").exists()


def test_cached_approval_and_threshold_in_donor_are_not_imported(tmp_path):
    master, donors = batches()
    master["items"][0]["classification"] = {"puntaje_minimo":101}
    batch, reviewed = write(tmp_path, master, donors)
    out = import_checkpoint_reviews(batch, reviewed, tmp_path/"new")
    merged = json.loads(Path(out["checkpoint"]).read_text())
    assert merged["items"][0]["classification"]["puntaje_minimo"] == 101
    assert not merged["items"][0]["classification"]["aprobado"]


def test_a_name_match_alone_is_not_a_human_confirmation(tmp_path):
    master, donors = batches(); donors["items"][0].pop("financial_audio_review")
    batch, reviewed = write(tmp_path, master, donors)
    with pytest.raises(ValueError, match="no contiene confirmaciones"):
        import_checkpoint_reviews(batch, reviewed, tmp_path/"new")


def test_different_existing_confirmation_is_preserved_and_import_rejected(tmp_path):
    master, donors = batches()
    master["items"][0]["financial_audio_review"] = deepcopy(donors["items"][0]["financial_audio_review"])
    master["items"][0]["financial_audio_review"]["balance"]["confirmed_amount"] = 1500
    batch, reviewed = write(tmp_path, master, donors); before = batch.read_bytes()
    with pytest.raises(ValueError, match="confirmación diferente"):
        import_checkpoint_reviews(batch, reviewed, tmp_path/"new")
    assert batch.read_bytes() == before


def test_reordered_reviewed_calls_are_matched_by_filename(tmp_path):
    master, donors = batches(); donors["items"].reverse()
    batch, reviewed = write(tmp_path, master, donors)
    assert import_checkpoint_reviews(batch, reviewed, tmp_path/"new")["outcomes"] == {"APROBADO":4}


def test_transfer_does_not_waive_confirmed_disclosure(tmp_path):
    from app.services.financial_review_service import make_financial_review
    from app.services.checkpoint_human_review import unreviewed_negotiation
    master, donors = batches()
    text = "Familiar: Soy su madre.\nGestor: Tiene una deuda, un saldo de 1300 pesos."
    for data in (master, donors):
        data["items"][0]["result"]["texto"] = text
        data["items"][0]["result"]["segmentos"] = [{"inicio":0,"fin":20,"texto":text}]
    i = donors["items"][0]
    i["financial_audio_review"] = make_financial_review(i["result"], unreviewed_negotiation(i), balance_amount=1600, audio_verified=True)
    batch, reviewed = write(tmp_path, master, donors)
    out = import_checkpoint_reviews(batch, reviewed, tmp_path/"new")
    merged = json.loads(Path(out["checkpoint"]).read_text())
    assert merged["items"][0]["classification"]["clasificacion"] == "FALLA_GRAVE"


def test_source_revocation_cannot_silently_remove_a_valid_confirmation(tmp_path):
    master, donors = batches()
    master["items"][0]["financial_audio_review"] = deepcopy(donors["items"][0]["financial_audio_review"])
    donors["items"][0]["financial_audio_review"] = None
    batch, reviewed = write(tmp_path, master, donors)
    with pytest.raises(ValueError, match="revocación"):
        import_checkpoint_reviews(batch, reviewed, tmp_path/"new")


def test_corrected_account_excel_is_a_copy_with_only_the_confirmed_name_changed(tmp_path):
    master, donors = batches();batch, reviewed = write(tmp_path, master, donors)
    excel = tmp_path/"accounts.xlsx"
    book = Workbook();sheet = book.active;sheet.title = "Relacion_Audios"
    sheet.append(["archivo_original","numero_de_tarjeta","nombre_final","saldo_total"])
    sheet.append(["synthetic.wav",CARD,master["items"][2]["relation"]["nombre_final"],50000])
    sheet.append(["other.wav",OTHER,"untouched",100]);book.save(excel);before = excel.read_bytes()
    out = import_checkpoint_reviews(batch, reviewed, tmp_path/"new", relation_excel=excel)
    book = load_workbook(out["corrected_excel"]);rows = list(book.active.values);book.close()
    assert rows[1] == ("synthetic.wav",CARD,"000"+CARD+"_C4_synthetic",50000)
    assert rows[2] == ("other.wav",OTHER,"untouched",100)
    assert excel.read_bytes() == before


def test_same_file_and_nonempty_destination_are_rejected(tmp_path):
    master, donors = batches();batch, reviewed = write(tmp_path, master, donors)
    with pytest.raises(ValueError, match="diferente"):
        import_checkpoint_reviews(batch, batch, tmp_path/"new")
    target = tmp_path/"existing";target.mkdir();saved = target/"keep.txt";saved.write_text("keep")
    with pytest.raises(ValueError, match="no se sobrescriben"):
        import_checkpoint_reviews(batch, reviewed, target)
    assert saved.read_text() == "keep"
