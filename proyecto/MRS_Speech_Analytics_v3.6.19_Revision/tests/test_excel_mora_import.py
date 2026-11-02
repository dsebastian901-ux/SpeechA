"""Mora real de Excel: alias, cruce de cuentas y preservación de evidencia."""
from copy import deepcopy
import json
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook

from app.services.relation_service import RelationService
from app.services.checkpoint_mora_update import fill_missing_mora
from app.services.checkpoint_reanalysis import reanalyze_checkpoint


def payload():
    return {'items':[{'file_name':'sample.ogg','result':{'archivo':'sample.ogg',
        'texto':'Habla MRS. Soy su hermana. ¿Podré dejar un recado? Sí. Que se comunique al número telefónico 5551234567.',
        'segmentos':[],'duracion_segundos':120},'relation':{
            'archivo_original':'sample.ogg','gestor':'example','estatus_grabacion':'MENSAJE CON FAMILIAR',
            'nombre_final':'0001234567890123456_REF','numero_de_tarjeta':'1234567890123456',
            'asignacion':'','fecha':'2026-10-31','pago_minimo':'1000','saldo_total':'3000'},
        'classification':{'puntaje_minimo':80}}]}


def excel(tmp_path,mora=1,*,data=None,column='mora',extra_columns=None):
    relation=deepcopy(data or payload()['items'][0]['relation'])
    relation.pop('asignacion',None)
    relation[column]=mora
    relation.update(extra_columns or {})
    wb=Workbook();s=wb.active;s.title='Relacion_Audios'
    s.append(list(relation));s.append(list(relation.values()))
    p=tmp_path/'accounts.xlsx';wb.save(p);wb.close();return p


@pytest.mark.parametrize('header',['mora','Móra','asignacion','Asignación'])
@pytest.mark.parametrize('value',[1,'C2'])
def test_reader_accepts_real_mora_header_and_legacy_assignment_header(tmp_path,header,value):
    p=excel(tmp_path,value,column=header)
    service=RelationService();info=service.load_excel(p)
    assert not info['errores']
    assert service.records[0].asignacion==str(value)


@pytest.mark.parametrize('first,second,expected',[('',1,'1'),('1PV',1,'1PV'),('C2','C2','C2')])
def test_equivalent_or_one_empty_mora_column_does_not_conflict(tmp_path,first,second,expected):
    p=excel(tmp_path,second,extra_columns={'asignacion':first})
    service=RelationService();info=service.load_excel(p)
    assert not info['errores']
    assert service.records[0].asignacion==expected


def test_conflicting_mora_columns_are_not_silently_selected(tmp_path):
    p=excel(tmp_path,1,extra_columns={'asignacion':'C2'})
    service=RelationService();info=service.load_excel(p)
    assert not service.records
    assert any('no coinciden' in r for e in info['errores'] for r in e['errores'])


def test_duplicate_mora_header_is_rejected(tmp_path):
    p=excel(tmp_path);w=load_workbook(p);s=w.active;s.cell(1,s.max_column+1).value='mora';w.save(p);w.close()
    with pytest.raises(ValueError,match='duplicada'):RelationService().load_excel(p)


def test_import_changes_only_missing_mora_and_records_verified_origin(tmp_path):
    original=payload();original['items'][0]['financial_audio_review']={'balance':{'confirmed_amount':123}}
    p=excel(tmp_path);raw=p.read_bytes();snapshot=deepcopy(original)
    updated,summary=fill_missing_mora(original,p)
    assert original==snapshot and p.read_bytes()==raw
    old=original['items'][0];new=updated['items'][0]
    assert new['relation']==dict(old['relation'],asignacion='1')
    assert new['result']==old['result'] and new['financial_audio_review']==old['financial_audio_review']
    assert summary['filled']==1 and not summary['whisper_executed']
    assert new['mora_excel_import']['canonical_value']=='1PV'
    assert not new['mora_excel_import']['new_human_audio_confirmation']


def test_prior_output_name_correction_is_preserved_instead_of_replaced_from_excel(tmp_path):
    original=payload();relation=original['items'][0]['relation'];source=dict(relation,nombre_final='0001234567890123000_REF')
    original['items'][0]['output_name_review']={'source':'EXPLICIT_USER_CHOICE_NUMERO_DE_TARJETA',
        'previous_name':source['nombre_final'],'corrected_name':relation['nombre_final'],'reviewed_at':'previous'}
    p=excel(tmp_path,data=source)
    updated,_=fill_missing_mora(original,p)
    assert updated['items'][0]['output_name_review']==original['items'][0]['output_name_review']
    assert updated['items'][0]['relation']['nombre_final']==relation['nombre_final']


@pytest.mark.parametrize('field,value',[
    ('numero_de_tarjeta','1234567890123999'),('fecha','2026-10-30'),
    ('estatus_grabacion','CONVENIO'),('pago_minimo','2000'),('saldo_total','4000'),
    ('gestor','different'),('nombre_final','0001234567890123999_REF'),
])
def test_different_account_context_is_rejected_before_creating_output(tmp_path,field,value):
    original=payload();p=excel(tmp_path,data=dict(original['items'][0]['relation'],**{field:value}))
    source=tmp_path/'checkpoint.json';source.write_text(json.dumps(original));raw=source.read_bytes();xraw=p.read_bytes()
    with pytest.raises(ValueError):reanalyze_checkpoint(source,tmp_path/'out',relation_excel=p)
    assert source.read_bytes()==raw and p.read_bytes()==xraw and not (tmp_path/'out').exists()


@pytest.mark.parametrize('mora',['','Z9',None])
def test_missing_or_unknown_source_mora_does_not_fill_the_checkpoint(tmp_path,mora):
    p=excel(tmp_path,mora)
    with pytest.raises(ValueError,match='mora vacía o no reconocida'):fill_missing_mora(payload(),p)


def test_existing_mora_is_preserved_and_repeated_import_is_idempotent(tmp_path):
    p=excel(tmp_path)
    first,_=fill_missing_mora(payload(),p)
    second,summary=fill_missing_mora(first,p)
    assert second['items']==first['items']
    assert summary['filled']==0 and summary['already_present']==1


def test_existing_mora_contradiction_is_rejected(tmp_path):
    original=payload();original['items'][0]['relation']['asignacion']='C2'
    p=excel(tmp_path)
    with pytest.raises(ValueError,match='contradice'):fill_missing_mora(original,p)


def test_duplicate_target_excel_rows_cannot_be_imported(tmp_path):
    p=excel(tmp_path);w=load_workbook(p);s=w.active;s.append([c.value for c in s[2]]);w.save(p);w.close()
    with pytest.raises(ValueError,match='fila única'):fill_missing_mora(payload(),p)


def test_mora_import_activates_existing_low_mora_visit_rule(tmp_path):
    original=payload();original['items'][0]['result']['texto']='Habla MRS. Soy yo. Gestor: Vamos a realizar una visita a su domicilio.'
    original['items'][0]['relation']['estatus_grabacion']='NO DEFINE FECHA DE PAGO'
    p=excel(tmp_path,data=original['items'][0]['relation']);source=tmp_path/'checkpoint.json';source.write_text(json.dumps(original))
    result=reanalyze_checkpoint(source,tmp_path/'out',relation_excel=p)
    item=json.loads(Path(result['checkpoint']).read_text())['items'][0]
    assert item['analytics_metadata']['negotiation_validation']['mora_payment_policy']['mora']=='1PV'
    assert any(p['id']=='PEN012' and p['invalidates'] for p in item['analytics_metadata']['penalties'])
    assert item['classification']['clasificacion']=='FALLA_GRAVE'


def test_mora_import_activates_month_cap_for_liquidation_without_retranscription(tmp_path):
    p=excel(tmp_path,'C2');original=payload();source=tmp_path/'checkpoint.json';source.write_text(json.dumps(original));raw=source.read_bytes()
    result=reanalyze_checkpoint(source,tmp_path/'out',relation_excel=p)
    item=json.loads(Path(result['checkpoint']).read_text())['items'][0]
    policy=item['analytics_metadata']['negotiation_validation']['mora_payment_policy']
    assert policy['mode']=='MES_DE_LLAMADA' and policy['window_end']=='2026-10-31'
    assert item['result']==original['items'][0]['result'] and source.read_bytes()==raw


def test_cli_imports_excel_without_whisper_or_additional_readings(tmp_path,monkeypatch,capsys):
    from tools.reanalyze_checkpoint import main
    p=excel(tmp_path);source=tmp_path/'checkpoint.json';source.write_text(json.dumps(payload()))
    monkeypatch.setattr(sys,'argv',['reanalyze_checkpoint.py','--checkpoint',str(source),
        '--relation-excel',str(p),'--output-dir',str(tmp_path/'out')])
    main();saved=json.loads((tmp_path/'out/checkpoint_reanalizado.json').read_text())
    assert saved['mora_excel_update']['filled']==1
    assert saved['items'][0]['classification']['clasificacion']=='APROBADO'
    assert 'Moras recuperadas del Excel: 1' in capsys.readouterr().out


def test_context_bound_asr_record_is_preserved_without_silently_rebinding_it(tmp_path):
    from app.services.audio_block_evidence import binding,record_digest,validate_evidence
    original=payload();item=original['items'][0]
    record={'schema':1,'kind':'AUTOMATED_ASR_ONLY','binding':binding(item['result'],item['relation']),
            'human_verified':False,'speaker_roles_confirmed':False,
            'segments':[{'text':'Lectura adicional.','start':0,'end':5}],'tail_segments':[]}
    record['record_sha256']=record_digest(record)
    assert validate_evidence(record,item['result'],item['relation'],'TERCERO')['valid']
    item['audio_block_evidence']=record
    p=excel(tmp_path);source=tmp_path/'checkpoint.json';source.write_text(json.dumps(original))
    result=reanalyze_checkpoint(source,tmp_path/'out',relation_excel=p)
    after=json.loads(Path(result['checkpoint']).read_text())['items'][0]
    assert after['audio_block_evidence']==record
    assert not after['analytics_metadata']['audio_block_evidence_validation']['valid']
    assert not after['classification']['aprobado']
