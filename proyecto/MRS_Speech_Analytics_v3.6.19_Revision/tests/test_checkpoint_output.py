"""Reproducir rutas acumuladas sin depender de habilitar rutas largas en Windows."""
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services import checkpoint_output as output
from app.services.checkpoint_amount_recovery import recover_checkpoint_amounts
from app.services.checkpoint_human_review import review_checkpoint
from app.services.checkpoint_reanalysis import reanalyze_checkpoint


@pytest.mark.parametrize("operation", ["review", "recover", "reanalyze"])
def test_repeated_operations_escape_deep_source_and_use_independent_folders(tmp_path, monkeypatch, operation):
    home = tmp_path / "usuario"
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(output, "WINDOWS", True)
    source = tmp_path / ("origen_" * 15) / ("Recuperacion_MRS_r16_" * 6) / "resultados" / "checkpoint.json"
    source.parent.mkdir(parents=True)
    payload = {"items": [{"file_name": "synthetic.wav", "result": {
        "archivo": "synthetic.wav", "duracion_segundos": 90,
        "texto": "Habla MRS. Soy su madre. Dígale que se comunique. Le paso el recado."}}]}
    source.write_text(json.dumps(payload), encoding="utf-8")
    assert len(str(source)) > 260
    raw = source.read_bytes()

    def run(path):
        if operation == "review":
            return review_checkpoint(path, {})
        if operation == "recover":
            # No se necesita ni debe ejecutarse reconocimiento para este texto.
            return recover_checkpoint_amounts(path, whisper_service=object())
        return reanalyze_checkpoint(path)

    first = run(source)
    first_bytes = Path(first["checkpoint"]).read_bytes()
    second = run(first["checkpoint"])
    base = home / "MRS_Resultados"
    folders = list(base.iterdir())
    assert len(folders) == 2
    assert all(folder.parent == base for folder in folders)
    assert first["checkpoint"] != second["checkpoint"]
    for result in (first, second):
        assert len(str(result["report"]).encode("utf-16-le")) // 2 < output.WINDOWS_PATH_BUDGET
        assert result["outcomes"] == {"APROBADO": 1}
        book = load_workbook(result["report"], read_only=True)
        assert "NOTA_REANALISIS" in book.sheetnames
        book.close()
        saved = json.loads(Path(result["checkpoint"]).read_text(encoding="utf-8"))
        assert saved["items"][0]["result"] == payload["items"][0]["result"]
        assert saved["reanalysis"]["whisper_executed"] is False
    assert source.read_bytes() == raw
    assert Path(first["checkpoint"]).read_bytes() == first_bytes


def test_long_explicit_output_is_rejected_before_creating_files(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "WINDOWS", True)
    target = tmp_path / ("carpeta_" * 20) / "resultados"
    with pytest.raises(ValueError, match="demasiado larga para Windows"):
        output.create_checkpoint_output(target, "Validacion_MRS_r17_", nested_results=True)
    assert not target.parent.exists()


def test_existing_output_is_never_overwritten(tmp_path):
    target = tmp_path / "existing"
    target.mkdir()
    file = target / "checkpoint.json"
    file.write_text("original")
    with pytest.raises(ValueError, match="no se sobrescriben"):
        output.create_checkpoint_output(target, "Validacion_MRS_r17_")
    assert file.read_text() == "original"


def test_long_user_home_is_rejected_before_creating_default_folder(tmp_path, monkeypatch):
    home = tmp_path / ("usuario_" * 24)
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: home))
    monkeypatch.setattr(output, "WINDOWS", True)
    with pytest.raises(ValueError, match="carpeta corta"):
        output.create_checkpoint_output(None, "Validacion_MRS_r17_", nested_results=True)
    assert not home.exists()


def test_windows_length_counts_utf16_units(tmp_path, monkeypatch):
    monkeypatch.setattr(output, "WINDOWS", True)
    target = tmp_path / ("😀" * 40)
    longest = target / "resultados" / "04_Reportes" / "Reporte_Reanalisis_Reglas_nuevo.xlsx"
    monkeypatch.setattr(output, "WINDOWS_PATH_BUDGET", len(str(longest)) + 1)
    with pytest.raises(ValueError, match="demasiado larga"):
        output.create_checkpoint_output(target, "Validacion_MRS_r17_", nested_results=True)
    assert not target.exists()
